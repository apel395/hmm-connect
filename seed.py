#!/usr/bin/env python3
"""Import supplier and item master data into HMM-Connect.

This script is intended for controlled EC2/admin execution, not Lambda.
It is idempotent: existing suppliers/products are updated instead of duplicated.

Examples:
    python seed.py --items import_data/Book1.xlsx --suppliers import_data/FORMAT_DATA_SUPPLIER.xlsx
    python seed.py --items import_data/Book1.xlsx --suppliers import_data/FORMAT_DATA_SUPPLIER.xlsx --apply

Without --apply the transaction is rolled back after showing a preview.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - execution guidance
    raise SystemExit(
        "openpyxl is required for this EC2-only seed utility. "
        "Install it with: python -m pip install openpyxl"
    ) from exc

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.item_group import ItemGroup
from app.models.item_subgroup import ItemSubgroup
from app.models.product import Product
from app.models.supplier import Supplier, SupplierRow


@dataclass
class Counters:
    created: int = 0
    updated: int = 0
    skipped: int = 0

    def add(self, other: "Counters") -> None:
        self.created += other.created
        self.updated += other.updated
        self.skipped += other.skipped


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", text(value).upper())


def decimal_value(value: object, default: str = "0") -> Decimal:
    raw = text(value).replace(",", "")
    if not raw:
        return Decimal(default)
    try:
        return Decimal(raw)
    except InvalidOperation:
        return Decimal(default)


def supplier_prefix(name: str) -> str:
    upper = name.upper().strip()
    if upper.startswith("PT"):
        return "PT"
    if upper.startswith("CV"):
        return "CV"
    if upper.startswith("UD"):
        return "UD"
    return "PR"


def next_supplier_code(db: Session, name: str) -> str:
    prefix = supplier_prefix(name)
    existing = [row[0] for row in db.query(Supplier.code).filter(Supplier.code.ilike(f"{prefix}%")).all()]
    maximum = 0
    for code in existing:
        match = re.fullmatch(rf"{prefix}(\d+)", (code or "").upper())
        if match:
            maximum = max(maximum, int(match.group(1)))
    return f"{prefix}{maximum + 1:04d}"


def find_supplier(db: Session, code: str, name: str) -> Supplier | None:
    if code:
        found = db.query(Supplier).filter(func.upper(Supplier.code) == code.upper()).first()
        if found:
            return found
    target = normalized(name)
    return next((supplier for supplier in db.query(Supplier).all() if normalized(supplier.name) == target), None)


def upsert_supplier(
    db: Session,
    *,
    code: str,
    name: str,
    address: str = "",
    bank_name: str = "",
    bank_account_number: str = "",
    bank_account_name: str = "",
    pic_name: str = "",
    phone: str = "",
    unit: str = "",
    item: str = "",
) -> Counters:
    counters = Counters()
    if not name:
        counters.skipped += 1
        return counters

    supplier = find_supplier(db, code, name)
    if supplier is None:
        supplier = Supplier(
            code=code or next_supplier_code(db, name),
            name=name,
            address=address or None,
            bank_name=bank_name or None,
            bank_account_number=bank_account_number or None,
            bank_account_name=bank_account_name or None,
            is_active=True,
        )
        db.add(supplier)
        db.flush()
        counters.created += 1
    else:
        changed = False
        for field, value in {
            "name": name,
            "address": address or supplier.address,
            "bank_name": bank_name or supplier.bank_name,
            "bank_account_number": bank_account_number or supplier.bank_account_number,
            "bank_account_name": bank_account_name or supplier.bank_account_name,
            "is_active": True,
        }.items():
            if value is not None and getattr(supplier, field) != value:
                setattr(supplier, field, value)
                changed = True
        counters.updated += int(changed)
        counters.skipped += int(not changed)

    if any((pic_name, phone, unit, item)):
        row_key = tuple(normalized(value) for value in (pic_name, phone, unit, item))
        existing_keys = {
            tuple(normalized(value) for value in (row.pic_name, row.phone, row.unit, row.item))
            for row in supplier.rows
        }
        if row_key not in existing_keys:
            supplier.rows.append(
                SupplierRow(
                    pic_name=pic_name or None,
                    phone=phone or None,
                    unit=unit or None,
                    item=item or None,
                )
            )
    return counters


def import_supplier_template(db: Session, path: Path) -> Counters:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["DATA SUPPLIER"] if "DATA SUPPLIER" in workbook.sheetnames else workbook.active
    counters = Counters()
    current: dict[str, str] = {}

    for row in sheet.iter_rows(min_row=7, values_only=True):
        sequence, name, address, pic, phone, unit, item, bank, account, owner = row[:10]
        if name:
            current = {
                "name": text(name),
                "address": text(address),
                "bank_name": text(bank),
                "bank_account_number": text(account),
                "bank_account_name": text(owner),
            }
        if not current.get("name"):
            continue
        if sequence is None and not any((pic, phone, unit, item)):
            continue
        counters.add(
            upsert_supplier(
                db,
                code="",
                **current,
                pic_name=text(pic),
                phone=text(phone),
                unit=text(unit),
                item=text(item),
            )
        )
    return counters


def import_book1_suppliers(db: Session, path: Path) -> Counters:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["MASTER"] if "MASTER" in workbook.sheetnames else workbook.active
    counters = Counters()
    for row in sheet.iter_rows(min_row=6, values_only=True):
        code = text(row[11] if len(row) > 11 else None)
        name = text(row[12] if len(row) > 12 else None)
        if not name:
            continue
        counters.add(
            upsert_supplier(
                db,
                code=code,
                name=name,
                phone=text(row[13] if len(row) > 13 else None),
                pic_name=text(row[14] if len(row) > 14 else None),
                address=text(row[15] if len(row) > 15 else None),
                unit=text(row[16] if len(row) > 16 else None),
                item="",
            )
        )
    return counters


def setup_codes(path: Path) -> tuple[dict[str, str], dict[tuple[str, str], str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    groups = {"FOOD": "1", "BEVERAGE": "2", "MATERIAL": "3", "ASSET": "4"}
    subgroups: dict[tuple[str, str], str] = {}
    if "SETUP" not in workbook.sheetnames:
        return groups, subgroups
    sheet = workbook["SETUP"]
    for row in sheet.iter_rows(values_only=True):
        code = text(row[5] if len(row) > 5 else None)
        name = text(row[6] if len(row) > 6 else None)
        if not re.fullmatch(r"\d+-\d+", code) or not name:
            continue
        group_code = code.split("-", 1)[0]
        group_name = next((key for key, value in groups.items() if value == group_code), group_code)
        subgroups[(normalized(group_name), normalized(name))] = code
    return groups, subgroups


def upsert_group_and_subgroup(
    db: Session,
    group_code: str,
    group_name: str,
    subgroup_code: str,
    subgroup_name: str,
) -> None:
    group = db.query(ItemGroup).filter(ItemGroup.code == group_code).first()
    if group is None:
        group = ItemGroup(code=group_code, name=group_name, is_active=True)
        db.add(group)
        db.flush()
    else:
        group.name = group_name
        group.is_active = True

    subgroup = db.query(ItemSubgroup).filter(ItemSubgroup.code == subgroup_code).first()
    if subgroup is None:
        db.add(ItemSubgroup(code=subgroup_code, name=subgroup_name, group_id=group.id, is_active=True))
    else:
        subgroup.name = subgroup_name
        subgroup.group_id = group.id
        subgroup.is_active = True


def import_items(db: Session, path: Path) -> Counters:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["MASTER"] if "MASTER" in workbook.sheetnames else workbook.active
    group_codes, subgroup_codes = setup_codes(path)
    counters = Counters()

    for row in sheet.iter_rows(min_row=6, values_only=True):
        name = text(row[2] if len(row) > 2 else None)
        code = text(row[3] if len(row) > 3 else None)
        group_name = text(row[4] if len(row) > 4 else None).upper()
        subgroup_name = text(row[5] if len(row) > 5 else None).upper()
        if not name or not code or not group_name or not subgroup_name:
            continue

        group_code = group_codes.get(group_name, normalized(group_name)[:10])
        subgroup_code = subgroup_codes.get(
            (normalized(group_name), normalized(subgroup_name)),
            f"{group_code}-{normalized(subgroup_name)[:6]}",
        )[:10]
        upsert_group_and_subgroup(db, group_code, group_name, subgroup_code, subgroup_name)

        values = {
            "name": name,
            "group_code": group_code,
            "group_name": group_name,
            "subgroup_code": subgroup_code,
            "subgroup_name": subgroup_name,
            "uom": text(row[6] if len(row) > 6 else None) or "Pcs",
            "usage_uom": text(row[7] if len(row) > 7 else None) or text(row[6]) or "Pcs",
            "conversion": decimal_value(row[8] if len(row) > 8 else None, "1"),
            "is_active": True,
        }
        product = db.query(Product).filter(func.upper(Product.code) == code.upper()).first()
        if product is None:
            db.add(Product(code=code, cost_price=0, selling_price=0, **values))
            counters.created += 1
        else:
            changed = False
            for field, value in values.items():
                if getattr(product, field) != value:
                    setattr(product, field, value)
                    changed = True
            counters.updated += int(changed)
            counters.skipped += int(not changed)
    return counters


def validate_paths(paths: Iterable[str]) -> list[Path]:
    result: list[Path] = []
    for raw in paths:
        path = Path(raw).expanduser().resolve()
        if not path.is_file():
            raise SystemExit(f"Input file not found: {path}")
        result.append(path)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Seed supplier and item master data from Excel files")
    parser.add_argument("--items", action="append", default=[], help="Item workbook path, e.g. Book1.xlsx")
    parser.add_argument("--suppliers", action="append", default=[], help="Supplier workbook path")
    parser.add_argument("--apply", action="store_true", help="Commit changes. Without this flag the run is preview-only.")
    args = parser.parse_args()

    item_files = validate_paths(args.items)
    supplier_files = validate_paths(args.suppliers)
    if not item_files and not supplier_files:
        parser.error("provide at least one --items or --suppliers file")

    db = SessionLocal()
    try:
        supplier_total = Counters()
        item_total = Counters()
        for path in item_files:
            item_total.add(import_items(db, path))
            supplier_total.add(import_book1_suppliers(db, path))
        for path in supplier_files:
            supplier_total.add(import_supplier_template(db, path))

        print(
            f"Items: created={item_total.created}, updated={item_total.updated}, "
            f"unchanged/skipped={item_total.skipped}"
        )
        print(
            f"Suppliers: created={supplier_total.created}, updated={supplier_total.updated}, "
            f"unchanged/skipped={supplier_total.skipped}"
        )
        if args.apply:
            db.commit()
            print("Seed committed successfully.")
        else:
            db.rollback()
            print("Preview complete; transaction rolled back. Re-run with --apply to commit.")
        return 0
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
