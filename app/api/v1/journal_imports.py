import base64
import io
import re
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, HTTPException
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.dependencies.authorization import require_permission
from app.models.business_unit import BusinessUnit
from app.models.chart_of_account import ChartOfAccount
from app.models.department import Department
from app.models.import_batch import ImportBatch, ImportRow
from app.models.location import Location
from app.models.supplier import Supplier
from app.models.user import User
from app.services.journal_import_jobs import (
    head_journal_import_object,
    presigned_upload_url,
    trigger_journal_import_job,
)
from app.services.location_access import assigned_location_id


router = APIRouter(prefix="/journal-imports", tags=["Journal Imports"])
journal_import_user = require_permission("journal.manage")

ALLOWED_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/csv",
    "application/csv",
    "application/octet-stream",
}

TEMPLATE_HEADERS = [
    "journal_group",
    "posting_date",
    "description",
    "location_code",
    "business_unit_code (optional)",
    "department_code (optional)",
    "account_code",
    "line_description",
    "debit",
    "credit",
    "supplier_code (optional)",
    "external_reference (optional)",
]


class JournalImportCreate(BaseModel):
    batch_name: str = Field(min_length=1, max_length=150)
    notes: str | None = None


class JournalImportUploadUrl(BaseModel):
    filename: str = Field(min_length=1, max_length=255)
    content_type: str = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class JournalImportConfirm(BaseModel):
    post_valid: bool = False


def _get_batch(db: Session, job_id: int, user: User) -> ImportBatch:
    batch = (
        db.query(ImportBatch)
        .filter(
            ImportBatch.id == job_id,
            ImportBatch.created_by_id == user.id,
            ImportBatch.import_type == "JOURNAL_ENTRY",
            ImportBatch.source == "JOURNAL_BULK_IMPORT",
        )
        .first()
    )
    if not batch:
        raise HTTPException(404, "Journal import job not found")
    return batch


def _serialize(batch: ImportBatch) -> dict:
    return {
        "id": batch.id,
        "batch_number": batch.batch_number,
        "batch_name": batch.batch_name,
        "status": batch.status,
        "source": batch.source,
        "original_file_name": batch.original_file_name,
        "total_rows": batch.total_rows,
        "valid_rows": batch.valid_rows,
        "warning_rows": batch.warning_rows,
        "error_rows": batch.error_rows,
        "imported_rows": batch.imported_rows,
        "failed_rows": batch.failed_rows,
        "total_documents": batch.total_documents,
        "success_count": batch.success_count,
        "failed_count": batch.failed_count,
        "period_from": batch.period_from,
        "period_to": batch.period_to,
        "worker_command_id": batch.worker_command_id,
        "worker_started_at": batch.worker_started_at,
        "worker_completed_at": batch.worker_completed_at,
        "worker_error_message": batch.worker_error_message,
        "created_at": batch.created_at,
        "validated_at": batch.validated_at,
        "imported_at": batch.imported_at,
        "notes": batch.notes,
    }


def _decimal(value) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    raw = re.sub(r"(?i)rp", "", str(value)).replace(" ", "")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
    elif raw.count(".") > 1:
        raw = raw.replace(".", "")
    elif raw.count(",") > 1:
        raw = raw.replace(",", "")
    elif "," in raw:
        left, right = raw.split(",", 1)
        raw = f"{left}.{right}" if len(right) in {1, 2} else left + right
    elif "." in raw:
        left, right = raw.split(".", 1)
        if len(right) == 3 and left not in {"0", "-0"}:
            raw = left + right
    try:
        return Decimal(raw)
    except InvalidOperation:
        return Decimal("0")


def _xlsx_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _master_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[tuple]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.sheet_state = "hidden"


def _add_list_validation(sheet, formula: str, cell_range: str) -> None:
    validation = DataValidation(type="list", formula1=f"={formula}", allow_blank=True)
    validation.error = "Select a value from current ERP master data."
    validation.errorTitle = "Invalid master-data code"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _build_template(db: Session, user: User) -> bytes:
    location_query = db.query(Location).filter(Location.is_active.is_(True))
    scoped_location = assigned_location_id(user)
    if scoped_location is not None:
        location_query = location_query.filter(Location.id == scoped_location)
    locations = location_query.order_by(Location.code).all()
    units = db.query(BusinessUnit).filter(BusinessUnit.is_active.is_(True)).order_by(BusinessUnit.code).all()
    departments = db.query(Department).filter(Department.is_active.is_(True)).order_by(Department.code).all()
    accounts = db.query(ChartOfAccount).filter(
        ChartOfAccount.is_active.is_(True),
        ChartOfAccount.is_posting.is_(True),
    ).order_by(ChartOfAccount.account_number).all()
    suppliers = db.query(Supplier).filter(Supplier.is_active.is_(True)).order_by(Supplier.code).all()

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Easy Bulk Journal Import"])
    instructions.append(["1", "Enter journal_group only on the first line of a journal. Blank group cells continue the journal above."])
    instructions.append(["2", "Enter date, description, location and optional dimensions only on the first line; blank cells inherit them."])
    instructions.append(["3", "Enter one account line per row. Each row must contain either debit or credit."])
    instructions.append(["4", "Upload first validates only. Review the preview, then import valid journals as Draft or Post them."])
    instructions.append(["5", "One invalid journal does not block other valid journals."])
    instructions.column_dimensions["A"].width = 8
    instructions.column_dimensions["B"].width = 110
    instructions["A1"].font = Font(bold=True, size=14)

    sheet = workbook.create_sheet("Journal Import")
    sheet.append(TEMPLATE_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:L501"
    widths = [16, 15, 30, 18, 24, 22, 18, 30, 18, 18, 22, 28]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    sheet["A1"].comment = Comment("Enter a new group only on the first row. Blank rows continue the group above.", "HMM Connect")
    sheet["B1"].comment = Comment("Accepted: YYYY-MM-DD, DD/MM/YYYY, DD-MM-YYYY, or a real Excel date.", "HMM Connect")
    sheet["I1"].comment = Comment("Indonesian and international number formats are accepted.", "HMM Connect")

    _master_sheet(workbook, "Master Locations", ["code", "name", "business_unit"], [
        (row.code, row.name, row.business_unit.code if row.business_unit else "") for row in locations
    ])
    _master_sheet(workbook, "Master Business Units", ["code", "name"], [(row.code, row.name) for row in units])
    _master_sheet(workbook, "Master Departments", ["code", "name"], [(row.code, row.name) for row in departments])
    _master_sheet(workbook, "Master Accounts", ["code", "name"], [(row.account_number, row.account_name) for row in accounts])
    _master_sheet(workbook, "Master Suppliers", ["code", "name"], [(row.code, row.name) for row in suppliers])

    ranges = {
        "D2:D501": ("LocationCodes", "Master Locations", len(locations)),
        "E2:E501": ("BusinessUnitCodes", "Master Business Units", len(units)),
        "F2:F501": ("DepartmentCodes", "Master Departments", len(departments)),
        "G2:G501": ("AccountCodes", "Master Accounts", len(accounts)),
        "K2:K501": ("SupplierCodes", "Master Suppliers", len(suppliers)),
    }
    for cell_range, (range_name, master, count) in ranges.items():
        if count:
            workbook.defined_names.add(DefinedName(
                range_name,
                attr_text=f"'{master}'!$A$2:$A${count + 1}",
            ))
            _add_list_validation(sheet, range_name, cell_range)

    red_fill = PatternFill("solid", fgColor="FEE2E2")
    sheet.conditional_formatting.add("I2:J501", FormulaRule(formula=["AND($I2>0,$J2>0)"], fill=red_fill))
    return _xlsx_bytes(workbook)


@router.get("/template")
def template(
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    payload = _build_template(db, user)
    return {
        "filename": "hmm-journal-import-easy-template.xlsx",
        "content_base64": base64.b64encode(payload).decode("ascii"),
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "format": "XLSX",
    }


@router.post("", status_code=201)
def create_job(
    data: JournalImportCreate,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    from uuid import uuid4

    batch = ImportBatch(
        batch_number=f"JIMP-{uuid4().hex[:12].upper()}",
        source="JOURNAL_BULK_IMPORT",
        description=data.notes,
        batch_name=data.batch_name,
        import_type="JOURNAL_ENTRY",
        import_mode="PREVIEW",
        source_type="JOURNAL_BULK_IMPORT",
        status="DRAFT",
        created_by_id=user.id,
        imported_by_id=user.id,
        notes=data.notes,
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return _serialize(batch)


@router.post("/{job_id}/upload-url")
def upload_url(
    job_id: int,
    data: JournalImportUploadUrl,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    if batch.status not in {"DRAFT", "UPLOADED"}:
        raise HTTPException(409, "This import job is no longer accepting a file")
    suffix = data.filename.lower().rsplit(".", 1)[-1] if "." in data.filename else ""
    if suffix not in {"xlsx", "csv"}:
        raise HTTPException(415, "Only XLSX and CSV files are supported")
    if data.content_type not in ALLOWED_TYPES:
        raise HTTPException(415, "Unsupported upload content type")
    key, url = presigned_upload_url(
        user_id=user.id,
        job_id=job_id,
        filename=data.filename,
        content_type=data.content_type,
    )
    batch.file_name = data.filename
    batch.original_file_name = data.filename
    batch.stored_file_path = key
    db.commit()
    return {
        "job_id": job_id,
        "upload_url": url,
        "object_key": key,
        "expires_in": settings.JOURNAL_IMPORT_UPLOAD_URL_EXPIRES_SECONDS,
        "content_type": data.content_type,
    }


@router.post("/{job_id}/complete")
def complete_upload(
    job_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    if batch.status not in {"DRAFT", "UPLOADED"}:
        raise HTTPException(409, "This import job cannot be completed")
    if not batch.stored_file_path:
        raise HTTPException(409, "Upload URL has not been created")
    try:
        metadata = head_journal_import_object(batch.stored_file_path)
    except Exception as exc:
        raise HTTPException(409, "Uploaded file was not found in storage") from exc
    size = int(metadata.get("ContentLength") or 0)
    maximum = settings.JOURNAL_IMPORT_MAX_SIZE_MB * 1024 * 1024
    if size <= 0:
        raise HTTPException(409, "Uploaded file is empty")
    if size > maximum:
        raise HTTPException(413, f"Journal import file exceeds {settings.JOURNAL_IMPORT_MAX_SIZE_MB} MB")
    batch.status = "UPLOADED"
    batch.worker_error_message = None
    db.commit()
    db.refresh(batch)
    return _serialize(batch)


@router.post("/{job_id}/start", status_code=202)
def start_validation(
    job_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    if batch.status != "UPLOADED":
        raise HTTPException(409, "Complete the file upload before validation")
    batch.status = "QUEUED"
    batch.worker_error_message = None
    db.commit()
    try:
        command_id = trigger_journal_import_job(job_id, mode="validate")
    except Exception as exc:
        batch.status = "FAILED"
        batch.worker_error_message = str(exc)[:2000]
        batch.worker_completed_at = datetime.now(timezone.utc)
        db.commit()
        raise HTTPException(503, "Journal validation worker could not be triggered") from exc
    batch.worker_command_id = command_id
    db.commit()
    db.refresh(batch)
    return _serialize(batch)


@router.post("/{job_id}/confirm", status_code=202)
def confirm_import(
    job_id: int,
    data: JournalImportConfirm,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    if batch.status != "VALIDATED":
        raise HTTPException(409, "Validate and review this import before confirming")
    valid_rows = db.query(ImportRow).filter(
        ImportRow.import_batch_id == batch.id,
        ImportRow.status == "VALID",
    ).count()
    if not valid_rows:
        raise HTTPException(409, "There are no valid journals to import")
    batch.status = "IMPORT_QUEUED"
    batch.worker_error_message = None
    db.commit()
    try:
        command_id = trigger_journal_import_job(job_id, mode="import", post_valid=data.post_valid)
    except Exception as exc:
        batch.status = "VALIDATED"
        batch.worker_error_message = str(exc)[:2000]
        db.commit()
        raise HTTPException(503, "Journal import worker could not be triggered") from exc
    batch.worker_command_id = command_id
    db.commit()
    db.refresh(batch)
    return _serialize(batch)


@router.get("/{job_id}/preview")
def get_preview(
    job_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    rows = db.query(ImportRow).filter(ImportRow.import_batch_id == batch.id).order_by(ImportRow.id).all()
    groups = defaultdict(list)
    for row in rows:
        data = row.normalized_data_json or {}
        groups[data.get("_group_key") or row.document_number or str(row.id)].append(row)
    result = []
    for group_rows in groups.values():
        first = group_rows[0].normalized_data_json or {}
        debit = sum((_decimal((row.normalized_data_json or {}).get("debit")) for row in group_rows), Decimal("0"))
        credit = sum((_decimal((row.normalized_data_json or {}).get("credit")) for row in group_rows), Decimal("0"))
        errors = list(dict.fromkeys(message for row in group_rows for message in (row.error_messages_json or [])))
        result.append({
            "journal_group": first.get("_group_label"),
            "posting_date": first.get("posting_date"),
            "description": first.get("description"),
            "location": first.get("location_code") or first.get("location"),
            "line_count": len(group_rows),
            "debit": debit,
            "credit": credit,
            "difference": debit - credit,
            "status": "ERROR" if errors else "VALID",
            "errors": errors,
        })
    return {"job_id": batch.id, "status": batch.status, "journals": result}


@router.get("/{job_id}/errors")
def get_errors(
    job_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    rows = db.query(ImportRow).filter(
        ImportRow.import_batch_id == batch.id,
        ImportRow.status.in_({"ERROR", "FAILED"}),
    ).order_by(ImportRow.row_number).limit(1000).all()
    return {
        "job_id": batch.id,
        "rows": [{
            "row_number": row.row_number,
            "sheet_name": row.sheet_name,
            "document_number": row.document_number,
            "journal_group": (row.normalized_data_json or {}).get("_group_label"),
            "status": row.status,
            "errors": row.error_messages_json or [],
            "warnings": row.warning_messages_json or [],
            "data": row.normalized_data_json or {},
        } for row in rows],
    }


@router.get("/{job_id}/correction-file")
def correction_file(
    job_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    batch = _get_batch(db, job_id, user)
    rows = db.query(ImportRow).filter(
        ImportRow.import_batch_id == batch.id,
        ImportRow.status.in_({"ERROR", "FAILED"}),
    ).order_by(ImportRow.id).all()
    if not rows:
        raise HTTPException(404, "This import has no rejected rows")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Correction"
    headers = [*TEMPLATE_HEADERS, "validation_errors"]
    sheet.append(headers)
    keys = ["journal_group", "posting_date", "description", "location_code", "business_unit_code", "department_code", "account_code", "line_description", "debit", "credit", "supplier_code", "external_reference"]
    previous_group = None
    for row in rows:
        data = row.normalized_data_json or {}
        group = data.get("_group_label")
        output_group = group if group != previous_group else ""
        previous_group = group
        values = [output_group, *[data.get(key, "") for key in keys[1:]]]
        values.append(" | ".join(row.error_messages_json or []))
        sheet.append(values)
    sheet.freeze_panes = "A2"
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22 if index != len(headers) else 70
    payload = _xlsx_bytes(workbook)
    return {
        "filename": f"{batch.batch_number}-corrections.xlsx",
        "content_base64": base64.b64encode(payload).decode("ascii"),
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


@router.get("/{job_id}")
def get_job(
    job_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(journal_import_user),
):
    return _serialize(_get_batch(db, job_id, user))
