from datetime import date
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.dependencies.authorization import require_permission
from app.models.accounting_setting import AccountingSetting
from app.models.chart_of_account import ChartOfAccount
from app.models.financial_report_setting import FinancialReportSetting
from app.services.audit import record_audit
from app.services.financial_reporting_service import FinancialReportingService
from app.services.location_access import resolve_location_id
from app.api.v1.profit_loss import serialize

router = APIRouter(prefix="/financial-reports", tags=["Financial Reports"])
view = require_permission("finance.view")
manage = require_permission("finance.manage")
TYPES = "^(PROFIT_LOSS|BALANCE_SHEET)$"


class ReportSetupInput(BaseModel):
    selected_account_ids: list[int] | None = None
    header_only: bool = False
    current_year_earnings_account_id: int | None = None


@router.get("/setup")
def setup(db: Session = Depends(get_db), user=Depends(view)):
    rows = db.query(ChartOfAccount).filter(ChartOfAccount.is_active.is_(True)).order_by(ChartOfAccount.account_number).all()
    settings = db.query(FinancialReportSetting).all()
    earnings = db.query(AccountingSetting).filter_by(setting_key="CURRENT_YEAR_EARNINGS_ACCOUNT").first()
    return {"accounts": [{"id": a.id, "code": a.account_number, "name": a.account_name,
             "type": a.account_type, "normal_balance": a.normal_balance, "parent_id": a.parent_account_id, "is_posting": a.is_posting} for a in rows],
            "settings": [{"report_type": s.report_type, "selected_account_ids": s.selected_account_ids, "header_only": s.header_only} for s in settings],
            "current_year_earnings_account_id": earnings.account_id if earnings else None}


@router.put("/setup/{report_type}")
def save_setup(data: ReportSetupInput, report_type: str, db: Session = Depends(get_db), user=Depends(manage)):
    if report_type not in ("PROFIT_LOSS", "BALANCE_SHEET"):
        raise HTTPException(422, "Unsupported report type")
    if data.selected_account_ids is not None:
        ids = set(data.selected_account_ids)
        allowed = ("REVENUE", "INCOME", "EXPENSE") if report_type == "PROFIT_LOSS" else ("ASSET", "LIABILITY", "EQUITY")
        found = db.query(ChartOfAccount.id).filter(ChartOfAccount.id.in_(ids), ChartOfAccount.account_type.in_(allowed)).all()
        if len(found) != len(ids):
            raise HTTPException(422, "Selected accounts do not belong to this report")
    if data.current_year_earnings_account_id is not None:
        account = db.get(ChartOfAccount, data.current_year_earnings_account_id)
        if not account or not account.is_active or not account.is_posting or account.account_type != "EQUITY" or account.normal_balance != "CREDIT":
            raise HTTPException(422, "Current-year earnings requires an active credit equity account")
        mapping = db.query(AccountingSetting).filter_by(setting_key="CURRENT_YEAR_EARNINGS_ACCOUNT").first()
        if not mapping:
            mapping = AccountingSetting(setting_key="CURRENT_YEAR_EARNINGS_ACCOUNT")
            db.add(mapping)
        mapping.account_id = account.id
    setting = db.query(FinancialReportSetting).filter_by(report_type=report_type).with_for_update().first()
    if not setting:
        setting = FinancialReportSetting(report_type=report_type)
        db.add(setting)
    before = {"selected_account_ids": setting.selected_account_ids, "header_only": setting.header_only}
    setting.selected_account_ids = sorted(set(data.selected_account_ids)) if data.selected_account_ids is not None else None
    setting.header_only, setting.updated_by_id = data.header_only, user.id
    record_audit(db, user, "FINANCIAL_REPORT_SETUP_UPDATED", "FINANCIAL_REPORT", f"Updated {report_type} display",
                 before_data=before, after_data=data.model_dump())
    db.commit()
    return {"message": "Report setup saved"}


@router.get("/comparison")
def comparison(report_type: str = Query(pattern=TYPES), as_of_date: date = Query(),
               location_id: int | None = None, business_unit_id: int | None = None,
               department_id: int | None = None, budget_id: int | None = None,
               kap: bool = False, include_zero: bool = False,
               db: Session = Depends(get_db), user=Depends(view)):
    scope = resolve_location_id(user, location_id)
    return serialize(FinancialReportingService(db).build(report_type=report_type, as_of_date=as_of_date,
        location_id=scope, business_unit_id=business_unit_id, department_id=department_id,
        budget_id=budget_id, kap=kap, include_zero=include_zero))


@router.get("/comparison/export")
def export_comparison(report_type: str = Query(pattern=TYPES), as_of_date: date = Query(),
                      location_id: int | None = None, business_unit_id: int | None = None,
                      department_id: int | None = None, budget_id: int | None = None,
                      kap: bool = False, include_zero: bool = False,
                      db: Session = Depends(get_db), user=Depends(view)):
    scope = resolve_location_id(user, location_id)
    report = FinancialReportingService(db).build(report_type=report_type, as_of_date=as_of_date,
        location_id=scope, business_unit_id=business_unit_id, department_id=department_id,
        budget_id=budget_id, kap=kap, include_zero=include_zero)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for period in ("MTD", "YTD"):
        sheet = workbook.create_sheet(period)
        sheet.append([report_type, as_of_date.isoformat(), "Location", scope or "Consolidated", "Unit", business_unit_id or "All"])
        metrics = ("current", "budget", "last_month", "last_year", "variance") + (("movement",) if report_type == "BALANCE_SHEET" else ())
        sheet.append(["Account", "Name", "Current", "Budget", "Last month", "Same period last year", "Variance"] + (["Period movement"] if report_type == "BALANCE_SHEET" else []))
        for section, totals in report["section_totals"].items():
            values = totals[period]
            sheet.append([section, "Total", *[float(values[m]) if values[m] is not None else None for m in metrics]])
            for row in report["rows"]:
                if row["section"] != section:
                    continue
                values = row["periods"][period]
                sheet.append([row["account_code"], row["account_name"], *[float(values[m]) if values[m] is not None else None for m in metrics]])
        if report_type == "PROFIT_LOSS":
            v = report["section_totals"]
            sheet.append(["NET", "Net profit", *[None if any(v[s][period][m] is None for s in v) else float(v["REVENUE"][period][m] - v["COST_OF_REVENUE"][period][m] - v["EXPENSES"][period][m]) for m in metrics]])
        for cell in sheet[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E293B")
        sheet.freeze_panes = "C3"
        sheet.auto_filter.ref = f"A2:{'H' if report_type == 'BALANCE_SHEET' else 'G'}{sheet.max_row}"
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 48
        for column in "CDEFGH":
            sheet.column_dimensions[column].width = 24
            for cell in sheet[column][2:]:
                cell.number_format = '#,##0.00;[Red](#,##0.00)'
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{report_type.lower()}-{as_of_date}.xlsx"'})
