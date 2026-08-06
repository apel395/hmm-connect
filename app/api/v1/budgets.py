import csv,io
from datetime import date,datetime,timezone
from decimal import Decimal
from math import ceil
from fastapi import APIRouter,Depends,File,Form,HTTPException,Query,UploadFile
from fastapi.responses import Response
from openpyxl import Workbook,load_workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.dependencies.auth import get_current_user
from app.models.budget import Budget,BudgetLine
from app.models.chart_of_account import ChartOfAccount
from app.models.department import Department
from app.models.location import Location
from app.models.business_unit import BusinessUnit
from app.schemas.budget import BudgetCreate,BudgetUpdate,RejectInput,BudgetLineInput
from app.services.audit import record_audit
from app.services.budget_service import BudgetService,BudgetVsActualService,ZERO
budgets_router=APIRouter(prefix="/budgets",tags=["Budgets"]);reports_router=APIRouter(prefix="/reports/budget-vs-actual",tags=["Budget vs Actual"])
def access(u):
 if not {r.name for r in u.roles}.intersection({"FINANCE","ADMIN"}):raise HTTPException(403,"Finance access required")
def audit(db,u,a,x):record_audit(db,u,a,"BUDGET",f"{a.replace('_',' ').title()} {x.budget_number}",entity_id=x.id,reference_document=x.budget_number)
@budgets_router.get("")
def listing(page:int=1,page_size:int=20,search:str|None=None,fiscal_year:int|None=None,status:str|None=None,business_unit_id:int|None=None,department_id:int|None=None,location_id:int|None=None,account_id:int|None=None,db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user);q=db.query(Budget)
 if search:q=q.filter(or_(Budget.budget_number.ilike(f"%{search}%"),Budget.name.ilike(f"%{search}%")))
 if fiscal_year:q=q.filter(Budget.fiscal_year==fiscal_year)
 if status:q=q.filter(Budget.status==status)
 if any((business_unit_id,department_id,location_id,account_id)):q=q.join(Budget.lines)
 if business_unit_id:q=q.filter(BudgetLine.business_unit_id==business_unit_id)
 if department_id:q=q.filter(BudgetLine.department_id==department_id)
 if location_id:q=q.filter(BudgetLine.location_id==location_id)
 if account_id:q=q.filter(BudgetLine.account_id==account_id)
 total=q.distinct().count();svc=BudgetService(db);return {"items":[svc.serialize(x) for x in q.distinct().order_by(Budget.id.desc()).offset((page-1)*page_size).limit(page_size)],"page":page,"page_size":page_size,"total_items":total,"total_pages":ceil(total/page_size) if total else 0}
@budgets_router.post("")
def create(data:BudgetCreate,db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user);svc=BudgetService(db)
 try:x=svc.create(data,user)
 except ValueError as e:raise HTTPException(400,str(e))
 audit(db,user,"BUDGET_CREATED",x);db.commit();db.refresh(x);return svc.serialize(x)
@budgets_router.get("/availability")
def availability(account_id:int,transaction_date:date,amount:Decimal,business_unit_id:int,department_id:int|None=None,location_id:int|None=None,db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user);r=BudgetVsActualService(db).report(fiscal_year=transaction_date.year,date_from=transaction_date.replace(day=1),date_to=transaction_date,account_id=account_id,department_id=department_id,location_id=location_id,business_unit_id=business_unit_id,include_zero=True);x=r["items"][0] if r["items"] else None;available=Decimal(x["available_after_commitments"]) if x else ZERO;return {"matching_budget":None,"matching_budget_line":x,"budget_amount":x["budget_amount"] if x else ZERO,"actual_amount":x["actual_amount"] if x else ZERO,"committed_amount":x["committed_po"] if x else ZERO,"available_amount":available,"requested_amount":amount,"projected_remaining":available-amount,"status":"NO_BUDGET" if not x else "OVER_BUDGET" if available<amount else x["status"]}
@budgets_router.get("/template")
def template(user=Depends(get_current_user)):
 access(user);out=io.StringIO();csv.writer(out).writerow(["fiscal_year","month","business_unit_code","location_code","department_code","account_code","budget_amount","notes"]);return Response(out.getvalue(),media_type="text/csv",headers={"Content-Disposition":'attachment; filename="budget-template.csv"'})
def rows(content,name):
 if name.lower().endswith(".csv"):return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
 if name.lower().endswith(".xlsx"):
  data=list(load_workbook(io.BytesIO(content),read_only=True,data_only=True).active.iter_rows(values_only=True));return [dict(zip([str(x or "") for x in data[0]],r)) for r in data[1:]]
 raise ValueError("Only CSV and XLSX files are supported")
@budgets_router.post("/import")
async def import_budget(file:UploadFile=File(),preview:bool=Form(True),name:str|None=Form(None),db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user)
 try:raw=rows(await file.read(),file.filename or "budget.csv")
 except ValueError as e:raise HTTPException(415,str(e))
 parsed=[];errors=[];seen=set();years=set()
 for n,r in enumerate(raw,2):
  try:
   year=int(r["fiscal_year"]);month=int(r["month"]);a=db.query(ChartOfAccount).filter(ChartOfAccount.account_number==str(r["account_code"]),ChartOfAccount.is_active.is_(True)).one();bu_code=str(r["business_unit_code"]).strip();bu=db.query(BusinessUnit).filter(BusinessUnit.code==bu_code,BusinessUnit.is_active.is_(True)).one();department_code=str(r.get("department_code") or "").strip();location_code=str(r.get("location_code") or "").strip();d=db.query(Department).filter(Department.code==department_code).first() if department_code else None;l=db.query(Location).filter(Location.code==location_code).first() if location_code else None
   if department_code and not d:raise ValueError(f"Unknown department code {department_code}")
   if location_code and not l:raise ValueError(f"Unknown location code {location_code}")
   if l and l.business_unit_id!=bu.id:raise ValueError("Location does not belong to Business Unit")
   if d and d.business_unit_id not in (None,bu.id):raise ValueError("Department does not belong to Business Unit")
   key=(a.id,d.id if d else None,l.id if l else None,month,bu.id)
   if key in seen:raise ValueError("Duplicate line dimensions")
   seen.add(key);years.add(year);parsed.append(BudgetLineInput(account_id=a.id,business_unit_id=bu.id,department_id=d.id if d else None,location_id=l.id if l else None,month=month,budget_amount=Decimal(str(r["budget_amount"])),notes=r.get("notes")))
  except Exception as e:errors.append({"row":n,"error":str(e)})
 result={"rows_detected":len(raw),"rows_accepted":len(parsed),"errors":errors}
 if preview or errors:return result
 if len(years)!=1:raise HTTPException(400,"Import requires one fiscal year")
 y=next(iter(years));x=BudgetService(db).create(BudgetCreate(name=name or file.filename or f"Imported Budget {y}",fiscal_year=y,date_from=date(y,1,1),date_to=date(y,12,31),lines=parsed),user);audit(db,user,"BUDGET_CREATED",x);db.commit();return {**result,"budget_id":x.id,"budget_number":x.budget_number}
@budgets_router.get("/{id}")
def detail(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user)
 try:return BudgetService(db).serialize(BudgetService(db).get(id))
 except ValueError as e:raise HTTPException(404,str(e))
@budgets_router.get("/{id}/allowed-actions")
def allowed(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):access(user);return BudgetService.actions(BudgetService(db).get(id))
@budgets_router.patch("/{id}")
def update(id:int,data:BudgetUpdate,db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user);svc=BudgetService(db);x=svc.get(id)
 if x.status!="DRAFT":raise HTTPException(409,"Only draft budgets can be edited")
 for k,v in data.model_dump(exclude_unset=True,exclude={"lines"}).items():setattr(x,k,v)
 if data.lines is not None:
  check=BudgetCreate(name=x.name,fiscal_year=x.fiscal_year,date_from=x.date_from,date_to=x.date_to,currency=x.currency,description=x.description,lines=data.lines);svc.validate(check);x.lines=[BudgetLine(**l.model_dump()) for l in data.lines]
 audit(db,user,"BUDGET_UPDATED",x);db.commit();return svc.serialize(x)
def transition(id,target,source,action,db,user):
 access(user);svc=BudgetService(db);x=svc.get(id)
 if x.status not in source:raise HTTPException(409,"Action is not allowed")
 x.status=target;now=datetime.now(timezone.utc)
 if target=="APPROVED":x.approved_by_id=user.id;x.approved_at=now
 if target=="SUBMITTED":x.submitted_at=now
 if target=="LOCKED":x.locked_at=now
 if target=="CANCELLED":x.cancelled_at=now
 audit(db,user,action,x);db.commit();return svc.serialize(x)
@budgets_router.post("/{id}/submit")
def submit(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):return transition(id,"SUBMITTED",{"DRAFT"},"BUDGET_SUBMITTED",db,user)
@budgets_router.post("/{id}/approve")
def approve(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):return transition(id,"APPROVED",{"SUBMITTED"},"BUDGET_APPROVED",db,user)
@budgets_router.post("/{id}/reject")
def reject(id:int,data:RejectInput,db:Session=Depends(get_db),user=Depends(get_current_user)):return transition(id,"REJECTED",{"SUBMITTED"},"BUDGET_REJECTED",db,user)
@budgets_router.post("/{id}/lock")
def lock(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):return transition(id,"LOCKED",{"APPROVED"},"BUDGET_LOCKED",db,user)
@budgets_router.post("/{id}/cancel")
def cancel(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):return transition(id,"CANCELLED",{"DRAFT","SUBMITTED","APPROVED"},"BUDGET_CANCELLED",db,user)
@budgets_router.get("/{id}/print")
def printable(id:int,db:Session=Depends(get_db),user=Depends(get_current_user)):return {**detail(id,db,user),"document_type":"Budget","printed_at":datetime.now(timezone.utc)}
@reports_router.get("")
def report(budget_id:int|None=None,fiscal_year:int|None=None,date_from:date|None=None,date_to:date|None=None,account_id:int|None=None,account_type:str|None=None,department_id:int|None=None,location_id:int|None=None,business_unit_id:int|None=None,include_zero:bool=False,page:int=1,page_size:int=50,search:str|None=None,group_by:str="ACCOUNT",db:Session=Depends(get_db),user=Depends(get_current_user)):access(user);return BudgetVsActualService(db).report(budget_id,fiscal_year,date_from,date_to,account_id,account_type,department_id,location_id,group_by,include_zero,search,page,page_size,business_unit_id)
@reports_router.get("/export")
def export(format:str="xlsx",fiscal_year:int|None=None,group_by:str="ACCOUNT",db:Session=Depends(get_db),user=Depends(get_current_user)):
 access(user);items=BudgetVsActualService(db).report(fiscal_year=fiscal_year,group_by=group_by,include_zero=True,page_size=100000)["items"];heads=["Group","Budget","Planned","Committed","Payable","Actual","Remaining","Available","Variance","Utilization %","Status"];matrix=[[x["group_label"],x["budget_amount"],x["planned_pr"],x["committed_po"],x["payable_invoice"],x["actual_amount"],x["remaining_budget"],x["available_after_commitments"],x["variance_amount"],x["utilization_percentage"],x["status"]] for x in items]
 if format=="csv":out=io.StringIO();w=csv.writer(out);w.writerow(heads);w.writerows(matrix);return Response(out.getvalue(),media_type="text/csv",headers={"Content-Disposition":'attachment; filename="budget-vs-actual.csv"'})
 wb=Workbook();ws=wb.active;ws.append(heads)
 for row in matrix:ws.append(row)
 out=io.BytesIO();wb.save(out);return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="budget-vs-actual.xlsx"'})
