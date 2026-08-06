import csv
from datetime import date, datetime, timezone
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.dependencies.auth import get_current_user
from app.services.trial_balance_service import TrialBalanceService


router = APIRouter(prefix="/trial-balance", tags=["Financial Reports"])
MONEY_FIELDS = (
    "opening_debit",
    "opening_credit",
    "period_debit",
    "period_credit",
    "closing_debit",
    "closing_credit",
    "difference",
)


def access(user):
    if not {role.name for role in user.roles}.intersection({"FINANCE", "ADMIN"}):
        raise HTTPException(403, "Finance report access required")


def serialize(report: dict) -> dict:
    for row in report["items"]:
        for field in MONEY_FIELDS[:-1]:
            row[field] = f"{row[field]:.2f}"
    for field in MONEY_FIELDS:
        if field in report["summary"]:
            report["summary"][field] = f"{report['summary'][field]:.2f}"
    return report


def build_report(
    db,
    date_from,
    date_to,
    page,
    page_size,
    search,
    account_type,
    account_id,
    parent_account_id,
    include_zero_balance,
    include_inactive,
    location_id,
    department_id,
    business_unit_id,
    sort_by,
    sort_order,
    paginate=True,
):
    try:
        return TrialBalanceService(db).report(
            date_from=date_from,
            date_to=date_to,
            page=page,
            page_size=page_size,
            search=search,
            account_type=account_type,
            account_id=account_id,
            parent_account_id=parent_account_id,
            include_zero_balance=include_zero_balance,
            include_inactive=include_inactive,
            location_id=location_id,
            department_id=department_id,
            business_unit_id=business_unit_id,
            sort_by=sort_by,
            sort_order=sort_order,
            paginate=paginate,
        )
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.get("")
def get_trial_balance(
    date_from: date,
    date_to: date,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    search: str | None = None,
    account_type: str | None = None,
    account_id: int | None = None,
    parent_account_id: int | None = None,
    include_zero_balance: bool = False,
    include_inactive: bool = False,
    location_id: int | None = None,
    department_id: int | None = None,
    business_unit_id: int | None = None,
    sort_by: str = "account_code",
    sort_order: str = Query("asc", pattern="^(asc|desc)$"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    access(user)
    return serialize(
        build_report(
            db, date_from, date_to, page, page_size, search, account_type, account_id,
            parent_account_id, include_zero_balance, include_inactive, location_id,
            department_id, business_unit_id, sort_by, sort_order,
        )
    )


@router.get("/export")
def export_trial_balance(
    date_from: date,
    date_to: date,
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    search: str | None = None,
    account_type: str | None = None,
    account_id: int | None = None,
    parent_account_id: int | None = None,
    include_zero_balance: bool = False,
    include_inactive: bool = False,
    location_id: int | None = None,
    department_id: int | None = None,
    business_unit_id: int | None = None,
    sort_by: str = "account_code",
    sort_order: str = Query("asc", pattern="^(asc|desc)$"),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    access(user)
    report = build_report(
        db, date_from, date_to, 1, 200, search, account_type, account_id,
        parent_account_id, include_zero_balance, include_inactive, location_id,
        department_id, business_unit_id, sort_by, sort_order, paginate=False,
    )
    headers = [
        "Account Code", "Account Name", "Account Type", "Normal Balance",
        "Opening Debit", "Opening Credit", "Period Debit", "Period Credit",
        "Closing Debit", "Closing Credit", "Abnormal Balance",
    ]
    fields = [
        "account_code", "account_name", "account_type", "normal_balance",
        "opening_debit", "opening_credit", "period_debit", "period_credit",
        "closing_debit", "closing_credit", "is_abnormal_balance",
    ]
    filename = f"trial-balance-{date_from}-{date_to}.{format}"
    if format == "csv":
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in report["items"]:
            writer.writerow([row[field] for field in fields])
        summary = report["summary"]
        writer.writerow([
            "TOTAL", "", "", "", summary["opening_debit"], summary["opening_credit"],
            summary["period_debit"], summary["period_credit"], summary["closing_debit"],
            summary["closing_credit"], "",
        ])
        return StreamingResponse(
            iter([output.getvalue()]), media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trial Balance"
    sheet.append(["HMM-Connect — Trial Balance"])
    sheet.append([f"Period: {date_from} to {date_to}"])
    sheet.append([f"Generated: {datetime.now(timezone.utc).isoformat()}"])
    filters = f"Search={search or 'All'}; Account Type={account_type or 'All'}; Account={account_id or 'All'}; Parent={parent_account_id or 'All'}; Location={location_id or 'All'}; Department={department_id or 'All'}"
    sheet.append([f"Filters: {filters}"])
    sheet.append([])
    sheet.append(headers)
    header_row = 6
    thin = Side(style="thin", color="D9E0E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F1F5F9")
        cell.border = border
    for row in report["items"]:
        sheet.append([row[field] for field in fields])
    summary = report["summary"]
    sheet.append([
        "TOTAL", "", "", "", summary["opening_debit"], summary["opening_credit"],
        summary["period_debit"], summary["period_credit"], summary["closing_debit"],
        summary["closing_credit"], "",
    ])
    total_row = sheet.max_row
    for row in sheet.iter_rows(min_row=header_row, max_row=total_row):
        for cell in row:
            cell.border = border
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=total_row, min_col=5, max_col=10):
        for cell in row:
            cell.number_format = '#,##0.00'
    sheet.freeze_panes = "A7"
    for column, width in {"A": 16, "B": 32, "C": 16, "D": 16, "E": 18, "F": 18, "G": 18, "H": 18, "I": 18, "J": 18, "K": 18}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
