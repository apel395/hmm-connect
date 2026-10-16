import argparse
import time
from datetime import datetime, timezone
from pathlib import Path

import boto3
from openpyxl import Workbook

from app.core.config import settings
from app.core.database import SessionLocal
from app.models.export_job import ExportJob
from app.models.product import Product
from app.models.user import User
from app.services.general_ledger_report import (
    build_ledger_query,
    ledger_rows_with_balance,
    ledger_totals,
)
from app.services.location_access import resolve_location_id
from app.services.item_reference_cost import receipt_costs, reference_payload


MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _generate_products_xlsx(db, output_path: Path, params=None, user=None) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Products")
    ws.append([
        "Code",
        "Name",
        "Group Code",
        "Group Name",
        "Subgroup Code",
        "Subgroup Name",
        "Purchase UOM",
        "Usage UOM",
        "Conversion",
        "Reference Cost",
        "Reference Method",
        "Catalog",
        "Selling Price",
        "Minimum Stock",
        "Reorder Level",
        "Maximum Stock",
        "Lead Time Days",
        "Tax Category",
        "Taxable",
        "Tax Rate",
        "Active", "Item Type", "Department",
    ])

    catalog = (params or {}).get("catalog_type")
    if catalog not in {None, "CATERING", "HOTEL"}:
        raise ValueError("Unknown item catalog")
    item_type = (params or {}).get("item_type")
    if item_type not in {None, "BUY", "SELL"}:
        raise ValueError("Unknown item type")
    query = db.query(Product)
    if item_type:
        query = query.filter(Product.item_type == item_type)
    if catalog:
        query = query.filter(Product.catalog_type == catalog)
    # Read price history in batches, so exports do not issue one query per item.
    last_id = 0
    while True:
        rows = query.filter(Product.id > last_id).order_by(Product.id).limit(500).all()
        if not rows:
            break
        history = receipt_costs(db, [row.id for row in rows], user) if user else {}
        from app.services.item_departments import payloads
        departments = payloads(db, [row.id for row in rows])
        for row in rows:
            ref = reference_payload(row, history.get(row.id))
            ws.append([
                row.code, row.name, row.group_code, row.group_name,
                row.subgroup_code, row.subgroup_name, row.uom, row.usage_uom,
                float(row.conversion or 0), float(ref["reference_cost"]),
                row.reference_cost_method, row.catalog_type,
                float(row.selling_price or 0), float(row.minimum_stock or 0),
                float(row.reorder_level or 0), float(row.maximum_stock or 0),
                row.lead_time_days, row.tax_category, bool(row.is_taxable),
                float(row.tax_rate or 0), bool(row.is_active), row.item_type, departments.get(row.id, {}).get("department_name", "Shared"),
            ])
        last_id = rows[-1].id
    wb.save(output_path)


def _generate_general_ledger_xlsx(db, output_path: Path, params: dict | None) -> None:
    data = params or {}

    def _date(value):
        if not value:
            return None
        try:
            return datetime.fromisoformat(str(value)).date()
        except (TypeError, ValueError):
            return None

    def _int(value):
        if value in (None, ""):
            return None
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    def _string(value):
        value = str(value or "")
        return value.strip() or None

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("General Ledger")
    ws.append([
        "Date",
        "Journal No",
        "Account",
        "Description",
        "Document Number",
        "Location",
        "Department",
        "Supplier",
        "Business Unit",
        "Debit",
        "Credit",
        "Running Balance",
    ])

    q = build_ledger_query(
        db,
        account_id=_int(data.get("account_id")),
        account_ids=_string(data.get("account_ids")),
        account_from_id=_int(data.get("account_from_id")),
        account_to_id=_int(data.get("account_to_id")),
        search_type=_string(data.get("search_type")) or "ALL",
        date_from=_date(data.get("date_from")),
        date_to=_date(data.get("date_to")),
        reference_type=_string(data.get("reference_type")),
        supplier_id=_int(data.get("supplier_id")),
        department_id=_int(data.get("department_id")),
        location_id=_int(data.get("location_id")),
        business_unit_id=_int(data.get("business_unit_id")),
        search=_string(data.get("search")),
    )
    for line, balance in ledger_rows_with_balance(q).yield_per(1000):
        ws.append([
            line.journal_entry.posting_date,
            line.journal_entry.journal_number,
            line.account.account_name,
            line.description,
            line.journal_entry.document_number,
            line.location.name if line.location else None,
            line.department.name if line.department else None,
            line.supplier.name if line.supplier else None,
            line.business_unit.name if line.business_unit else None,
            float(line.debit or 0),
            float(line.credit or 0),
            float(balance or 0),
        ])
    total_debit, total_credit = ledger_totals(q)
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "TOTAL", float(total_debit), float(total_credit), ""])
    wb.save(output_path)


def run_job(job_id: str) -> None:
    db = SessionLocal()
    output_path: Path | None = None
    try:
        job = db.get(ExportJob, job_id)
        if not job or job.status not in {"PENDING", "PROCESSING"}:
            return

        job.status = "PROCESSING"
        job.started_at = datetime.now(timezone.utc)
        job.error_message = None
        db.commit()

        if job.export_type == "products":
            filename = f"MMH-Master-Products-{datetime.now(timezone.utc):%Y%m%d-%H%M%S}.xlsx"
        elif job.export_type == "general_ledger":
            filename = f"MMH-General-Ledger-{datetime.now(timezone.utc):%Y%m%d-%H%M%S}.xlsx"
        else:
            raise ValueError(f"Unsupported export type: {job.export_type}")

        tmp_dir = Path(settings.EXPORT_TMP_DIR)
        tmp_dir.mkdir(parents=True, exist_ok=True)
        output_path = tmp_dir / f"{job.id}.xlsx"

        if job.export_type == "products":
            user = db.get(User, job.user_id)
            if not user:
                raise ValueError("Export owner no longer exists")
            _generate_products_xlsx(db, output_path, job.params, user)
        elif job.export_type == "general_ledger":
            params = dict(job.params or {})
            user = db.get(User, job.user_id)
            if user:
                params["location_id"] = resolve_location_id(user, params.get("location_id"))
            _generate_general_ledger_xlsx(db, output_path, params)

        bucket = settings.EXPORT_S3_BUCKET or settings.S3_BUCKET
        if not bucket:
            raise RuntimeError("EXPORT_S3_BUCKET or S3_BUCKET must be configured")

        key = f"{settings.EXPORT_S3_PREFIX.strip('/')}/{job.user_id}/{job.id}/{filename}"
        boto3.client("s3", region_name=settings.AWS_REGION or None).upload_file(
            str(output_path),
            bucket,
            key,
            ExtraArgs={
                "ContentType": MIME_XLSX,
                "ContentDisposition": f'attachment; filename="{filename}"',
            },
        )

        job.file_key = key
        job.filename = filename
        job.status = "READY"
        job.completed_at = datetime.now(timezone.utc)
        db.commit()
    except Exception as exc:
        db.rollback()
        job = db.get(ExportJob, job_id)
        if job:
            job.status = "FAILED"
            job.error_message = str(exc)[:2000]
            job.completed_at = datetime.now(timezone.utc)
            db.commit()
        raise
    finally:
        if output_path and output_path.exists():
            output_path.unlink(missing_ok=True)
        db.close()


def next_pending_job_id() -> str | None:
    """Return the oldest pending job.

    HMM currently runs one export-worker service on the EC2 host, so one consumer
    is sufficient and keeps the queue semantics intentionally simple.
    """
    with SessionLocal() as db:
        job = (
            db.query(ExportJob)
            .filter(ExportJob.status == "PENDING")
            .order_by(ExportJob.created_at.asc(), ExportJob.id.asc())
            .first()
        )
        return job.id if job else None


def poll(interval: float) -> None:
    print(f"HMM export worker polling every {interval:g}s", flush=True)
    while True:
        job_id = next_pending_job_id()
        if not job_id:
            time.sleep(interval)
            continue
        try:
            print(f"Processing export job {job_id}", flush=True)
            run_job(job_id)
            print(f"Export job {job_id} ready", flush=True)
        except Exception as exc:
            print(f"Export job {job_id} failed: {exc}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="HMM Connect XLS export worker")
    parser.add_argument("--job-id")
    parser.add_argument("--poll", action="store_true")
    parser.add_argument("--interval", type=float, default=2.0)
    args = parser.parse_args()

    if args.poll:
        poll(max(0.5, args.interval))
        return
    if not args.job_id:
        parser.error("either --job-id or --poll is required")
    run_job(args.job_id)


if __name__ == "__main__":
    main()
