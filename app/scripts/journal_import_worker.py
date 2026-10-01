from __future__ import annotations

import argparse
import csv
import os
import re
import tempfile
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation

import boto3
from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.api.v1.journals import post_locked_journal
from app.core.config import settings
from app.core.database import SessionLocal
from app.models.business_unit import BusinessUnit
from app.models.chart_of_account import ChartOfAccount
from app.models.department import Department
from app.models.import_batch import ImportBatch, ImportRow
from app.models.journal_entry import JournalEntry
from app.models.journal_line import JournalLine
from app.models.location import Location
from app.models.supplier import Supplier
from app.models.user import User
from app.services.activity_service import ActivityService
from app.services.document_number import next_document_number
from app.services.location_access import enforce_location_access


REQUIRED_HEADERS = {"posting_date", "account_code", "debit", "credit"}
MAX_ROWS = 5000
MAX_DOCUMENTS = 500
ZERO = Decimal("0")

HEADER_ALIASES = {
    "group": "journal_group",
    "journal_id": "journal_group",
    "journal_no": "journal_group",
    "document_no": "document_number",
    "document": "document_number",
    "date": "posting_date",
    "journal_date": "posting_date",
    "reference": "reference_type",
    "ref_type": "reference_type",
    "ref_id": "reference_id",
    "business_unit_name": "business_unit",
    "bu_name": "business_unit",
    "department_name": "department",
    "location_name": "location",
    "account": "account_code",
    "account_number": "account_code",
    "description_line": "line_description",
    "supplier_name": "supplier",
    "external_ref": "external_reference",
}

INHERITED_FIELDS = (
    "posting_date",
    "description",
    "reference_type",
    "reference_id",
    "business_unit",
    "business_unit_code",
    "department",
    "department_code",
    "location",
    "location_code",
    "external_reference",
)


def text(value) -> str:
    return str(value).strip() if value is not None else ""


def json_safe(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def header(value) -> str:
    value = text(value).lower().replace("-", "_").replace(" ", "_")
    if value.endswith("_(optional)"):
        value = value[:-11]
    return HEADER_ALIASES.get(value, value)


def money(value, field: str) -> Decimal:
    if value is None or text(value) == "":
        return ZERO
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))

    raw = re.sub(r"(?i)rp", "", text(value)).replace(" ", "")
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif raw.count(".") > 1:
        raw = raw.replace(".", "")
    elif raw.count(",") > 1:
        raw = raw.replace(",", "")
    elif "," in raw:
        left, right = raw.split(",", 1)
        raw = f"{left}.{right}" if len(right) in {1, 2} else left + right
    elif "." in raw:
        left, right = raw.split(".", 1)
        if len(right) == 3 and left not in {"0", "-0"}:
            raw = left + right

    try:
        return Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"{field}: invalid amount {value!r}") from exc


def parse_date(value, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    if not raw:
        raise ValueError(f"{field}: required on the first line of each journal group")
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{field}: invalid date {value!r}")


def parse_int(value, field: str, default: int = 0) -> int:
    raw = text(value)
    if not raw:
        return default
    try:
        return int(raw)
    except ValueError as exc:
        raise ValueError(f"{field}: invalid integer {value!r}") from exc


def read_rows(path: str, original_name: str) -> tuple[list[tuple[int, str, dict]], set[str]]:
    suffix = os.path.splitext(original_name.lower())[1]
    result: list[tuple[int, str, dict]] = []
    all_headers: set[str] = set()

    if suffix == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            raw_headers = next(reader, None)
            if not raw_headers:
                raise ValueError("The import file is empty")
            headers = [header(value) for value in raw_headers]
            all_headers.update(key for key in headers if key)
            for row_number, values in enumerate(reader, 2):
                if not any(text(value) for value in values):
                    continue
                result.append((row_number, "CSV", {
                    key: json_safe(value) for key, value in zip(headers, values) if key
                }))
    elif suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible" or sheet.title.lower().startswith(("instruction", "master ")):
                    continue
                values = sheet.iter_rows(values_only=True)
                raw_headers = next(values, None)
                if not raw_headers:
                    continue
                headers = [header(value) for value in raw_headers]
                all_headers.update(key for key in headers if key)
                for row_number, cells in enumerate(values, 2):
                    if not any(value is not None and text(value) for value in cells):
                        continue
                    result.append((row_number, sheet.title, {
                        key: json_safe(value) for key, value in zip(headers, cells) if key
                    }))
        finally:
            workbook.close()
    else:
        raise ValueError("Only XLSX and CSV files are supported")

    if not result:
        raise ValueError("The import file contains no data rows")
    if len(result) > MAX_ROWS:
        raise ValueError(f"Import exceeds the {MAX_ROWS}-row limit")
    return result, all_headers


def normalize(raw: dict) -> dict:
    return {
        "journal_group": text(raw.get("journal_group")),
        "document_number": text(raw.get("document_number")),
        "posting_date": json_safe(raw.get("posting_date")),
        "description": text(raw.get("description")),
        "reference_type": text(raw.get("reference_type")).upper() or "MANUAL",
        "reference_id": text(raw.get("reference_id")) or "0",
        "business_unit": text(raw.get("business_unit")),
        "business_unit_code": text(raw.get("business_unit_code")).upper(),
        "department": text(raw.get("department")),
        "department_code": text(raw.get("department_code")).upper(),
        "location": text(raw.get("location")),
        "location_code": text(raw.get("location_code")).upper(),
        "account_code": text(raw.get("account_code")).upper(),
        "line_description": text(raw.get("line_description")),
        "debit": json_safe(raw.get("debit")) if raw.get("debit") is not None else "0",
        "credit": json_safe(raw.get("credit")) if raw.get("credit") is not None else "0",
        "supplier": text(raw.get("supplier")),
        "supplier_code": text(raw.get("supplier_code")).upper(),
        "external_reference": text(raw.get("external_reference")),
        "status": "DRAFT",
    }


def resolve_location(db: Session, data: dict) -> Location | None:
    code = text(data.get("location_code")).upper()
    name = text(data.get("location"))
    if code:
        return db.query(Location).filter(func.upper(Location.code) == code).first()
    if name:
        return db.query(Location).filter(func.lower(Location.name) == name.lower()).first()
    return None


def resolve_business_unit(db: Session, data: dict, location: Location | None) -> BusinessUnit | None:
    code = text(data.get("business_unit_code")).upper()
    name = text(data.get("business_unit"))
    if code:
        return db.query(BusinessUnit).filter(func.upper(BusinessUnit.code) == code).first()
    if name:
        return db.query(BusinessUnit).filter(func.lower(BusinessUnit.name) == name.lower()).first()
    return location.business_unit if location else None


def resolve_department(db: Session, data: dict) -> Department | None:
    code = text(data.get("department_code")).upper()
    name = text(data.get("department"))
    if code:
        return db.query(Department).filter(func.upper(Department.code) == code).first()
    if name:
        return db.query(Department).filter(func.lower(Department.name) == name.lower()).first()
    return None


def resolve_supplier(db: Session, data: dict) -> tuple[Supplier | None, str | None]:
    code = text(data.get("supplier_code")).upper()
    name = text(data.get("supplier"))
    if code:
        supplier = db.query(Supplier).filter(func.upper(Supplier.code) == code).first()
        return supplier, None if supplier else f"unknown supplier code {code}"
    if not name:
        return None, None
    suppliers = db.query(Supplier).filter(func.lower(Supplier.name) == name.lower()).all()
    if len(suppliers) > 1:
        return None, f"ambiguous supplier name {name}; use supplier_code"
    return (suppliers[0], None) if suppliers else (None, f"unknown supplier {name}")


def validate_group(db: Session, user: User, batch: ImportBatch, rows: list[ImportRow]) -> list[str]:
    errors: list[str] = []
    first = rows[0].normalized_data_json or {}
    label = text(first.get("_group_label")) or text(first.get("document_number"))

    try:
        posting_date = parse_date(first.get("posting_date"), "posting_date")
    except ValueError as exc:
        errors.append(str(exc))
        posting_date = None

    try:
        reference_id = parse_int(first.get("reference_id"), "reference_id")
        if reference_id < 0:
            errors.append("reference_id cannot be negative")
    except ValueError as exc:
        errors.append(str(exc))

    if text(first.get("reference_type")).upper() != "MANUAL":
        errors.append("Bulk journal reference_type must be MANUAL; source-controlled journals belong in their source module")

    location = resolve_location(db, first)
    if not location:
        errors.append(f"Unknown location: {text(first.get('location_code')) or text(first.get('location')) or '(blank)'}")
    else:
        if not location.is_active:
            errors.append(f"Location is inactive: {location.code}")
        try:
            enforce_location_access(user, location.id)
        except Exception:
            errors.append(f"You do not have access to location {location.code}")

    business_unit = resolve_business_unit(db, first, location)
    if not business_unit:
        errors.append("Business unit could not be determined from the location")
    else:
        if not business_unit.is_active:
            errors.append(f"Business unit is inactive: {business_unit.code}")
        if location and location.business_unit_id and location.business_unit_id != business_unit.id:
            errors.append(f"Business unit {business_unit.code} does not match location {location.code}")

    department = resolve_department(db, first)
    department_supplied = bool(text(first.get("department_code")) or text(first.get("department")))
    if department_supplied and not department:
        errors.append(f"Unknown department: {text(first.get('department_code')) or text(first.get('department'))}")
    elif department:
        if not department.is_active:
            errors.append(f"Department is inactive: {department.code}")
        if business_unit and getattr(department, "business_unit_id", None) not in {None, business_unit.id}:
            errors.append(f"Department {department.code} does not belong to business unit {business_unit.code}")

    supplied_document = text(first.get("_supplied_document_number"))
    if supplied_document and db.query(JournalEntry).filter(
        JournalEntry.document_number == supplied_document
    ).first():
        errors.append(f"Document number already exists: {supplied_document}")

    for row in rows[1:]:
        data = row.normalized_data_json or {}
        for field in INHERITED_FIELDS:
            if text(data.get(field)) != text(first.get(field)):
                errors.append(f"Row {row.row_number}: {field} conflicts with journal group {label}; start a new journal_group")

    debit_total = ZERO
    credit_total = ZERO
    for row in rows:
        data = row.normalized_data_json or {}
        try:
            debit = money(data.get("debit"), "debit")
            credit = money(data.get("credit"), "credit")
        except ValueError as exc:
            errors.append(f"Row {row.row_number}: {exc}")
            continue
        if debit < ZERO or credit < ZERO:
            errors.append(f"Row {row.row_number}: debit and credit cannot be negative")
        if (debit > ZERO) == (credit > ZERO):
            errors.append(f"Row {row.row_number}: enter either debit or credit, not both and not zero/zero")
        debit_total += debit
        credit_total += credit

        account_code = text(data.get("account_code")).upper()
        account = db.query(ChartOfAccount).filter(ChartOfAccount.account_number == account_code).first()
        if not account or not account.is_active or not account.is_posting:
            errors.append(f"Row {row.row_number}: invalid active posting account {account_code or '(blank)'}")

        supplier, supplier_error = resolve_supplier(db, data)
        if supplier_error:
            errors.append(f"Row {row.row_number}: {supplier_error}")
        elif supplier and not supplier.is_active:
            errors.append(f"Row {row.row_number}: supplier {supplier.code} is inactive")

    if debit_total != credit_total:
        difference = debit_total - credit_total
        side = "credit" if difference > ZERO else "debit"
        errors.append(f"Journal group {label} needs {side} {abs(difference)} to balance")

    if posting_date:
        batch.period_from = posting_date if batch.period_from is None else min(batch.period_from, posting_date)
        batch.period_to = posting_date if batch.period_to is None else max(batch.period_to, posting_date)
    return list(dict.fromkeys(errors))


def build_groups(db: Session, batch: ImportBatch, source_rows: list[tuple[int, str, dict]]) -> dict[str, list[ImportRow]]:
    groups: dict[str, list[ImportRow]] = defaultdict(list)
    current_key: str | None = None
    group_headers: dict[str, dict] = {}
    group_numbers: dict[str, int] = {}

    for row_number, sheet_name, raw in source_rows:
        normalized = normalize(raw)
        explicit_group = normalized["journal_group"]
        explicit_document = normalized["document_number"]

        if explicit_group:
            current_key = f"GROUP:{explicit_group}"
        elif explicit_document:
            current_key = f"DOCUMENT:{explicit_document}"
        elif current_key is None:
            current_key = "GROUP:1"

        if current_key not in group_numbers:
            group_numbers[current_key] = len(group_numbers) + 1
        ordinal = group_numbers[current_key]
        group_label = explicit_group or group_headers.get(current_key, {}).get("_group_label") or str(ordinal)

        if current_key not in group_headers:
            group_headers[current_key] = dict(normalized)
            group_headers[current_key]["_group_label"] = group_label
            group_headers[current_key]["_supplied_document_number"] = explicit_document
        else:
            first = group_headers[current_key]
            for field in INHERITED_FIELDS:
                if not text(normalized.get(field)):
                    normalized[field] = first.get(field)

        normalized["_group_key"] = current_key
        normalized["_group_label"] = group_headers[current_key]["_group_label"]
        normalized["_supplied_document_number"] = group_headers[current_key]["_supplied_document_number"]
        tracking_number = explicit_document or group_headers[current_key].get("_supplied_document_number")
        if not tracking_number:
            tracking_number = f"{batch.batch_number}-G{ordinal:04d}"
        normalized["document_number"] = tracking_number

        row = ImportRow(
            import_batch_id=batch.id,
            row_number=row_number,
            sheet_name=sheet_name,
            document_type="JOURNAL_ENTRY",
            raw_data_json=raw,
            normalized_data_json=normalized,
            status="PENDING",
            document_number=tracking_number,
        )
        db.add(row)
        groups[current_key].append(row)
    return groups


def validate_job(db: Session, batch: ImportBatch, user: User) -> None:
    bucket = settings.EXPORT_S3_BUCKET or settings.S3_BUCKET
    if not bucket or not batch.stored_file_path:
        raise RuntimeError("Journal import storage is not configured")

    suffix = os.path.splitext(batch.original_file_name or "")[1].lower() or ".xlsx"
    fd, temp_path = tempfile.mkstemp(prefix=f"hmm-journal-import-{batch.id}-", suffix=suffix, dir="/tmp")
    os.close(fd)
    try:
        boto3.client("s3", region_name=settings.AWS_REGION or None).download_file(
            bucket, batch.stored_file_path, temp_path
        )
        source_rows, headers = read_rows(temp_path, batch.original_file_name or "")
    finally:
        try:
            os.unlink(temp_path)
        except FileNotFoundError:
            pass

    missing = REQUIRED_HEADERS - headers
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(sorted(missing)))

    batch.rows.clear()
    batch.period_from = None
    batch.period_to = None
    db.flush()
    groups = build_groups(db, batch, source_rows)
    db.flush()
    if len(groups) > MAX_DOCUMENTS:
        raise ValueError(f"Import exceeds the {MAX_DOCUMENTS}-journal limit")

    valid_documents = 0
    error_documents = 0
    for rows in groups.values():
        errors = validate_group(db, user, batch, rows)
        status = "ERROR" if errors else "VALID"
        if errors:
            error_documents += 1
        else:
            valid_documents += 1
        for row in rows:
            row.status = status
            row.error_messages_json = errors
            row.warning_messages_json = []

    batch.total_rows = len(source_rows)
    batch.total_documents = len(groups)
    batch.valid_rows = sum(row.status == "VALID" for rows in groups.values() for row in rows)
    batch.error_rows = sum(row.status == "ERROR" for rows in groups.values() for row in rows)
    batch.warning_rows = 0
    batch.success_count = valid_documents
    batch.failed_count = error_documents
    batch.status = "VALIDATED"
    batch.validated_at = datetime.now(timezone.utc)
    batch.worker_completed_at = datetime.now(timezone.utc)
    batch.worker_error_message = None
    db.commit()


def import_valid_groups(db: Session, batch: ImportBatch, user: User, post_valid: bool) -> None:
    rows = db.query(ImportRow).filter(ImportRow.import_batch_id == batch.id).order_by(ImportRow.id).all()
    groups: dict[str, list[int]] = defaultdict(list)
    invalid_documents = set()
    for row in rows:
        data = row.normalized_data_json or {}
        key = text(data.get("_group_key")) or row.document_number or str(row.id)
        if row.status == "VALID":
            groups[key].append(row.id)
        elif row.status in {"ERROR", "FAILED"}:
            invalid_documents.add(key)

    imported_documents = 0
    imported_rows = 0
    runtime_failed = 0
    for row_ids in groups.values():
        try:
            group_rows = db.query(ImportRow).filter(ImportRow.id.in_(row_ids)).order_by(ImportRow.id).all()
            first = group_rows[0].normalized_data_json or {}
            location = resolve_location(db, first)
            business_unit = resolve_business_unit(db, first, location)
            department = resolve_department(db, first)
            if not location or not business_unit:
                raise ValueError("Location or business unit no longer exists")

            supplied_document = text(first.get("_supplied_document_number"))
            journal = JournalEntry(
                journal_number=next_document_number(db, "MAN", digits=6),
                posting_date=parse_date(first.get("posting_date"), "posting_date"),
                status="DRAFT",
                document_type="BULK_JOURNAL_IMPORT",
                document_number=supplied_document or first.get("document_number"),
                reference_type="MANUAL",
                reference_id=0,
                description=text(first.get("description")) or None,
                created_by_id=user.id,
                source_type="JOURNAL_IMPORT",
                external_reference=(text(first.get("external_reference")) or f"JOURNAL_IMPORT:{batch.id}:{first.get('_group_label')}")[:100],
                import_batch_id=batch.id,
                imported_by_id=user.id,
                imported_at=datetime.now(timezone.utc),
                business_unit_id=business_unit.id,
                department_id=department.id if department else None,
                location_id=location.id,
            )
            for line_number, row in enumerate(group_rows, 1):
                data = row.normalized_data_json or {}
                account = db.query(ChartOfAccount).filter(
                    ChartOfAccount.account_number == text(data.get("account_code")).upper()
                ).one()
                supplier, supplier_error = resolve_supplier(db, data)
                if supplier_error:
                    raise ValueError(supplier_error)
                journal.lines.append(JournalLine(
                    line_number=line_number,
                    account_id=account.id,
                    debit=money(data.get("debit"), "debit"),
                    credit=money(data.get("credit"), "credit"),
                    description=text(data.get("line_description")) or None,
                    business_unit_id=business_unit.id,
                    department_id=department.id if department else None,
                    location_id=location.id,
                    supplier_id=supplier.id if supplier else None,
                ))

            db.add(journal)
            db.flush()
            ActivityService.log(
                db=db, user=user, action="JOURNAL_IMPORT_CREATED",
                entity_type="JOURNAL_ENTRY", entity_id=journal.id,
                summary=f"Imported {journal.journal_number} from {batch.batch_number}",
                details={"batch_number": batch.batch_number, "journal_group": first.get("_group_label")},
                import_batch_id=batch.id,
            )
            if post_valid:
                post_locked_journal(db, journal=journal, user=user)
            imported_at = datetime.now(timezone.utc)
            for row in group_rows:
                row.status = "IMPORTED"
                row.entity_type = "JOURNAL_ENTRY"
                row.entity_id = journal.id
                row.imported_at = imported_at
            db.commit()
            imported_documents += 1
            imported_rows += len(group_rows)
        except Exception as exc:
            db.rollback()
            failed_rows = db.query(ImportRow).filter(ImportRow.id.in_(row_ids)).all()
            message = str(getattr(exc, "detail", exc))[:1000]
            for row in failed_rows:
                row.status = "FAILED"
                row.error_messages_json = list(dict.fromkeys([*(row.error_messages_json or []), message]))
            db.commit()
            runtime_failed += 1

    batch = db.get(ImportBatch, batch.id)
    batch.imported_rows = imported_rows
    batch.failed_rows = db.query(ImportRow).filter(
        ImportRow.import_batch_id == batch.id,
        ImportRow.status.in_({"ERROR", "FAILED"}),
    ).count()
    batch.success_count = imported_documents
    batch.failed_count = len(invalid_documents) + runtime_failed
    batch.status = "IMPORTED_WITH_ERRORS" if batch.failed_count else "IMPORTED"
    batch.imported_at = datetime.now(timezone.utc)
    batch.worker_completed_at = datetime.now(timezone.utc)
    batch.worker_error_message = None
    ActivityService.log(
        db=db, user=user, action="JOURNAL_IMPORT_COMPLETED",
        entity_type="IMPORT_BATCH", entity_id=batch.id,
        summary=f"Imported {imported_documents} journals from {batch.batch_number}",
        details={"documents": imported_documents, "rows": imported_rows, "failed_documents": batch.failed_count},
        import_batch_id=batch.id,
    )
    db.commit()


def run_job(job_id: int, mode: str = "validate", post_valid: bool = False) -> None:
    db = SessionLocal()
    try:
        batch = db.get(ImportBatch, job_id)
        if not batch or batch.source != "JOURNAL_BULK_IMPORT":
            return
        user = db.get(User, batch.created_by_id)
        if not user:
            raise RuntimeError("Import creator no longer exists")

        expected = {"validate": {"QUEUED", "PROCESSING"}, "import": {"IMPORT_QUEUED", "IMPORTING"}}
        if batch.status not in expected[mode]:
            return
        batch.status = "PROCESSING" if mode == "validate" else "IMPORTING"
        batch.worker_started_at = datetime.now(timezone.utc)
        batch.worker_error_message = None
        db.commit()

        if mode == "validate":
            validate_job(db, batch, user)
        else:
            import_valid_groups(db, batch, user, post_valid)
    except Exception as exc:
        db.rollback()
        failed_batch = db.get(ImportBatch, job_id)
        if failed_batch:
            failed_batch.status = "FAILED"
            failed_batch.worker_error_message = str(getattr(exc, "detail", exc))[:4000]
            failed_batch.worker_completed_at = datetime.now(timezone.utc)
            db.commit()
        raise
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="HMM Connect journal import worker")
    parser.add_argument("--job-id", type=int, required=True)
    parser.add_argument("--mode", choices=("validate", "import"), default="validate")
    parser.add_argument("--post-valid", action="store_true")
    args = parser.parse_args()
    run_job(args.job_id, mode=args.mode, post_valid=args.post_valid)


if __name__ == "__main__":
    main()
