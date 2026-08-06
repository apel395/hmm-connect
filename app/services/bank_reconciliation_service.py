import csv, io, json, re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload
from app.models.bank_reconciliation import BankReconciliation, BankStatementLine, BankReconciliationMatch
from app.models.chart_of_account import ChartOfAccount
from app.models.journal_entry import JournalEntry
from app.models.journal_line import JournalLine
from app.services.document_number import next_document_number

ZERO=Decimal("0.00")
def dec(value):
    if value in (None,""):return ZERO
    text=str(value).strip().replace("Rp","").replace(" ","")
    if "," in text and "." in text:text=text.replace(".","").replace(",",".")
    elif "," in text:text=text.replace(",",".")
    try:return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:raise ValueError(f"Invalid amount: {value}")
def normal(value):return re.sub(r"[^A-Z0-9]","",str(value or "").upper())

class BankReconciliationService:
    def __init__(self,db:Session):self.db=db
    def get(self,id,lock=False):
        q=self.db.query(BankReconciliation).filter(BankReconciliation.id==id)
        if not lock:q=q.options(joinedload(BankReconciliation.account))
        row=(q.with_for_update(of=BankReconciliation) if lock else q).first()
        if not row:raise ValueError("Reconciliation not found")
        return row
    def validate_account(self,id):
        a=self.db.get(ChartOfAccount,id)
        if not a or not a.is_active or not a.is_posting or a.account_type!="ASSET" or a.account_subtype!="BANK":raise ValueError("An active posting BANK account is required")
        return a
    def system_rows(self,r):
        matched=self.db.query(BankReconciliationMatch.journal_line_id,func.coalesce(func.sum(BankReconciliationMatch.matched_amount),0).label("matched")).join(BankReconciliation).filter(BankReconciliation.status!="VOID").group_by(BankReconciliationMatch.journal_line_id).subquery()
        rows=self.db.query(JournalLine,JournalEntry,func.coalesce(matched.c.matched,0)).join(JournalEntry,JournalEntry.id==JournalLine.journal_entry_id).outerjoin(matched,matched.c.journal_line_id==JournalLine.id).filter(JournalLine.account_id==r.account_id,JournalEntry.status=="POSTED",JournalEntry.posting_date>=r.statement_start_date,JournalEntry.posting_date<=r.statement_end_date).order_by(JournalEntry.posting_date,JournalLine.id).all()
        return [{"journal_line_id":l.id,"journal_id":j.id,"journal_number":j.journal_number,"posting_date":j.posting_date,"reference_type":j.reference_type,"reference_number":j.document_number or j.external_reference,"description":l.description or j.description,"counterparty":l.supplier.name if l.supplier else None,"debit":l.debit,"credit":l.credit,"amount":max(Decimal(l.debit),Decimal(l.credit)),"matched_amount":Decimal(m),"remaining_amount":max(Decimal(l.debit),Decimal(l.credit))-Decimal(m),"cleared":Decimal(m)>=max(Decimal(l.debit),Decimal(l.credit)),"outstanding_type":self.outstanding(j,l),"source_route":f"/journal-entries/{j.id}"} for l,j,m in rows]
    @staticmethod
    def outstanding(j,l):
        ref=(j.reference_type or "").upper()
        if "TRANSFER" in ref:return "TRANSFER_IN_TRANSIT"
        if Decimal(l.credit)>0:return "OUTSTANDING_PAYMENT"
        if Decimal(l.debit)>0:return "DEPOSIT_IN_TRANSIT"
        return "UNIDENTIFIED"
    def recalc(self,r):
        for line in r.statement_lines:
            if line.match_status=="IGNORED":continue
            allocated=sum((Decimal(x.matched_amount) for x in line.matches),ZERO);amount=max(Decimal(line.debit),Decimal(line.credit))
            line.match_status="UNMATCHED" if allocated==ZERO else "MATCHED" if allocated==amount else "PARTIAL"
        before=self.db.query(func.coalesce(func.sum(JournalLine.debit-JournalLine.credit),0)).join(JournalEntry).filter(JournalLine.account_id==r.account_id,JournalEntry.status=="POSTED",JournalEntry.posting_date<r.statement_start_date).scalar()
        period=self.db.query(func.coalesce(func.sum(JournalLine.debit),0),func.coalesce(func.sum(JournalLine.credit),0)).join(JournalEntry).filter(JournalLine.account_id==r.account_id,JournalEntry.status=="POSTED",JournalEntry.posting_date>=r.statement_start_date,JournalEntry.posting_date<=r.statement_end_date).one()
        r.system_opening_balance=Decimal(before);r.total_system_debit=Decimal(period[0]);r.total_system_credit=Decimal(period[1]);r.system_closing_balance=r.system_opening_balance+r.total_system_debit-r.total_system_credit
        r.total_statement_debit=sum((Decimal(x.debit) for x in r.statement_lines),ZERO);r.total_statement_credit=sum((Decimal(x.credit) for x in r.statement_lines),ZERO)
        r.matched_amount=sum((Decimal(x.matched_amount) for x in r.matches),ZERO)
        r.unmatched_statement_amount=sum((max(Decimal(x.debit),Decimal(x.credit))-sum((Decimal(m.matched_amount) for m in x.matches),ZERO) for x in r.statement_lines if x.match_status!="IGNORED"),ZERO)
        r.unmatched_system_amount=sum((max(Decimal(x["remaining_amount"]),ZERO) for x in self.system_rows(r)),ZERO)
        adjustment_types={"BANK_FEE","BANK_INTEREST","TAX_WITHHOLDING","OTHER_BANK_ADJUSTMENT"}
        r.adjustment_amount=sum((Decimal(x.matched_amount) for x in r.matches if x.journal_line.journal_entry.reference_type in adjustment_types),ZERO)
        r.difference=(Decimal(r.statement_closing_balance)-Decimal(r.system_closing_balance)).quantize(Decimal("0.01"));self.db.flush();return r
    def serialize(self,r,detail=False):
        self.recalc(r)
        data={"id":r.id,"reconciliation_number":r.reconciliation_number,"account_id":r.account_id,"account_code":r.account.account_number,"account_name":r.account.account_name,"statement_start_date":r.statement_start_date,"statement_end_date":r.statement_end_date,"statement_opening_balance":r.statement_opening_balance,"statement_closing_balance":r.statement_closing_balance,"system_opening_balance":r.system_opening_balance,"system_closing_balance":r.system_closing_balance,"total_statement_debit":r.total_statement_debit,"total_statement_credit":r.total_statement_credit,"total_system_debit":r.total_system_debit,"total_system_credit":r.total_system_credit,"matched_amount":r.matched_amount,"unmatched_statement_amount":r.unmatched_statement_amount,"unmatched_system_amount":r.unmatched_system_amount,"adjustment_amount":r.adjustment_amount,"difference":r.difference,"matched_count":sum(x.match_status in ("MATCHED","ADJUSTMENT_CREATED") for x in r.statement_lines),"unmatched_count":sum(x.match_status=="UNMATCHED" for x in r.statement_lines),"status":r.status,"notes":r.notes,"source_filename":r.source_filename,"created_by":r.created_by.full_name if r.created_by else None,"completed_by":r.completed_by.full_name if r.completed_by else None,"completed_at":r.completed_at,"allowed_actions":{"update":r.status in ("DRAFT","IN_PROGRESS"),"match":r.status in ("DRAFT","IN_PROGRESS"),"complete":r.status in ("DRAFT","IN_PROGRESS") and r.difference==ZERO,"void":r.status in ("DRAFT","IN_PROGRESS"),"print":True}}
        return data
    def create(self,data,user):
        self.validate_account(data.account_id)
        if data.statement_start_date>data.statement_end_date:raise ValueError("Statement period is invalid")
        r=BankReconciliation(reconciliation_number=next_document_number(self.db,"BR"),created_by_id=user.id,status="DRAFT",**data.model_dump());self.db.add(r);self.db.flush();self.recalc(r);return r
    def parse(self,content,filename,mapping):
        suffix=filename.lower().rsplit(".",1)[-1]
        if suffix=="csv":records=list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        elif suffix=="xlsx":
            ws=load_workbook(io.BytesIO(content),read_only=True,data_only=True).active;rows=list(ws.iter_rows(values_only=True));headers=[str(x or "").strip() for x in rows[0]];records=[dict(zip(headers,row)) for row in rows[1:]]
        else:raise ValueError("Only CSV and XLSX files are supported")
        def value(row,key):return row.get(mapping.get(key,key))
        parsed=[];errors=[]
        for number,row in enumerate(records,2):
            try:
                raw_date=value(row,"transaction_date");d=raw_date.date() if isinstance(raw_date,datetime) else datetime.fromisoformat(str(raw_date)).date()
                debit,credit=dec(value(row,"debit")),dec(value(row,"credit"));amount=dec(value(row,"amount"))
                if not debit and not credit and amount: credit=amount if amount>0 else ZERO;debit=abs(amount) if amount<0 else ZERO
                if (debit>0)==(credit>0):raise ValueError("Provide exactly one debit or credit")
                parsed.append({"transaction_date":d,"value_date":None,"description":value(row,"description"),"reference_number":value(row,"reference_number"),"debit":debit,"credit":credit,"balance":dec(value(row,"balance")) if value(row,"balance") not in (None,"") else None,"external_transaction_id":value(row,"external_transaction_id"),"counterparty_name":value(row,"counterparty_name"),"raw_data_json":{str(k):str(v)[:500] for k,v in row.items()}})
            except Exception as e:errors.append({"row":number,"error":str(e)})
        return parsed,errors,len(records)
    def add_match(self,r,alloc,user):
        if r.status not in ("DRAFT","IN_PROGRESS"):raise ValueError("Completed reconciliation is locked")
        s=self.db.get(BankStatementLine,alloc.statement_line_id);j=self.db.get(JournalLine,alloc.journal_line_id)
        if not s or s.reconciliation_id!=r.id or not j or j.account_id!=r.account_id or j.journal_entry.status!="POSTED":raise ValueError("Invalid match selection")
        s_remaining=max(Decimal(s.debit),Decimal(s.credit))-sum((Decimal(x.matched_amount) for x in s.matches),ZERO)
        used=Decimal(self.db.query(func.coalesce(func.sum(BankReconciliationMatch.matched_amount),0)).join(BankReconciliation).filter(BankReconciliationMatch.journal_line_id==j.id,BankReconciliation.status!="VOID").scalar())
        j_remaining=max(Decimal(j.debit),Decimal(j.credit))-used
        same=(Decimal(s.debit)>0 and Decimal(j.credit)>0) or (Decimal(s.credit)>0 and Decimal(j.debit)>0)
        if not same or alloc.matched_amount>s_remaining or alloc.matched_amount>j_remaining:raise ValueError("Match exceeds remaining amount or direction differs")
        m=BankReconciliationMatch(reconciliation_id=r.id,matched_by_id=user.id,**alloc.model_dump());self.db.add(m);self.db.flush();total=sum((Decimal(x.matched_amount) for x in s.matches),ZERO);s.match_status="MATCHED" if total==max(Decimal(s.debit),Decimal(s.credit)) else "PARTIAL";s.matched_journal_line_id=j.id if len(s.matches)==1 else None;s.matched_at=datetime.now(timezone.utc);r.status="IN_PROGRESS";return m
    def auto_match(self,r,user,tolerance=3,confirm_suggested=False):
        created=[];suggestions=[];system=[x for x in self.system_rows(r) if x["remaining_amount"]>0]
        for s in [x for x in r.statement_lines if x.match_status=="UNMATCHED"]:
            amount=max(Decimal(s.debit),Decimal(s.credit));candidates=[x for x in system if x["remaining_amount"]==amount and ((s.debit>0 and x["credit"]>0) or (s.credit>0 and x["debit"]>0)) and abs((s.transaction_date-x["posting_date"]).days)<=tolerance]
            exact=[x for x in candidates if normal(s.external_transaction_id or s.reference_number) and normal(s.external_transaction_id or s.reference_number) in (normal(x["reference_number"]),normal(x["journal_number"]))]
            pick=(exact or candidates)
            if len(pick)==1:
                kind="AUTO_EXACT" if exact else "AUTO_SUGGESTED";score=Decimal("100") if exact else Decimal("80")
                if kind=="AUTO_EXACT" or confirm_suggested:
                    from app.schemas.bank_reconciliation import MatchAllocation
                    created.append(self.add_match(r,MatchAllocation(statement_line_id=s.id,journal_line_id=pick[0]["journal_line_id"],matched_amount=amount,match_type=kind,confidence_score=score),user).id);pick[0]["remaining_amount"]=ZERO
                else:suggestions.append({"statement_line_id":s.id,"journal_line_id":pick[0]["journal_line_id"],"matched_amount":amount,"confidence_score":score})
        return created,suggestions
