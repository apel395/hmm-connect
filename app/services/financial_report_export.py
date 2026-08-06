import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side


def flatten_sections(report):
    rows=[]
    for section in report["sections"]:
        rows.append((section["label"], section["total"], section.get("comparative_total"), section.get("variance_amount"), section.get("variance_percentage"), "section"))
        for group in section["groups"]:
            rows.append((f"  {group['label']}", group["total"], group.get("comparative_total"), group.get("variance_amount"), group.get("variance_percentage"), "group"))
            for account in group["accounts"]:
                code=account.get("account_code") or ""
                label=f"    {code}{' — ' if code else ''}{account['account_name']}"
                rows.append((label,account["current_amount"],account.get("comparative_amount"),account.get("variance_amount"),account.get("variance_percentage"),"account"))
    return rows


def export_statement(report,report_name,period,filters,format,filename,summary_rows):
    headers=["Account / Section","Current Amount","Comparative Amount","Variance","Variance %"]
    rows=flatten_sections(report)
    if format=="csv":
        output=StringIO();writer=csv.writer(output);writer.writerow(headers)
        for label,current,comparative,variance,percentage,_ in rows:writer.writerow([label,current,comparative,variance,percentage])
        for label,value in summary_rows:writer.writerow([label,value,"","",""])
        return StreamingResponse(iter([output.getvalue()]),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":f'attachment; filename="{filename}.csv"'})
    workbook=Workbook();sheet=workbook.active;sheet.title=report_name[:31]
    sheet.append(["HMM-Connect"]);sheet.append([report_name]);sheet.append([period]);sheet.append([f"Generated: {datetime.now(timezone.utc).isoformat()}"]);sheet.append([f"Filters: {filters}"]);sheet.append([]);sheet.append(headers);header_row=7
    thin=Side(style="thin",color="D9E0E7");border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for cell in sheet[header_row]:cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="F1F5F9");cell.border=border
    for label,current,comparative,variance,percentage,kind in rows:
        sheet.append([label,current,comparative,variance,percentage]);row=sheet.max_row
        if kind in {"section","group"}:
            for cell in sheet[row]:cell.font=Font(bold=True)
    for label,value in summary_rows:
        sheet.append([label,value]);
        for cell in sheet[sheet.max_row]:cell.font=Font(bold=True)
    for row in sheet.iter_rows(min_row=header_row,max_row=sheet.max_row):
        for cell in row:cell.border=border
    for row in sheet.iter_rows(min_row=header_row+1,max_row=sheet.max_row,min_col=2,max_col=4):
        for cell in row:cell.number_format='#,##0.00'
    for cell in sheet["E"][header_row:]:cell.number_format='0.00'
    sheet.freeze_panes="A8";sheet.column_dimensions["A"].width=46
    for column in ("B","C","D","E"):sheet.column_dimensions[column].width=20
    output=BytesIO();workbook.save(output);output.seek(0)
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{filename}.xlsx"'})
