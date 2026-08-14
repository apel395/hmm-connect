from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.chart_of_account import ChartOfAccount
from app.models.department import Department
from app.models.goods_receipt import GoodsReceipt
from app.models.import_batch import ImportBatch, ImportRow
from app.models.journal_entry import JournalEntry
from app.models.journal_line import JournalLine
from app.models.location import Location
from app.models.payment import Payment
from app.models.product import Product
from app.models.purchase_order import PurchaseOrder
from app.models.purchase_order_item import PurchaseOrderItem
from app.models.purchase_requisition import PurchaseRequisition
from app.models.purchase_requisition_item import PurchaseRequisitionItem
from app.models.supplier import Supplier
from app.models.supplier_invoice import SupplierInvoice
from app.models.user import User
from app.models.warehouse import Warehouse
from app.schemas.goods_receipt import GoodsReceiptCreate, GoodsReceiptItemCreate
from app.schemas.payment import PaymentAllocationInput, PaymentCreate
from app.schemas.supplier_invoice import SupplierInvoiceInput, SupplierInvoiceLineInput
from app.services.activity_service import ActivityService
from app.services.goods_receipt_service import GoodsReceiptService
from app.services.supplier_invoice_service import SupplierInvoiceService

DOC_SLUGS = {
    "purchase-requisition": "PURCHASE_REQUISITION", "purchase-order": "PURCHASE_ORDER",
    "goods-receipt": "GOODS_RECEIPT", "supplier-invoice": "SUPPLIER_INVOICE",
    "payment": "PAYMENT", "journal-entry": "JOURNAL_ENTRY",
}
KEYS = {
    "PURCHASE_REQUISITION": "document_number", "PURCHASE_ORDER": "document_number",
    "GOODS_RECEIPT": "receipt_number", "SUPPLIER_INVOICE": "internal_invoice_number",
    "PAYMENT": "payment_number", "JOURNAL_ENTRY": "journal_number",
}
SPECS = {
    "PURCHASE_REQUISITION": ["document_number","document_date","required_date","department_code","location_code","requested_by_username","status","item_code","item_description","quantity","uom","estimated_unit_cost","notes","external_reference"],
    "PURCHASE_ORDER": ["document_number","creation_type","source_pr_number","supplier_code","order_date","estimated_delivery_date","department_code","location_code","currency","status","item_code","item_description","quantity","uom","unit_cost","discount_amount","tax_amount","notes","external_reference"],
    "GOODS_RECEIPT": ["receipt_number","receipt_date","po_number","destination_type","warehouse_code","delivery_note_number","status","item_code","item_description","received_quantity","notes","external_reference"],
    "SUPPLIER_INVOICE": ["internal_invoice_number","supplier_invoice_number","supplier_code","invoice_date","received_date","due_date","po_number","receipt_number","department_code","location_code","currency","status","item_code","description","quantity","uom","unit_price","account_code","discount_amount","tax_amount","withholding_tax_amount","notes","external_reference"],
    "PAYMENT": ["payment_number","payment_date","supplier_code","payment_method","payment_account_code","currency","total_amount","reference_number","cheque_number","status","supplier_invoice_number","allocated_amount","notes","external_reference"],
    "JOURNAL_ENTRY": ["journal_number","journal_date","description","reference_type","reference_number","status","account_code","line_description","debit","credit","department_code","location_code","supplier_code","notes"],
}
REQUIRED = {
    "PURCHASE_REQUISITION": {"document_number","document_date","department_code","location_code","quantity"},
    "PURCHASE_ORDER": {"document_number","creation_type","supplier_code","order_date","quantity","unit_cost"},
    "GOODS_RECEIPT": {"receipt_number","receipt_date","po_number","destination_type","received_quantity"},
    "SUPPLIER_INVOICE": {"internal_invoice_number","supplier_invoice_number","supplier_code","invoice_date","description","quantity","unit_price"},
    "PAYMENT": {"payment_number","payment_date","supplier_code","payment_method","payment_account_code","currency","total_amount","supplier_invoice_number","allocated_amount"},
    "JOURNAL_ENTRY": {"journal_number","journal_date","reference_type","account_code","debit","credit"},
}
EXAMPLES = {
    "PURCHASE_REQUISITION": ["PR-HIST-0001","2026-01-10","2026-01-20","DEV-PUR","DEV-PKU","requester","APPROVED","DEV-ITEM-001","Cleaning Chemical",10,"LTR",150000,"Legacy paper PR","BOX-001"],
    "PURCHASE_ORDER": ["PO-HIST-0001","DIRECT","","DEV-SUP-001","2026-01-12","2026-01-20","DEV-PUR","DEV-PKU","IDR","APPROVED","DEV-ITEM-001","Cleaning Chemical",10,"LTR",150000,0,0,"Legacy direct PO","BOX-001"],
    "GOODS_RECEIPT": ["GR-HIST-0001","2026-01-20","PO-HIST-0001","WAREHOUSE","DEV-WH-01","DN-001","POSTED","DEV-ITEM-001","Cleaning Chemical",10,"Legacy receipt","BOX-001"],
    "SUPPLIER_INVOICE": ["INV-HIST-0001","SUP-INV-001","DEV-SUP-001","2026-01-21","2026-01-21","2026-02-20","PO-HIST-0001","GR-HIST-0001","DEV-FIN","DEV-PKU","IDR","POSTED","DEV-ITEM-001","Cleaning Chemical",1,"EA",1500000,"120100",0,0,0,"Legacy invoice","BOX-001"],
    "PAYMENT": ["PAY-HIST-0001","2026-02-15","DEV-SUP-001","BANK_TRANSFER","110100","IDR",1500000,"BANK-REF-001","","POSTED","SUP-INV-001",1500000,"Legacy payment","BOX-001"],
    "JOURNAL_ENTRY": ["JV-HIST-0001","2026-01-31","Legacy adjustment","HISTORICAL_IMPORT","BOX-001","POSTED","510200","Office supplies",1000000,0,"DEV-FIN","DEV-PKU","DEV-SUP-001","Balanced counterpart required"],
}
STATUSES = {
    "PURCHASE_REQUISITION": {"DRAFT","APPROVED","REJECTED","CANCELLED"},
    "PURCHASE_ORDER": {"DRAFT","APPROVED","PARTIALLY_RECEIVED","RECEIVED","CANCELLED"},
    "GOODS_RECEIPT": {"DRAFT","POSTED"}, "SUPPLIER_INVOICE": {"DRAFT","VERIFIED","POSTED"},
    "PAYMENT": {"DRAFT","APPROVED","POSTED"}, "JOURNAL_ENTRY": {"DRAFT","POSTED"},
}


def template_bytes(doc_type: str) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Import"
    ws.append(SPECS[doc_type]); ws.append(EXAMPLES[doc_type])
    notes = wb.create_sheet("Notes and Instructions")
    notes.append(["Historical Import Template", doc_type]); notes.append(["Date format", "YYYY-MM-DD"])
    notes.append(["Required columns", ", ".join(sorted(REQUIRED[doc_type]))]); notes.append(["Allowed statuses", ", ".join(sorted(STATUSES[doc_type]))])
    notes.append(["References", "Use exact business document numbers and master-data codes. Missing references are never fabricated."])
    notes.append(["Multiple lines", f"Repeat {KEYS[doc_type]} for each line in the same document."])
    output = io.BytesIO(); wb.save(output); return output.getvalue()


def json_value(value):
    if isinstance(value, (datetime, date)): return value.isoformat()
    if isinstance(value, Decimal): return str(value)
    return value


def parse_file(path: str, original_name: str) -> list[tuple[int, str, dict]]:
    suffix = Path(original_name).suffix.lower(); rows = []
    if suffix == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            for number, raw in enumerate(csv.DictReader(handle), 2): rows.append((number, "CSV", {str(k).strip().lower(): json_value(v) for k,v in raw.items() if k}))
    elif suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True); sheet = book[book.sheetnames[0]]
        values = sheet.iter_rows(values_only=True); headers = [str(x).strip().lower() if x is not None else "" for x in next(values, [])]
        for number, cells in enumerate(values, 2):
            if not any(value is not None and str(value).strip() for value in cells): continue
            rows.append((number, sheet.title, {header: json_value(value) for header,value in zip(headers,cells) if header}))
    else: raise HTTPException(415, "Only XLSX and CSV files are supported")
    return rows


def text(value): return str(value).strip() if value is not None else ""
def code(value): return text(value).upper()
def decimal_value(value, field, errors):
    try: return Decimal(text(value) or "0")
    except (InvalidOperation, ValueError): errors.append(f"{field}: invalid decimal"); return Decimal("0")
def date_value(value, field, errors, required=False):
    if not text(value):
        if required: errors.append(f"{field}: required")
        return None
    try: return date.fromisoformat(text(value)[:10])
    except ValueError: errors.append(f"{field}: invalid date; use YYYY-MM-DD"); return None


class HistoricalImportService:
    def __init__(self, db: Session, user): self.db=db; self.user=user
    def normalize_validate(self, batch: ImportBatch):
        if not batch.stored_file_path: raise HTTPException(409, "Upload a file first")
        parsed=parse_file(batch.stored_file_path,batch.original_file_name or "")
        expected=set(SPECS[batch.import_type]); missing=REQUIRED[batch.import_type]-(set(parsed[0][2]) if parsed else set())
        if missing: raise HTTPException(422, f"Missing required columns: {', '.join(sorted(missing))}")
        batch.rows.clear(); self.db.flush(); groups=defaultdict(list)
        for number,sheet,raw in parsed:
            normalized={k:(code(v) if k.endswith("_code") or k in {"status","currency","creation_type","destination_type","payment_method"} else text(v)) for k,v in raw.items() if k in expected}
            row=ImportRow(row_number=number,sheet_name=sheet,document_type=batch.import_type,raw_data_json=raw,normalized_data_json=normalized,status="PENDING",document_number=text(normalized.get(KEYS[batch.import_type])))
            batch.rows.append(row); groups[row.document_number].append(row)
        self.db.flush()
        for number,rows in groups.items():
            errors=[]; warnings=[]
            if not number: errors.append(f"{KEYS[batch.import_type]}: required")
            self._validate_group(batch.import_type,number,rows,errors,warnings)
            state="ERROR" if errors else "WARNING" if warnings else "VALID"
            for row in rows: row.status=state; row.error_messages_json=list(errors); row.warning_messages_json=list(warnings)
        batch.total_rows=len(batch.rows); batch.total_documents=len(groups); batch.valid_rows=sum(r.status=="VALID" for r in batch.rows); batch.warning_rows=sum(r.status=="WARNING" for r in batch.rows); batch.error_rows=sum(r.status=="ERROR" for r in batch.rows)
        batch.status="VALIDATED"; batch.validated_at=datetime.now(timezone.utc); self.db.commit(); self.db.refresh(batch); return batch

    def _master(self, model, field, value): return self.db.query(model).filter(getattr(model,field)==value).first() if value else None
    def _validate_group(self, kind, number, rows, errors, warnings):
        first=rows[0].normalized_data_json
        for required in REQUIRED[kind]:
            if not any(text(r.normalized_data_json.get(required)) for r in rows): errors.append(f"{required}: required")
        status=code(first.get("status") or "DRAFT")
        if status not in STATUSES[kind]: errors.append(f"status: unsupported {status}")
        model,key={"PURCHASE_REQUISITION":(PurchaseRequisition,"pr_number"),"PURCHASE_ORDER":(PurchaseOrder,"po_number"),"GOODS_RECEIPT":(GoodsReceipt,"receipt_number"),"SUPPLIER_INVOICE":(SupplierInvoice,"internal_invoice_number"),"PAYMENT":(Payment,"payment_number"),"JOURNAL_ENTRY":(JournalEntry,"journal_number")}[kind]
        if self.db.query(model).filter(getattr(model,key)==number).first(): errors.append(f"Duplicate document number: {number}")
        date_field={"PURCHASE_REQUISITION":"document_date","PURCHASE_ORDER":"order_date","GOODS_RECEIPT":"receipt_date","SUPPLIER_INVOICE":"invoice_date","PAYMENT":"payment_date","JOURNAL_ENTRY":"journal_date"}[kind]
        date_value(first.get(date_field),date_field,errors,True)
        for row in rows:
            data=row.normalized_data_json
            if kind in {"PURCHASE_REQUISITION","PURCHASE_ORDER","GOODS_RECEIPT","SUPPLIER_INVOICE"}:
                item=self._master(Product,"code",code(data.get("item_code")))
                if not item: errors.append(f"Row {row.row_number}: unknown item_code {text(data.get('item_code'))}; canonical records require an existing item")
            if kind in {"PURCHASE_REQUISITION","PURCHASE_ORDER","SUPPLIER_INVOICE","JOURNAL_ENTRY"}:
                for label,model in (("department_code",Department),("location_code",Location)):
                    if text(data.get(label)) and not self._master(model,"code",code(data.get(label))): errors.append(f"Row {row.row_number}: unknown {label} {text(data.get(label))}")
            if kind in {"PURCHASE_REQUISITION","PURCHASE_ORDER"}: decimal_value(data.get("quantity"),"quantity",errors)
            if kind=="PURCHASE_ORDER":
                if not self._master(Supplier,"code",code(first.get("supplier_code"))): errors.append(f"Unknown supplier {text(first.get('supplier_code'))}")
                creation=code(first.get("creation_type")); pr=text(first.get("source_pr_number"))
                if creation not in {"DIRECT","FROM_PR"}: errors.append("creation_type must be DIRECT or FROM_PR")
                if creation=="FROM_PR" and (not pr or not self.db.query(PurchaseRequisition).filter_by(pr_number=pr).first()): errors.append(f"Exact source PR not found: {pr}")
                decimal_value(data.get("unit_cost"),"unit_cost",errors)
            elif kind=="GOODS_RECEIPT":
                po=self.db.query(PurchaseOrder).filter_by(po_number=text(first.get("po_number"))).first()
                if not po: errors.append(f"Exact PO not found: {text(first.get('po_number'))}")
                if code(first.get("destination_type"))=="WAREHOUSE" and not self._master(Warehouse,"code",code(first.get("warehouse_code"))): errors.append("WAREHOUSE requires a known warehouse_code")
                for row in rows: decimal_value(row.normalized_data_json.get("received_quantity"),"received_quantity",errors)
                if po:
                    for row in rows:
                        product=self._master(Product,"code",code(row.normalized_data_json.get("item_code"))); po_item=next((x for x in po.items if product and x.product_id==product.id),None); qty=decimal_value(row.normalized_data_json.get("received_quantity"),"received_quantity",errors)
                        if not po_item: errors.append(f"Item is not on PO: {text(row.normalized_data_json.get('item_code'))}")
                        elif qty>Decimal(po_item.quantity)-Decimal(po_item.received_quantity or 0): errors.append(f"Over-receipt for {product.code}")
            elif kind=="SUPPLIER_INVOICE":
                supplier=self._master(Supplier,"code",code(first.get("supplier_code")))
                if not supplier: errors.append(f"Unknown supplier {text(first.get('supplier_code'))}")
                elif self.db.query(SupplierInvoice).filter_by(supplier_id=supplier.id,supplier_invoice_number=text(first.get("supplier_invoice_number"))).first(): errors.append("Duplicate supplier invoice number for supplier")
                for row in rows:
                    account=self._master(ChartOfAccount,"account_number",code(row.normalized_data_json.get("account_code")))
                    if not account or not account.is_active or not account.is_posting: errors.append(f"Invalid posting account {text(row.normalized_data_json.get('account_code'))}")
                    decimal_value(row.normalized_data_json.get("quantity"),"quantity",errors);decimal_value(row.normalized_data_json.get("unit_price"),"unit_price",errors)
            elif kind=="PAYMENT":
                supplier=self._master(Supplier,"code",code(first.get("supplier_code"))); account=self._master(ChartOfAccount,"account_number",code(first.get("payment_account_code")))
                if not supplier: errors.append(f"Unknown supplier {text(first.get('supplier_code'))}")
                if not account or account.account_type!="ASSET" or not account.is_active or not account.is_posting: errors.append("Invalid payment account")
                total=decimal_value(first.get("total_amount"),"total_amount",errors); allocated=Decimal("0")
                for row in rows:
                    inv_number=text(row.normalized_data_json.get("supplier_invoice_number")); invoice=self.db.query(SupplierInvoice).filter_by(supplier_id=supplier.id if supplier else -1,supplier_invoice_number=inv_number).first(); amount=decimal_value(row.normalized_data_json.get("allocated_amount"),"allocated_amount",errors);allocated+=amount
                    if not invoice or invoice.status!="POSTED": errors.append(f"Eligible invoice not found: {inv_number}")
                    elif amount>Decimal(invoice.balance_amount): errors.append(f"Over-allocation for {inv_number}")
                if status in {"APPROVED","POSTED"} and allocated!=total: errors.append("Total allocations must equal payment total")
            elif kind=="JOURNAL_ENTRY":
                debit=credit=Decimal("0")
                for row in rows:
                    data=row.normalized_data_json; account=self._master(ChartOfAccount,"account_number",code(data.get("account_code"))); d=decimal_value(data.get("debit"),"debit",errors); c=decimal_value(data.get("credit"),"credit",errors)
                    if not account or not account.is_active or not account.is_posting: errors.append(f"Invalid posting account {text(data.get('account_code'))}")
                    if d<0 or c<0 or (d>0)==(c>0): errors.append(f"Row {row.row_number}: enter debit or credit, not both")
                    debit+=d;credit+=c
                if status=="POSTED" and debit!=credit: errors.append(f"Unbalanced journal: debit {debit}, credit {credit}")
        # Duplicate spreadsheet lines are warnings; exact duplicates remain visible for confirmation.
        signatures=[str(sorted(r.normalized_data_json.items())) for r in rows]
        if len(signatures)!=len(set(signatures)): warnings.append("Duplicate spreadsheet rows detected")

    def confirm(self,batch:ImportBatch,confirm_warnings=False):
        if batch.status!="VALIDATED": raise HTTPException(409,"Batch must be validated before confirmation")
        if batch.import_mode=="DRY_RUN":
            ActivityService.log(db=self.db,user=self.user,action="HISTORICAL_IMPORT_DRY_RUN",entity_type="IMPORT_BATCH",entity_id=batch.id,summary=f"Dry-run validated {batch.batch_number}",import_batch_id=batch.id);self.db.commit();return batch
        if batch.import_mode=="STRICT" and batch.error_rows: raise HTTPException(409,"STRICT mode blocks the entire batch while errors exist")
        if batch.warning_rows and not confirm_warnings: raise HTTPException(409,"Confirm warnings before import")
        batch.status="IMPORTING";self.db.commit(); groups=defaultdict(list)
        for row in batch.rows: groups[row.document_number].append(row)
        imported=failed=imported_documents=failed_documents=0
        for number,rows in groups.items():
            if any(r.status=="ERROR" for r in rows):
                for r in rows:r.status="SKIPPED"
                failed+=len(rows);failed_documents+=1;continue
            try:
                entity=self._import_group(batch,rows)
                now=datetime.now(timezone.utc)
                for r in rows:r.status="IMPORTED";r.entity_type=batch.import_type;r.entity_id=entity.id;r.imported_at=now
                imported+=len(rows);imported_documents+=1
            except Exception as exc:
                self.db.rollback();batch=self.db.get(ImportBatch,batch.id);rows=[self.db.get(ImportRow,r.id) for r in rows]
                for r in rows:r.status="FAILED";r.error_messages_json=[str(exc)]
                failed+=len(rows);failed_documents+=1
                if batch.import_mode=="STRICT": batch.status="FAILED";batch.failed_rows=failed;self.db.commit();raise HTTPException(409,f"Import failed for {number}: {exc}") from exc
        batch.imported_rows=imported;batch.failed_rows=failed;batch.success_count=imported_documents;batch.failed_count=failed_documents;batch.imported_at=datetime.now(timezone.utc);batch.status="IMPORTED" if failed==0 else "PARTIALLY_IMPORTED"
        ActivityService.log(db=self.db,user=self.user,action="HISTORICAL_IMPORT_CONFIRMED",entity_type="IMPORT_BATCH",entity_id=batch.id,summary=f"Imported {batch.batch_number}",details={"imported_rows":imported,"failed_rows":failed},import_batch_id=batch.id);self.db.commit();self.db.refresh(batch);return batch

    def _lookup(self,model,field,value): return self.db.query(model).filter(getattr(model,field)==value).one()
    def _metadata(self,batch): return {"source_type":"HISTORICAL_IMPORT","import_batch_id":batch.id,"imported_by_id":self.user.id,"imported_at":datetime.now(timezone.utc)}
    def _import_group(self,batch,rows):
        from app.api.v1 import journals, payments
        kind=batch.import_type;first=rows[0].normalized_data_json;number=rows[0].document_number;meta=self._metadata(batch);status=code(first.get("status") or "DRAFT")
        if kind=="PURCHASE_REQUISITION":
            department=self._master(Department,"code",code(first.get("department_code")));location=self._master(Location,"code",code(first.get("location_code")));requester=self._master(User,"username",text(first.get("requested_by_username")))
            entity=PurchaseRequisition(pr_number=number,status=status,request_date=date.fromisoformat(first["document_date"][:10]),expiry_date=date.fromisoformat(first["required_date"][:10]) if text(first.get("required_date")) else None,department_id=department.id if department else None,location_id=location.id if location else None,requested_by_id=requester.id if requester else None,notes=first.get("notes"),external_reference=first.get("external_reference"),total_amount=Decimal("0"),**meta)
            for row in rows:
                d=row.normalized_data_json;p=self._lookup(Product,"code",code(d.get("item_code")));q=Decimal(d["quantity"]);price=Decimal(d.get("estimated_unit_cost") or 0);entity.items.append(PurchaseRequisitionItem(product_id=p.id,quantity=q,estimated_price=price,total_price=q*price));entity.total_amount+=q*price
        elif kind=="PURCHASE_ORDER":
            supplier=self._lookup(Supplier,"code",code(first.get("supplier_code")));department=self._master(Department,"code",code(first.get("department_code")));location=self._master(Location,"code",code(first.get("location_code")));pr=self.db.query(PurchaseRequisition).filter_by(pr_number=text(first.get("source_pr_number"))).first() if code(first.get("creation_type"))=="FROM_PR" else None
            entity=PurchaseOrder(po_number=number,supplier_id=supplier.id,department_id=department.id if department else None,location_id=location.id if location else None,purchase_requisition_id=pr.id if pr else None,status=status,order_date=date.fromisoformat(first["order_date"][:10]),expected_date=date.fromisoformat(first["estimated_delivery_date"][:10]) if text(first.get("estimated_delivery_date")) else None,currency=code(first.get("currency") or "IDR"),discount_amount=Decimal(first.get("discount_amount") or 0),tax_amount=Decimal(first.get("tax_amount") or 0),total_amount=Decimal("0"),notes=first.get("notes"),external_reference=first.get("external_reference"),**meta)
            for row in rows:
                d=row.normalized_data_json;p=self._lookup(Product,"code",code(d.get("item_code")));q=Decimal(d["quantity"]);cost=Decimal(d["unit_cost"]);pr_item=next((x for x in pr.items if x.product_id==p.id),None) if pr else None;entity.items.append(PurchaseOrderItem(product_id=p.id,purchase_requisition_item_id=pr_item.id if pr_item else None,description=d.get("item_description"),uom=d.get("uom") or p.uom,quantity=q,received_quantity=0,unit_price=cost,total_price=q*cost));entity.total_amount+=q*cost
        elif kind=="GOODS_RECEIPT":
            po=self._lookup(PurchaseOrder,"po_number",text(first.get("po_number")));warehouse=self._master(Warehouse,"code",code(first.get("warehouse_code")));items=[]
            for row in rows:
                d=row.normalized_data_json;p=self._lookup(Product,"code",code(d.get("item_code")));po_item=next(x for x in po.items if x.product_id==p.id);items.append(GoodsReceiptItemCreate(purchase_order_item_id=po_item.id,quantity=Decimal(d["received_quantity"])))
            entity=GoodsReceiptService.create(self.db,data=GoodsReceiptCreate(purchase_order_id=po.id,received_at=datetime.combine(date.fromisoformat(first["receipt_date"][:10]),datetime.min.time(),tzinfo=timezone.utc),source_type="HISTORICAL_IMPORT",external_reference=first.get("external_reference"),destination_type=code(first.get("destination_type")),warehouse_id=warehouse.id if warehouse else None,notes=first.get("notes"),items=items),user=self.user);entity.receipt_number=number
            entity.delivery_note_number=first.get("delivery_note_number")
            for k,v in meta.items():setattr(entity,k,v)
            self.db.commit()
            if status=="POSTED":entity=GoodsReceiptService.post(self.db,receipt_id=entity.id,user=self.user)
            return entity
        elif kind=="SUPPLIER_INVOICE":
            supplier=self._lookup(Supplier,"code",code(first.get("supplier_code")));po=self.db.query(PurchaseOrder).filter_by(po_number=text(first.get("po_number"))).first();receipt=self.db.query(GoodsReceipt).filter_by(receipt_number=text(first.get("receipt_number"))).first();department=self._master(Department,"code",code(first.get("department_code")));location=self._master(Location,"code",code(first.get("location_code")))
            lines=[]
            for row in rows:
                d=row.normalized_data_json;p=self._lookup(Product,"code",code(d.get("item_code")));a=self._lookup(ChartOfAccount,"account_number",code(d.get("account_code")));lines.append(SupplierInvoiceLineInput(product_id=p.id,description=d["description"],quantity=Decimal(d["quantity"]),uom=d.get("uom"),unit_price=Decimal(d["unit_price"]),account_id=a.id,department_id=department.id if department else None))
            svc=SupplierInvoiceService(self.db);entity=svc.create(SupplierInvoiceInput(supplier_invoice_number=first["supplier_invoice_number"],supplier_id=supplier.id,invoice_date=date.fromisoformat(first["invoice_date"][:10]),received_date=date.fromisoformat(first["received_date"][:10]) if text(first.get("received_date")) else None,due_date=date.fromisoformat(first["due_date"][:10]) if text(first.get("due_date")) else None,purchase_order_id=po.id if po else None,goods_receipt_id=receipt.id if receipt else None,department_id=department.id if department else None,location_id=location.id if location else None,currency=code(first.get("currency") or "IDR"),discount_amount=Decimal(first.get("discount_amount") or 0),tax_amount=Decimal(first.get("tax_amount") or 0),withholding_tax_amount=Decimal(first.get("withholding_tax_amount") or 0),notes=first.get("notes"),source_type="HISTORICAL_IMPORT",lines=lines),self.user);entity.internal_invoice_number=number;entity.external_reference=first.get("external_reference")
            for k,v in meta.items():setattr(entity,k,v)
            self.db.commit()
            if status in {"VERIFIED","POSTED"}:entity=svc.verify(entity,self.user)
            if status=="POSTED":
                journal=svc.generate_journal(entity,self.user)
                for k,v in meta.items():setattr(journal,k,v)
                self.db.commit();journals.post(journal.id,self.db,self.user)
            return entity
        elif kind=="PAYMENT":
            supplier=self._lookup(Supplier,"code",code(first.get("supplier_code")));account=self._lookup(ChartOfAccount,"account_number",code(first.get("payment_account_code")));alloc=[]
            for row in rows:
                d=row.normalized_data_json;inv=self.db.query(SupplierInvoice).filter_by(supplier_id=supplier.id,supplier_invoice_number=d["supplier_invoice_number"]).one();alloc.append(PaymentAllocationInput(supplier_invoice_id=inv.id,allocated_amount=Decimal(d["allocated_amount"])))
            result=payments.create(PaymentCreate(supplier_id=supplier.id,payment_date=date.fromisoformat(first["payment_date"][:10]),payment_method=code(first["payment_method"]),payment_account_id=account.id,currency=code(first.get("currency") or "IDR"),total_amount=Decimal(first["total_amount"]),reference_number=first.get("reference_number"),cheque_number=first.get("cheque_number"),notes=first.get("notes"),source_type="HISTORICAL_IMPORT",allocations=alloc),self.db,self.user);entity=self.db.get(Payment,result["id"]);entity.payment_number=number;entity.external_reference=first.get("external_reference")
            for k,v in meta.items():setattr(entity,k,v)
            self.db.commit()
            if status in {"APPROVED","POSTED"}:payments.approve(entity.id,self.db,self.user)
            if status=="POSTED":
                payments.generate(entity.id,self.db,self.user);self.db.refresh(entity)
                for k,v in meta.items():setattr(entity.journal_entry,k,v)
                self.db.commit();payments.post(entity.id,self.db,self.user)
            return entity
        else:
            entity=JournalEntry(journal_number=number,posting_date=date.fromisoformat(first["journal_date"][:10]),status="DRAFT",document_type="HISTORICAL_IMPORT",document_number=first.get("reference_number") or number,reference_type=first.get("reference_type") or "HISTORICAL_IMPORT",reference_id=batch.id,description=first.get("description"),created_by_id=self.user.id,posted_at=None,external_reference=first.get("reference_number"),**meta)
            for index,row in enumerate(rows,1):
                d=row.normalized_data_json;a=self._lookup(ChartOfAccount,"account_number",code(d.get("account_code")));department=self._master(Department,"code",code(d.get("department_code")));location=self._master(Location,"code",code(d.get("location_code")));supplier=self._master(Supplier,"code",code(d.get("supplier_code")));entity.lines.append(JournalLine(line_number=index,account_id=a.id,debit=Decimal(d.get("debit") or 0),credit=Decimal(d.get("credit") or 0),description=d.get("line_description"),department_id=department.id if department else None,location_id=location.id if location else None,supplier_id=supplier.id if supplier else None))
        self.db.add(entity);self.db.flush();ActivityService.log(db=self.db,user=self.user,action="HISTORICAL_IMPORT_CREATED",entity_type=kind,entity_id=entity.id,summary=f"Imported {number}",import_batch_id=batch.id);self.db.commit();self.db.refresh(entity)
        if kind=="JOURNAL_ENTRY" and status=="POSTED":journals.post(entity.id,self.db,self.user)
        return entity
