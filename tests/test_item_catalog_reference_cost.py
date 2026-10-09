from datetime import datetime
from decimal import Decimal as D
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.models import Base
from app.models.product import Product
from app.models.goods_receipt import GoodsReceipt, GoodsReceiptItem
from app.models.chart_of_account import ChartOfAccount
from app.models.journal_entry import JournalEntry
from app.models.inventory import Inventory
from app.models.location import Location
from app.models.business_unit import BusinessUnit
from app.api.v1.products import list_products, lookup_products
from app.services.item_reference_cost import receipt_costs, reference_payload
from app.services.goods_receipt_service import GoodsReceiptService
from app.services.direct_coa_service import DirectCoaService
from app.services.approval import WORKFLOW_CHAINS, approve_current_step

ADMIN = SimpleNamespace(id=1,roles=[SimpleNamespace(name="SYSTEM_ADMINISTRATOR")],default_location_id=None,location_id=None)

@pytest.fixture
def db():
    engine=create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        unit=BusinessUnit(id=1,code="OPS",name="Operations")
        session.add(unit);session.flush()
        session.add_all([Location(id=1,code="L1",name="One",business_unit_id=1),Location(id=2,code="L2",name="Two",business_unit_id=1)])
        session.add_all([Product(id=1,code="1-100001",name="Rice",group_code="1",group_name="FOOD",subgroup_code="1-1",subgroup_name="MEAT",uom="Kg",usage_uom="Kg",inventory_treatment="STOCK",tax_category="NON_TAXABLE",cost_price=5),Product(id=2,code="1-100002",name="Rice",group_code="1",group_name="FOOD",subgroup_code="1-1",subgroup_name="MEAT",uom="Kg",usage_uom="Kg",inventory_treatment="STOCK",tax_category="NON_TAXABLE",catalog_type="HOTEL",cost_price=7)])
        session.commit()
        yield session
    engine.dispose()


def receipt(db,ident,year,qty,cost,status="POSTED",location=1):
    doc=GoodsReceipt(receipt_number=f"GR-{ident}",received_at=datetime(year,1,ident),status=status,location_id=location,business_unit_id=1,destination_type="WAREHOUSE")
    doc.items=[GoodsReceiptItem(product_id=1,quantity=qty,unit_cost=cost)]
    db.add(doc);db.commit()


def test_catalog_menu_and_lookup_do_not_mix_identical_names(db):
    for catalog,ident in [("CATERING",1),("HOTEL",2)]:
        report=list_products(catalog_type=catalog,page=1,page_size=50,db=db,user=ADMIN)
        assert [row["id"] for row in report["items"]] == [ident]
        lookup=lookup_products(catalog_type=catalog,item_type="BUY",search="Rice",limit=30,db=db,user=ADMIN)
        assert [row["id"] for row in lookup] == [ident]
    db.delete(db.get(Product,2));db.commit()
    assert list_products(catalog_type="HOTEL",page=1,page_size=50,db=db,user=ADMIN)["total_items"] == 0


def test_prices_use_posted_2026_receiving_and_weighted_quantity_read_only(db):
    receipt(db,1,2025,100,900)
    receipt(db,2,2026,2,10)
    receipt(db,3,2026,3,20)
    receipt(db,4,2026,100,1000,status="DRAFT")
    history=receipt_costs(db,[1],ADMIN)[1]
    assert history["last_purchase_price"] == D("20")
    assert history["average_purchase_price"] == D("16")
    product=db.get(Product,1)
    assert reference_payload(product,history)["reference_cost"] == D("20")
    product.reference_cost_method="AVERAGE_PURCHASE"
    assert reference_payload(product,history)["reference_cost"] == D("16")
    product.reference_cost_method="MANUAL";product.cost_price=D("17")
    assert reference_payload(product,history)["reference_cost"] == D("17")
    assert db.query(Inventory).count() == db.query(JournalEntry).count() == 0


def test_price_history_honors_assigned_location_without_partitioning_catalog(db):
    receipt(db,2,2026,2,10)
    receipt(db,3,2026,3,20,location=2)
    scoped=SimpleNamespace(roles=[],default_location_id=1,location_id=1)
    assert receipt_costs(db,[1],scoped)[1]["last_purchase_price"] == D("10")
    assert receipt_costs(db,[1],ADMIN)[1]["last_purchase_price"] == D("20")
    assert list_products(catalog_type="HOTEL",page=1,page_size=50,db=db,user=scoped)["total_items"] == 1


def test_direct_use_expense_override_preserves_stock_and_fixed_asset_routing(db):
    expense=ChartOfAccount(account_number="700-1",account_name="Hotel use",account_type="EXPENSE",normal_balance="DEBIT",is_posting=True,is_active=True)
    asset=ChartOfAccount(account_number="100-1",account_name="Inventory",account_type="ASSET",normal_balance="DEBIT",is_posting=True,is_active=True)
    db.add_all([expense,asset]);db.commit()
    resolver=DirectCoaService(db);product=db.get(Product,1)
    direct=SimpleNamespace(destination_type="DIRECT_USE",location_id=1)
    assert GoodsReceiptService._receiving_debit_account(resolver,direct,product,expense.id).id == expense.id
    for destination,treatment,account in [("WAREHOUSE","STOCK",expense.id),("DIRECT_USE","FIXED_ASSET",expense.id),("DIRECT_USE","STOCK",asset.id)]:
        product.inventory_treatment=treatment
        with pytest.raises(HTTPException):
            GoodsReceiptService._receiving_debit_account(resolver,SimpleNamespace(destination_type=destination,location_id=1),product,account)
    expense.is_active=False;product.inventory_treatment="STOCK"
    with pytest.raises(HTTPException):GoodsReceiptService._receiving_debit_account(resolver,direct,product,expense.id)


def test_receiving_needs_two_different_people_even_for_admin(monkeypatch):
    import app.services.approval as approval
    assert len(WORKFLOW_CHAINS["RR_APPROVAL"]) == 2
    step=SimpleNamespace(status="PENDING",approved_by_id=None)
    previous=SimpleNamespace(status="APPROVED",approved_by_id=1)
    request=SimpleNamespace(entity_type="GOODS_RECEIPT",steps=[previous,step])
    monkeypatch.setattr(approval,"_lock_request",lambda db,request:request)
    monkeypatch.setattr(approval,"_select_current_step",lambda request,user:step)
    monkeypatch.setattr(approval,"_advance_stage",lambda db,request:None)
    db=SimpleNamespace(flush=lambda:None)
    with pytest.raises(HTTPException) as error:approve_current_step(db,request=request,user=ADMIN)
    assert error.value.status_code == 409 and step.status == "PENDING"
    other=SimpleNamespace(id=2)
    approve_current_step(db,request=request,user=other)
    assert step.status == "APPROVED" and step.approved_by_id == 2


def test_invoice_snapshot_retains_receiving_expense_without_master_expense(db):
    from app.models.supplier_invoice import SupplierInvoiceLine
    expense=ChartOfAccount(account_number="700-2",account_name="Hotel supplies",account_type="EXPENSE",normal_balance="DEBIT",is_posting=True,is_active=True)
    db.add(expense);db.flush()
    receipt(db,5,2026,1,11)
    doc=db.query(GoodsReceipt).filter(GoodsReceipt.receipt_number=="GR-5").one()
    doc.destination_type="DIRECT_USE";doc.items[0].expense_account_id=expense.id
    db.commit()
    line=SupplierInvoiceLine(product_id=1,goods_receipt_item_id=doc.items[0].id,quantity=1,unit_price=11,line_amount=11)
    invoice=SimpleNamespace(goods_receipt=None,location_id=1)
    assert db.get(Product,1).purchase_account_id is None
    account=DirectCoaService(db).snapshot_invoice_line(invoice,line)
    assert account.id == expense.id and line.account_id == expense.id


def test_stock_opname_requires_department_even_when_user_has_default(db):
    from app.services.stock_opname_service import StockOpnameService
    user=SimpleNamespace(default_location_id=1,location_id=1,default_department_id=1,department_id=1)
    with pytest.raises(HTTPException) as error:
        StockOpnameService._resolve_dimensions(db,data=SimpleNamespace(location_id=1,department_id=None),current_user=user)
    assert error.value.status_code == 422 and "department" in error.value.detail


def test_export_hotel_is_empty_until_hotel_registration(db,tmp_path):
    from app.scripts.export_worker import _generate_products_xlsx
    from openpyxl import load_workbook
    db.delete(db.get(Product,2));db.commit()
    path=tmp_path/'hotel.xlsx'
    _generate_products_xlsx(db,path,{"catalog_type":"HOTEL"},ADMIN)
    wb=load_workbook(path)
    assert list(wb.active.values)[0][0] == "Code"
    assert len(list(wb.active.values)) == 1
    _generate_products_xlsx(db,path,{"catalog_type":"CATERING"},ADMIN)
    wb=load_workbook(path)
    assert len(list(wb.active.values)) == 2
    assert list(wb.active.values)[1][0] == '1-100001'


def test_hotel_direct_use_requires_explicit_department_and_expense(db):
    product=db.get(Product,2)
    item=SimpleNamespace(product=product,expense_account_id=None)
    doc=SimpleNamespace(destination_type='DIRECT_USE',department_id=None,items=[item])
    with pytest.raises(HTTPException):GoodsReceiptService._validate_hotel_expenses(doc)
    doc.department_id=1
    with pytest.raises(HTTPException):GoodsReceiptService._validate_hotel_expenses(doc)
    item.expense_account_id=8
    GoodsReceiptService._validate_hotel_expenses(doc)
    product.inventory_treatment='FIXED_ASSET';item.expense_account_id=None
    GoodsReceiptService._validate_hotel_expenses(doc)


def test_same_expense_account_tracks_two_consuming_departments(db):
    from app.models.accounting_setting import AccountingSetting
    expense=ChartOfAccount(account_number='700-500-104',account_name='Kebersihan dan Pest Control',account_type='EXPENSE',normal_balance='DEBIT',is_active=True,is_posting=True)
    suspense=ChartOfAccount(account_number='200-100-999',account_name='AP Suspense',account_type='LIABILITY',normal_balance='CREDIT',is_active=True,is_posting=True)
    db.add_all([expense,suspense]);db.flush()
    db.add(AccountingSetting(setting_key='AP_SUSPENSE_ACCOUNT',account_id=suspense.id));db.commit()
    for department,day in [(11,6),(12,7)]:
        doc=GoodsReceipt(receipt_number=f'HOTEL-{department}',received_at=datetime(2026,1,day),destination_type='DIRECT_USE',department_id=department,location_id=1,business_unit_id=1,subtotal=20)
        doc.items=[GoodsReceiptItem(product_id=2,quantity=2,unit_cost=10,expense_account_id=expense.id)]
        db.add(doc);db.flush()
        journal=GoodsReceiptService._sync_accounting_journal(db,receipt=doc,user=ADMIN)
        db.flush()
        assert len(journal.lines)==2
        assert journal.lines[0].account_id == expense.id
        assert journal.lines[0].debit == 20
        assert all(line.department_id == department for line in journal.lines)
        assert journal.lines[1].account_id == suspense.id and journal.lines[1].credit == 20
    assert db.query(Inventory).count() == 0


def test_po_remaining_items_has_catalog_without_receiving_only_column(db):
    from app.models.purchase_order import PurchaseOrder
    from app.models.purchase_order_item import PurchaseOrderItem
    from app.api.v1.goods_receipts import remaining_items
    po=PurchaseOrder(po_number='PO-TEST')
    po.items=[PurchaseOrderItem(product=db.get(Product,2),quantity=5,received_quantity=2,unit_price=10)]
    rows=remaining_items(po)
    assert rows[0]['catalog_type']=='HOTEL' and rows[0]['remaining_quantity']==3
