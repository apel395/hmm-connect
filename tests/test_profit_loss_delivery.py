from datetime import date
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi import HTTPException
from app.models import Base
from app.models.business_unit import BusinessUnit
from app.models.location import Location
from app.models.department import Department
from app.models.location_department import LocationDepartment
from app.models.chart_of_account import ChartOfAccount
from app.models.journal_entry import JournalEntry
from app.models.journal_line import JournalLine
from app.models.accounting_setting import AccountingSetting
from app.services.financial_reporting_service import FinancialReportingService
from app.services.profit_loss_delivery import deliver, workbook_sheets
from app.services.xlsx_export import build_xlsx

@pytest.fixture
def records():
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, expire_on_commit=False)()
    hotel = BusinessUnit(code='HOTEL', name='Hotel')
    catering = BusinessUnit(code='CATERING', name='Project Catering')
    db.add_all([hotel, catering]); db.flush()
    h = Location(code='H', name='Hotel', business_unit_id=hotel.id)
    c1 = Location(code='C1', name='Location one', business_unit_id=catering.id)
    c2 = Location(code='C2', name='Location two', business_unit_id=catering.id)
    d = Department(code='HK', name='House Keeping', business_unit_id=hotel.id)
    accounts = [ChartOfAccount(account_number=f'700-{i}', account_name=name, account_type='EXPENSE', normal_balance='DEBIT', is_posting=True, financial_statement_group=group) for i,(name,group) in enumerate([('Cost','COST_OF_REVENUE'),('Salary','OPERATING_EXPENSE'),('Supplies','OPERATING_EXPENSE'),('Admin salary','OPERATING_EXPENSE'),('Admin expenses','OPERATING_EXPENSE'),('Other','OTHER_EXPENSE'),('Unknown',None)])]
    revenue = ChartOfAccount(account_number='400', account_name='Revenue', account_type='REVENUE', normal_balance='CREDIT', is_posting=True)
    db.add_all([h,c1,c2,d,revenue,*accounts]); db.flush()
    db.add(LocationDepartment(location_id=h.id,department_id=d.id))
    for acc, cat in zip(accounts[1:5],['OPERATING_PAYROLL','OPERATING_EXPENSE','OVERHEAD_PAYROLL','OVERHEAD_EXPENSE']):
        db.add(AccountingSetting(setting_key=f'PL_CATEGORY:{acc.id}', account_id=acc.id, description=cat))
    for idx,(loc,dept,mult) in enumerate([(h,d.id,1),(c1,None,2),(c2,None,3)]):
        j=JournalEntry(journal_number=f'PL-{idx}',posting_date=date(2026,10,5),status='POSTED',reference_type='MANUAL',reference_id=0,location_id=loc.id,business_unit_id=loc.business_unit_id,department_id=dept)
        j.lines=[JournalLine(account_id=revenue.id,debit=0,credit=1000*mult)] + [JournalLine(account_id=a.id,debit=amount*mult,credit=0) for a,amount in zip(accounts,[100,50,40,30,20,10,5])]
        db.add(j)
    db.commit();yield db,h,c1,c2,d,accounts
    db.close();engine.dispose()


def subtotal(report,key):
    return next(r for r in report['presentation_rows'] if r['key']==key)['periods']['MTD']['current']


def test_scopes_and_management_subtotals_keep_every_expense(records):
    db,h,c1,c2,d,accounts=records
    service=FinancialReportingService(db)
    hotel=deliver(service,as_of_date=date(2026,10,10),report_scope='HOTEL')
    assert subtotal(hotel,'GOI')==Decimal(810)
    assert subtotal(hotel,'GOP')==Decimal(760)
    assert subtotal(hotel,'NOP')==Decimal(745)
    assert len(hotel['group_reports'])==1
    assert hotel['group_reports'][0]['department_name']=='House Keeping'
    catering=deliver(service,as_of_date=date(2026,10,10),report_scope='CATERING')
    assert subtotal(catering,'NOP')==Decimal(3725)
    assert {r['scope']['location_id'] for r in catering['group_reports']}=={c1.id,c2.id}
    consolidated=deliver(service,as_of_date=date(2026,10,10))
    assert subtotal(consolidated,'NOP')==subtotal(hotel,'NOP')+subtotal(catering,'NOP')
    assert all(gap==0 for vals in consolidated['split_control'].values() for gap in vals.values())
    assert db.query(JournalEntry).count()==3


def test_excel_summary_first_one_sheet_per_dimension_both_periods(records):
    db,*_=records
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10))
    sheets=workbook_sheets(report)
    assert sheets[0][0]=='Summary'
    assert len(sheets)==7
    assert len(set(name.casefold() for name,_ in sheets))==7
    assert all(len(name)<=31 for name,_ in sheets)
    assert all('MTD Current' in rows[5] and 'YTD Current' in rows[5] for _,rows in sheets)
    with ZipFile(BytesIO(build_xlsx(sheets))) as x:
        root=ET.fromstring(x.read('xl/workbook.xml'))
        assert len(root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets'))==7
        assert b'700-0' not in x.read('xl/worksheets/sheet1.xml')


def test_location_scope_and_hidden_details_do_not_change_totals(records):
    db,h,c1,c2,d,accounts=records
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='CONSOLIDATED',location_id=c1.id,include_breakdown=False)
    assert subtotal(report,'NOP')==1490
    assert not report['group_reports']
    with pytest.raises(HTTPException):
        deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='HOTEL',location_id=c1.id)


def test_report_setup_mapping_is_explicit_and_totals_ignore_hidden_details(records):
    from app.models.financial_report_setting import FinancialReportSetting
    db,h,c1,c2,d,accounts=records
    db.add(FinancialReportSetting(report_type='PROFIT_LOSS',selected_account_ids=[accounts[0].id],header_only=False))
    db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='HOTEL')
    assert subtotal(report,'NOP')==745
    details=[r for r in report['presentation_rows'] if r['kind']=='detail']
    assert {aid for r in details for aid in r['account_ids']}=={accounts[0].id}
    assert next(r for r in report['presentation_rows'] if r['key']=='UNDER_EXPENSE')['periods']['MTD']['current']==15


def test_organization_gap_is_reported_not_redistributed(records):
    db,h,c1,c2,d,accounts=records
    journal=JournalEntry(journal_number='UNASSIGNED',posting_date=date(2026,10,5),status='POSTED',reference_type='MANUAL',reference_id=0)
    journal.lines=[JournalLine(account_id=accounts[0].id,debit=17,credit=0)]
    db.add(journal);db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10))
    assert report['split_control']['MTD']['COST_OF_REVENUE']==17
    assert subtotal(report,'NOP')==4453
    assert db.query(JournalEntry).count()==4


def test_exact_summary_order_and_drilldown_includes_zero_accounts(records):
    db,*_=records
    zero=ChartOfAccount(account_number='700-999',account_name='Zero operations expense',account_type='EXPENSE',normal_balance='DEBIT',is_posting=True,financial_statement_group='OPERATING_EXPENSE')
    db.add(zero);db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='HOTEL')
    assert [r['key'] for r in report['presentation_rows'] if r['kind']!='detail']==['REVENUE','COST_OF_REVENUE','OPERATING_PAYROLL','OPERATING_EXPENSE','GOI','OVERHEAD_PAYROLL','OVERHEAD_EXPENSE','GOP','UNDER_EXPENSE','NOP']
    row=next(r for r in report['related_accounts'] if r['account_id']==zero.id)
    assert row['periods']['MTD']['current']==0
    assert zero.id in next(r for r in report['presentation_rows'] if r['key']=='OPERATING_EXPENSE')['account_ids']
    assert zero.id in next(r for r in report['presentation_rows'] if r['key']=='NOP')['account_ids']


def test_excel_export_ignores_hidden_accounts_and_includes_zero_department_accounts(records):
    from app.models.financial_report_setting import FinancialReportSetting
    from app.api.v1.financial_reports import export_comparison
    from types import SimpleNamespace
    import asyncio
    from io import BytesIO
    from openpyxl import load_workbook
    db,h,*_=records
    zero=ChartOfAccount(account_number='700-ZERO',account_name='Unused department expense',account_type='EXPENSE',normal_balance='DEBIT',is_posting=True,financial_statement_group='OPERATING_EXPENSE')
    db.add(zero);db.add(FinancialReportSetting(report_type='PROFIT_LOSS',selected_account_ids=[],header_only=True));db.commit()
    admin=SimpleNamespace(id=1,roles=[SimpleNamespace(name='SYSTEM_ADMINISTRATOR',permissions=[])])
    response=export_comparison(report_type='PROFIT_LOSS',as_of_date=date(2026,10,10),report_scope='HOTEL',include_zero=False,kap=True,db=db,user=admin)
    async def collect():return b''.join([chunk async for chunk in response.body_iterator])
    workbook=load_workbook(BytesIO(asyncio.run(collect())))
    rows=list(workbook[next(name for name in workbook.sheetnames if 'House Keeping' in name)].values)
    assert any(str(row[0]).strip()=='Unused department expense' and row[1]==0 and row[10]==0 for row in rows)
    assert any(str(row[0]).strip()=='Cost' and row[1]==100 for row in rows)
    assert any(str(row[0]).strip()=='Salaries & Wages' and row[1]==80 for row in rows)
    assert any(str(row[0]).strip()=='Revenue' and row[1]==1000 for row in rows)
    assert workbook.sheetnames[0]=='Summary'


def test_template_export_subtotals_match_report_and_keep_missing_comparisons(records):
    from app.services.profit_loss_workbook import department_rows, summary_rows
    db, *_ = records
    report = deliver(FinancialReportingService(db), as_of_date=date(2026,10,10))
    for statement in [report, *report['unit_reports'], *report['group_reports']]:
        rows = department_rows(statement, 'Test', report['as_of_date'])
        net = next(r for r in rows if r and r[0] == 'Net Operating Profit / Loss')
        expected = next(r for r in statement['presentation_rows'] if r['key'] == 'NOP')
        assert net[1] == expected['periods']['MTD']['current']
        assert net[10] == expected['periods']['YTD']['current']
        assert net[3] is None  # no budget entered
        assert net[7] is None  # no prior-year financial evidence
    summary = summary_rows(report, 'Summary', report['group_reports'], report['as_of_date'])
    net = next(r for r in summary if r and r[0] == 'NOP')
    assert net[1] == subtotal(report, 'NOP')


def test_template_service_charge_is_counted_once_and_combined_payroll_stays_intact(records):
    from app.services.profit_loss_workbook import department_rows, service_charge
    db,h,c1,c2,d,_ = records
    sc = ChartOfAccount(account_number='400-SC', account_name='Service Charge Revenue', account_type='REVENUE', normal_balance='CREDIT', is_posting=True)
    wage = ChartOfAccount(account_number='600-SC', account_name='Service Charge Distribution', account_type='EXPENSE', normal_balance='DEBIT', is_posting=True, financial_statement_group='OPERATING_PAYROLL')
    db.add_all([sc,wage]);db.flush()
    journal = JournalEntry(journal_number='SC-QA', posting_date=date(2026,10,5), status='POSTED',reference_type='MANUAL',reference_id=0,location_id=h.id,business_unit_id=h.business_unit_id,department_id=d.id)
    journal.lines = [JournalLine(account_id=sc.id,debit=0,credit=100),JournalLine(account_id=wage.id,debit=60,credit=0)]
    db.add(journal);db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='HOTEL')
    rows=department_rows(report,'Hotel',report['as_of_date'])
    get=lambda label:next(r for r in rows if r and r[0]==label)
    assert get('Net Revenue Without Service Charge')[1] == 1000
    assert get('Service Charge Revenue')[1] == 100
    assert get('Net Revenue With Service Charge')[1] == 1100
    assert get('Service Charge Distribution')[1] == 60
    assert get('Net Operating Profit / Loss')[1] == subtotal(report,'NOP') == 785
    assert not service_charge({'account_name':'Salary, Wages & Service Charge'})


def test_template_variance_is_last_month_and_percentages_use_matching_revenue(records):
    from app.services.profit_loss_workbook import department_rows
    db,*_=records
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='HOTEL')
    # Distinct comparisons prove columns do not reuse the budget variance.
    for account in report['related_accounts']:
        for p in ('MTD','YTD'):
            account['periods'][p]['last_month']=account['periods'][p]['current']/2
            account['periods'][p]['budget']=account['periods'][p]['current']*2
    for row in report['presentation_rows']:
        for p in ('MTD','YTD'):
            row['periods'][p]['budget']=Decimal(0)
    rows=department_rows(report,'Hotel',report['as_of_date'])
    row=next(r for r in rows if r and r[0]=='Net Revenue With Service Charge')
    assert row[1:7] == [1000,1,2000,1,500,1]
    assert row[9] == 500
    assert row[14] == 500


def test_template_excel_layout_print_and_numeric_formats(records, tmp_path):
    from openpyxl import load_workbook
    db,*_=records
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10))
    workbook=load_workbook(BytesIO(build_xlsx(workbook_sheets(report))))
    for sheet in workbook:
        assert sheet.freeze_panes=='B7'
        assert sheet.print_title_rows=='$1:$6'
        assert sheet.page_setup.orientation=='landscape'
        assert sheet.sheet_properties.pageSetUpPr.fitToPage
        assert sheet.auto_filter.ref is None
        assert sheet['B6'].value=='MTD Current'
        headers=[c.value for c in sheet[6]]
        assert 'YTD Current' in headers and 'MTD Variance' in headers
        for row in sheet:
            for cell in row:
                assert cell.data_type != 'e'
        revenue=next(row for row in sheet if row[0].value in ('NET REVENUE WITH SERVICE CHARGE', 'Net Revenue With SC', 'Total Revenue', 'Pendapatan'))
        assert revenue[1].data_type=='n'
        if revenue[2].value is not None:
            assert '%' in revenue[2].number_format
    # Build the real production bytes for optional read-only visual QA.
    import os
    if path:=os.environ.get('PL_TEMPLATE_QA_OUTPUT'):
        from pathlib import Path
        Path(path).write_bytes(build_xlsx(workbook_sheets(report)))


def test_hotel_summary_matches_page_and_departments_have_stable_order(records):
    db,h,c1,c2,d,_=records
    for code,name in [('FO','Front Office'),('FB','Food and Beverage'),('ACC','Accounting')]:
        dept=Department(code=code,name=name,business_unit_id=h.business_unit_id)
        db.add(dept);db.flush()
        db.add(LocationDepartment(location_id=h.id,department_id=dept.id))
    db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='HOTEL')
    sheets=workbook_sheets(report)
    summary=sheets[0][1]
    assert next(r[1] for r in summary if r and r[0]=='GROSS OPERATING PROFIT')==subtotal(report,'GOP')
    assert next(r[1] for r in summary if r and r[0]=='NETT OPERATING PROFIT')==subtotal(report,'NOP')
    assert [name for name,_ in sheets[1:]]==['Hotel · Front Office','Hotel · Food and Beverage','Hotel · House Keeping','Hotel · Accounting','Hotel Utility']


def test_related_zeros_are_scoped_but_wrong_and_offsetting_postings_remain(records):
    db,h,c1,c2,d,_=records
    room=ChartOfAccount(account_number='400-100-100',account_name='Room Sales',account_type='REVENUE',normal_balance='CREDIT',is_posting=True)
    catering=ChartOfAccount(account_number='400-200-100',account_name='Pendapatan Project Catering',account_type='REVENUE',normal_balance='CREDIT',is_posting=True)
    fo=ChartOfAccount(account_number='600-FO',account_name='FO - Salary',account_type='EXPENSE',normal_balance='DEBIT',is_posting=True)
    hk=ChartOfAccount(account_number='600-HK',account_name='HK - Salary',account_type='EXPENSE',normal_balance='DEBIT',is_posting=True)
    db.add_all([room,catering,fo,hk]);db.commit()
    options=dict(as_of_date=date(2026,10,10),include_zero=True,all_accounts=True)
    hotel=deliver(FinancialReportingService(db),report_scope='HOTEL',**options)
    house=hotel['group_reports'][0]
    assert hk.id in {a['account_id'] for a in house['related_accounts']}
    assert not {room.id,catering.id,fo.id}&{a['account_id'] for a in house['related_accounts']}
    project=deliver(FinancialReportingService(db),report_scope='CATERING',**options)
    assert room.id not in {a['account_id'] for a in project['related_accounts']}
    assert catering.id in {a['account_id'] for a in project['related_accounts']}
    # Preserve misclassified actuals, even when debit/credit cancel to zero.
    for n,amount in enumerate((10,-10)):
        journal=JournalEntry(journal_number=f'WRONG-ROOM-{n}',posting_date=date(2026,10,5),status='POSTED',reference_type='MANUAL',reference_id=0,location_id=c1.id,business_unit_id=c1.business_unit_id)
        journal.lines=[JournalLine(account_id=room.id,debit=max(-amount,0),credit=max(amount,0))]
        db.add(journal)
    db.commit()
    project=deliver(FinancialReportingService(db),report_scope='CATERING',**options)
    assert room.id in {a['account_id'] for a in project['related_accounts']}
    assert db.query(JournalEntry).count()==5


def test_catering_workbook_has_one_location_schedule_and_no_hotel_boilerplate(records):
    db,h,c1,c2,d,accounts=records
    room=ChartOfAccount(account_number='400-100-100',account_name='Room Sales',account_type='REVENUE',normal_balance='CREDIT',is_posting=True)
    db.add(room);db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='CATERING',all_accounts=True,include_zero=True)
    sheets=workbook_sheets(report)
    assert [name for name,_ in sheets]==['Summary','Project Catering · Location one','Project Catering · Location two']
    for (name,rows),statement in zip(sheets[1:],report['group_reports']):
        labels=[r[0] for r in rows if r]
        assert 'Room Sales' not in [label.strip() for label in labels]
        assert 'OPERATING STATISTICS' not in labels
        assert 'Net Revenue With Service Charge' not in labels
        assert 'Service Charge Distribution' in labels
        assert rows.template_profile['source_sheet']=='PC LIPET'
        assert statement['scope_labels']['location_id'].upper() in rows[2][0]
        net=next(r for r in rows if r and r[0]=='Gross Operating Income / Profit / Loss')
        assert net[1]==subtotal(statement,'NOP')
        expected=next(r for r in statement['presentation_rows'] if r['key']=='NOP')['periods']
        assert net[rows.headers.index('YTD Current')]==expected['YTD']['current']
        assert net[rows.headers.index('MTD Budget')] is None
        assert net[rows.headers.index('MTD Last Year')] is None
    assert sum(next(r for r in rows if r and r[0]=='Gross Operating Income / Profit / Loss')[1] for _,rows in sheets[1:])==subtotal(report,'NOP')
    assert next(r[1] for r in sheets[0][1] if r and r[0]=='NOP')==subtotal(report,'NOP')
    assert db.query(JournalEntry).count()==3


def test_consolidated_export_dispatches_hotel_and_catering_layouts_separately(records):
    from openpyxl import load_workbook
    db,*_=records
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10))
    workbook=load_workbook(BytesIO(build_xlsx(workbook_sheets(report))))
    hotel=workbook[next(n for n in workbook.sheetnames if 'House Keeping' in n)]
    location=workbook[next(n for n in workbook.sheetnames if 'Location one' in n)]
    assert hotel['A3'].value=='HOUSE KEEPING DEPARTEMEN'
    assert location['A3'].value=='PROJECT CATERING DEPARTEMEN — LOCATION ONE'
    assert next(row[1].value for row in location if row[0].value=='Gross Operating Income / Profit / Loss')==1490
    assert location.freeze_panes=='B7'


def test_catering_service_charge_detail_is_once_and_cannot_inherit_other_location(records):
    db,h,c1,c2,d,_=records
    sc=ChartOfAccount(account_number='400-SC',account_name='Service Charge Revenue',account_type='REVENUE',normal_balance='CREDIT',is_posting=True)
    db.add(sc);db.flush()
    journal=JournalEntry(journal_number='PROJECT-SC',posting_date=date(2026,10,5),status='POSTED',reference_type='MANUAL',reference_id=0,location_id=c1.id,business_unit_id=c1.business_unit_id)
    journal.lines=[JournalLine(account_id=sc.id,debit=0,credit=50)]
    db.add(journal);db.commit()
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10),report_scope='CATERING')
    sheets=workbook_sheets(report)
    for (_,rows),expected in zip(sheets[1:],(50,0)):
        assert next(r for r in rows if r and r[0]=='Other Revenue Service Charge')[1]==expected
    assert next(r for r in sheets[1][1] if r and r[0]=='Gross Operating Income / Profit / Loss')[1]==1540


def test_removing_beban_statement_row_preserves_executive_expense_card(records):
    from types import SimpleNamespace
    from app.api.v1.financial_reports import executive_overview
    db,*_=records
    admin=SimpleNamespace(id=1,roles=[SimpleNamespace(name='SYSTEM_ADMINISTRATOR',permissions=[])])
    overview=executive_overview(as_of_date=date(2026,10,10),db=db,user=admin)
    expenses=next(r for r in overview['profit_loss']['presentation_rows'] if r['key']=='BEBAN')
    assert expenses['periods']['MTD']['current']=='1530.00'
    assert expenses['periods']['MTD']['budget'] is None
    report=deliver(FinancialReportingService(db),as_of_date=date(2026,10,10))
    assert not any(r['key']=='BEBAN' for r in report['presentation_rows'])
