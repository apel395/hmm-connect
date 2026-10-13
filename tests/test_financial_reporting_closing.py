from datetime import date
from decimal import Decimal
from types import SimpleNamespace
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from app.models import Base
from app.models.accounting_setting import AccountingSetting
from app.models.chart_of_account import ChartOfAccount
from app.models.fiscal_period import FiscalPeriod
from app.models.journal_entry import JournalEntry
from app.models.journal_line import JournalLine
from app.models.general_ledger import GeneralLedger
from app.models.opening_balance import OpeningBalance
from app.models.budget import Budget, BudgetLine
from app.models.financial_report_setting import FinancialReportSetting
from app.services.period_closing_service import PeriodClosingService
from app.services.financial_reporting_service import FinancialReportingService, periods
from app.services.financial_statement_service import FinancialStatementService
from app.services.trial_balance_service import TrialBalanceService, ZERO
from app.api.v1 import journals

USER = SimpleNamespace(id=1, default_location_id=None, location_id=None, roles=[])

@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    accounts = [ChartOfAccount(account_number=code, account_name=name, account_type=kind,
                 normal_balance=normal, financial_statement_group=group, is_posting=True, is_active=True)
                for code, name, kind, normal, group in (
                    ("400", "Pendapatan", "REVENUE", "CREDIT", "OPERATING_REVENUE"),
                    ("700-100", "Expense", "EXPENSE", "DEBIT", "OPERATING_EXPENSE"),
                    ("100-100", "Cash", "ASSET", "DEBIT", "CURRENT_ASSET"),
                    ("300-300-100", "Current year profit", "EQUITY", "CREDIT", "EQUITY"),
                    ("100-800-103", "Depreciation", "ASSET", "CREDIT", "NON_CURRENT_ASSET"),
                )]
    session.add_all(accounts)
    session.flush()
    session.add(AccountingSetting(setting_key="CURRENT_YEAR_EARNINGS_ACCOUNT", account_id=accounts[3].id))
    session.add_all([FiscalPeriod(fiscal_year=2026, fiscal_month=month, period_name=f"Month {month}",
                   start_date=date(2026, month, 1), end_date=date(2026, month, 31 if month == 1 else 28),
                   is_open=True, is_closed=False, allow_backdated_entries=False) for month in (1, 2)])
    session.commit()
    yield session, accounts
    session.close()


def activity(db, accounts, stamp, revenue, expense, location=1, status="POSTED", unit=1):
    journal = JournalEntry(journal_number=f"TEST-{stamp}-{location}-{db.query(JournalEntry).count()}",
        posting_date=stamp, status=status, reference_type="MANUAL", reference_id=0,
        business_unit_id=unit, location_id=location)
    journal.lines = [JournalLine(account_id=accounts[0].id, credit=Decimal(revenue), debit=0, business_unit_id=unit, location_id=location),
        JournalLine(account_id=accounts[1].id, debit=Decimal(expense), credit=0, business_unit_id=unit, location_id=location),
        JournalLine(account_id=accounts[2].id, debit=Decimal(revenue)-Decimal(expense), credit=0, business_unit_id=unit, location_id=location)]
    db.add(journal)
    db.commit()
    return journal


def test_close_reopen_and_reclose_are_balanced_and_preserve_pnl_and_opening(db):
    session, a = db
    activity(session, a, date(2026, 1, 10), "100", "30")
    activity(session, a, date(2026, 1, 11), "40", "60", location=2, unit=2)
    activity(session, a, date(2026, 1, 12), "900", "200", status="DRAFT")
    opening = OpeningBalance(fiscal_year=2026, account_id=a[2].id, location_id=1, amount=Decimal("1000"), balance_type="DEBIT")
    session.add(opening); session.commit()
    before = FinancialStatementService(session).balance_sheet(as_of_date=date(2026,1,31))["summary"]
    result = PeriodClosingService.close(session, 1, USER)
    session.commit()
    assert len(result["closing_journal_ids"]) == 2
    for jid in result["closing_journal_ids"]:
        j = session.get(JournalEntry, jid)
        assert j.posting_date == date(2026,1,31)
        assert j.fiscal_year == 2026
        assert sum(l.debit-l.credit for l in j.lines) == 0
        assert all(l.debit >= 0 and l.credit >= 0 for l in j.lines)
    assert PeriodClosingService.close(session, 1, USER)["closing_journal_ids"] == result["closing_journal_ids"]
    report = FinancialStatementService(session).profit_and_loss(date_from=date(2026,1,1), date_to=date(2026,1,31))
    assert report["summary"]["net_profit_after_tax"] == Decimal("50")
    after = FinancialStatementService(session).balance_sheet(as_of_date=date(2026,1,31))["summary"]
    assert before["total_assets"] == after["total_assets"]
    assert before["total_equity"] == after["total_equity"]
    assert session.get(OpeningBalance, opening.id).amount == Decimal("1000")
    PeriodClosingService.reopen(session, 1, USER, "Correct January")
    session.commit()
    assert session.query(GeneralLedger).filter_by(reference_type="PERIOD_CLOSE").count() == 0
    assert all(session.get(JournalEntry, jid).status == "VOID" for jid in result["closing_journal_ids"])
    result2 = PeriodClosingService.close(session, 1, USER)
    assert set(result2["closing_journal_ids"]).isdisjoint(result["closing_journal_ids"])


def test_current_year_equity_is_one_account_with_opening_and_unclosed_evidence(db):
    session, a = db
    a[3].account_name = "Laba/Rugi Ditahan Tahun Berjalan"
    opening = OpeningBalance(fiscal_year=2026, account_id=a[3].id, location_id=1,
                             amount=Decimal("500"), balance_type="CREDIT")
    session.add(opening); session.commit()
    activity(session, a, date(2026, 1, 10), "100", "30")
    report = FinancialStatementService(session).balance_sheet(as_of_date=date(2026, 1, 31))
    equity = [r for s in report["sections"] if s["code"] == "EQUITY" for g in s["groups"] for r in g["accounts"]]
    assert len(equity) == 1
    assert equity[0]["account_id"] == a[3].id
    assert equity[0]["closing_amount"] == Decimal("570")
    evidence = report["summary"]["earnings_reconciliation"]
    assert evidence["opening_balance"] == Decimal("500")
    assert evidence["posted_movement"] == ZERO
    assert evidence["unclosed_profit_loss"] == Decimal("70")
    PeriodClosingService.close(session, 1, USER); session.commit()
    after = FinancialStatementService(session).balance_sheet(as_of_date=date(2026, 1, 31))
    assert after["summary"]["total_equity"] == report["summary"]["total_equity"]
    assert after["summary"]["earnings_reconciliation"]["posted_movement"] == Decimal("70")
    assert after["summary"]["earnings_reconciliation"]["unclosed_profit_loss"] == ZERO
    assert session.get(OpeningBalance, opening.id).amount == Decimal("500")


def test_executive_equity_detail_and_kap_share_totals_and_actual_account_ids(db):
    session, a = db
    activity(session, a, date(2026, 1, 10), "100", "30")
    statements = FinancialReportingService(session)
    report = statements.build(report_type="BALANCE_SHEET", as_of_date=date(2026, 1, 31))
    kap = statements.build(report_type="BALANCE_SHEET", as_of_date=date(2026, 1, 31), kap=True)
    assert kap["section_totals"] == report["section_totals"]
    row = next(r for r in report["rows"] if r["section"] == "EQUITY")
    assert row["account_ids"] == [a[3].id]
    assert row["includes_unclosed_earnings"]
    for section in ("ASSETS", "LIABILITIES", "EQUITY"):
        total = sum(r["periods"]["YTD"]["current"] for r in report["executive_rows"] if r["section"] == section)
        assert total == report["section_totals"][section]["YTD"]["current"]


def test_mtd_ytd_scoped_comparisons_and_budget(db):
    session, a = db
    activity(session, a, date(2026,1,10), "100", "30")
    activity(session, a, date(2026,2,10), "200", "80")
    activity(session, a, date(2026,2,12), "900", "300", location=2, unit=2)
    activity(session, a, date(2025,2,10), "150", "60")
    b = Budget(budget_number="B1", name="Plan", fiscal_year=2026, date_from=date(2026,1,1), date_to=date(2026,12,31), status="APPROVED")
    b.lines = [BudgetLine(account_id=a[0].id, business_unit_id=1, location_id=1, month=m, budget_amount=Decimal("250")) for m in (1,2)]
    session.add(b); session.commit()
    result = FinancialReportingService(session).build(report_type="PROFIT_LOSS", as_of_date=date(2026,2,28), location_id=1, business_unit_id=1)
    revenue = next(r for r in result["rows"] if r["account_code"] == "400")
    assert revenue["periods"]["MTD"]["current"] == Decimal("200")
    assert revenue["periods"]["YTD"]["current"] == Decimal("300")
    assert revenue["periods"]["MTD"]["budget"] == Decimal("250")
    assert revenue["periods"]["YTD"]["budget"] == Decimal("500")
    assert revenue["periods"]["MTD"]["last_year"] == Decimal("150")
    assert result["summary"]["MTD"]["current"]["net_profit_after_tax"] == Decimal("120")


def test_kap_hides_leaf_accounts_without_changing_totals_and_contra_asset_sign(db):
    session, a = db
    parent = ChartOfAccount(account_number="700", account_name="Expenses", account_type="EXPENSE", normal_balance="DEBIT", is_posting=False)
    session.add(parent); session.flush(); a[1].parent_account_id=parent.id
    activity(session, a, date(2026,2,10), "100", "30")
    session.add(OpeningBalance(fiscal_year=2026, account_id=a[4].id, location_id=1, amount=Decimal("10"), balance_type="CREDIT"))
    session.commit()
    report = FinancialReportingService(session).build(report_type="PROFIT_LOSS", as_of_date=date(2026,2,28), kap=True)
    assert any(r["account_code"] == "700" for r in report["rows"])
    assert not any(r["account_code"] == "700-100" for r in report["rows"])
    bs = FinancialStatementService(session).balance_sheet(as_of_date=date(2026,2,28))
    assert bs["summary"]["total_assets"] == Decimal("60")
    session.add(FinancialReportSetting(report_type="PROFIT_LOSS", selected_account_ids=[a[0].id], header_only=False)); session.commit()
    selected = FinancialReportingService(session).build(report_type="PROFIT_LOSS", as_of_date=date(2026,2,28))
    assert len(selected["rows"]) == 1
    assert selected["section_totals"]["EXPENSES"]["MTD"]["current"] == Decimal("30")


def test_revenue_400_with_legacy_children_remains_visible_and_other_revenue_rejected(db):
    session, a = db
    child=ChartOfAccount(account_number="400-100", account_name="Legacy revenue", account_type="REVENUE", normal_balance="CREDIT", parent_account_id=a[0].id)
    session.add(child); session.commit()
    activity(session,a,date(2026,2,10),"100","30")
    report=TrialBalanceService(session).report(date_from=date(2026,1,1),date_to=date(2026,2,28),paginate=False)
    assert any(r["account_code"] == "400" for r in report["items"])
    with pytest.raises(HTTPException) as exc:
        journals.validate(session,[SimpleNamespace(account_id=child.id,debit=0,credit=Decimal("1"))])
    assert "400" in exc.value.detail


def test_comparison_leap_day_and_empty_period(db):
    assert periods(date(2028,2,29))["MTD"]["last_year"][1] == date(2027,2,28)
    session,_=db
    result=PeriodClosingService.close(session,1,USER)
    assert result["closing_journal_ids"] == []
    assert session.query(JournalEntry).count() == 0


def test_report_api_scopes_reads_and_excel_and_rejects_other_location(db):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.api.v1 import financial_reports
    import asyncio
    session,a=db
    activity(session,a,date(2026,2,10),"100","30")
    activity(session,a,date(2026,2,10),"900","300",location=2,unit=2)
    scoped=SimpleNamespace(id=1, default_location_id=1, location_id=1, roles=[])
    # Direct endpoint calls verify the same scope and exported numeric values.
    report=financial_reports.comparison(report_type="PROFIT_LOSS",as_of_date=date(2026,2,28),
        location_id=None,business_unit_id=None,department_id=None,budget_id=None,
        kap=False,include_zero=False,db=session,user=scoped)
    assert report["scope"]["location_id"] == 1
    assert report["summary"]["MTD"]["current"]["total_revenue"] == "100.00"
    with pytest.raises(HTTPException) as denied:
        financial_reports.comparison(report_type="PROFIT_LOSS",as_of_date=date(2026,2,28),
            location_id=2,business_unit_id=None,department_id=None,budget_id=None,
            kap=False,include_zero=False,db=session,user=scoped)
    assert denied.value.status_code==404
    response=financial_reports.export_comparison(report_type="PROFIT_LOSS",as_of_date=date(2026,2,28),
        location_id=None,business_unit_id=None,department_id=None,budget_id=None,
        kap=True,include_zero=False,db=session,user=scoped)
    async def collect():
        return b"".join([chunk async for chunk in response.body_iterator])
    workbook=load_workbook(BytesIO(asyncio.run(collect())))
    assert workbook.sheetnames == ["MTD","YTD"]
    values=list(workbook["MTD"].values)
    assert any(row[0]=="Pendapatan" and row[1]==100 for row in values)
    assert not any(row[2]==900 for row in values)


def test_overlapping_budgets_require_selection_and_missing_budget_is_unavailable(db):
    session,a=db
    activity(session,a,date(2026,2,10),"100","30")
    no_budget=FinancialReportingService(session).build(report_type="PROFIT_LOSS",as_of_date=date(2026,2,28))
    assert no_budget["section_totals"]["REVENUE"]["MTD"]["budget"] is None
    for index in (1,2):
        b=Budget(budget_number=f"OVER-{index}",name="Plan",fiscal_year=2026,date_from=date(2026,1,1),date_to=date(2026,12,31),status="APPROVED")
        b.lines=[BudgetLine(account_id=a[0].id,business_unit_id=1,location_id=1,month=2,budget_amount=Decimal("200"))]
        session.add(b)
    session.commit()
    with pytest.raises(HTTPException) as overlapping:
        FinancialReportingService(session).build(report_type="PROFIT_LOSS",as_of_date=date(2026,2,28))
    assert overlapping.value.status_code==422
    selected=FinancialReportingService(session).build(report_type="PROFIT_LOSS",as_of_date=date(2026,2,28),budget_id=1)
    assert selected["section_totals"]["REVENUE"]["MTD"]["budget"]==Decimal("200")


def test_balance_sheet_mtd_and_ytd_movement_and_point_in_time_budget(db):
    session,a=db
    activity(session,a,date(2026,1,10),"100","30")
    activity(session,a,date(2026,2,10),"200","80")
    session.add(OpeningBalance(fiscal_year=2026,account_id=a[2].id,location_id=1,amount=Decimal("1000"),balance_type="DEBIT"))
    b=Budget(budget_number="BS-BUD",name="BS Plan",fiscal_year=2026,date_from=date(2026,1,1),date_to=date(2026,12,31),status="LOCKED")
    b.lines=[BudgetLine(account_id=a[2].id,business_unit_id=1,location_id=1,month=m,budget_amount=Decimal("1500")) for m in (1,2)]
    session.add(b); session.commit()
    report=FinancialReportingService(session).build(report_type="BALANCE_SHEET",as_of_date=date(2026,2,28),location_id=1)
    cash=next(row for row in report["rows"] if row["account_code"]=="100-100")
    assert cash["periods"]["MTD"]["current"]==Decimal("1190")
    assert cash["periods"]["MTD"]["movement"]==Decimal("120")
    assert cash["periods"]["YTD"]["movement"]==Decimal("190")
    assert cash["periods"]["YTD"]["budget"]==Decimal("1500")


@pytest.mark.parametrize("kind", ["BALANCE_SHEET", "PROFIT_LOSS"])
def test_report_actuals_load_before_optional_setup_migration(db, kind):
    session, accounts = db
    activity(session, accounts, date(2026, 1, 10), "100", "30")
    opening = OpeningBalance(fiscal_year=2026, account_id=accounts[2].id, location_id=1,
                             amount=Decimal("1000"), balance_type="DEBIT")
    session.add(opening); session.commit()
    before = FinancialReportingService(session).build(report_type=kind, as_of_date=date(2026, 10, 6))
    FinancialReportSetting.__table__.drop(session.connection())
    session.commit()
    after = FinancialReportingService(session).build(report_type=kind, as_of_date=date(2026, 10, 6))
    assert after["metadata"]["report_setup_available"] is False
    assert after["section_totals"] == before["section_totals"]
    assert after["rows"] == before["rows"]
    assert session.get(OpeningBalance, opening.id).amount == Decimal("1000")


@pytest.mark.parametrize("kind", ["BALANCE_SHEET", "PROFIT_LOSS"])
def test_comparison_and_export_http_contract_without_setup_table(db, kind):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api.v1 import financial_reports
    from app.core.database import get_db
    session, accounts = db
    activity(session, accounts, date(2026, 10, 1), "100", "30")
    FinancialReportSetting.__table__.drop(session.connection()); session.commit()
    app = FastAPI()
    app.include_router(financial_reports.router)
    app.dependency_overrides[get_db] = lambda: session
    app.dependency_overrides[financial_reports.view] = lambda: USER
    app.dependency_overrides[financial_reports.manage] = lambda: USER
    with TestClient(app) as client:
        params = {"report_type": kind, "as_of_date": "2026-10-06", "kap": "false", "include_zero": "false"}
        response = client.get("/financial-reports/comparison", params=params)
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["periods"]["MTD"]["current"] == ["2026-10-01", "2026-10-06"]
        assert payload["rows"]
        export = client.get("/financial-reports/comparison/export", params=params)
        assert export.status_code == 200, export.text
        assert export.content[:2] == b"PK"
        setup = client.get("/financial-reports/setup")
        assert setup.status_code == 503
        save = client.put("/financial-reports/setup/" + kind, json={"header_only": True})
        assert save.status_code == 503


def test_actuals_do_not_require_optional_budget_tables(db):
    session, accounts = db
    activity(session, accounts, date(2026, 10, 1), "100", "30")
    BudgetLine.__table__.drop(session.connection())
    Budget.__table__.drop(session.connection()); session.commit()
    result = FinancialReportingService(session).build(report_type="BALANCE_SHEET", as_of_date=date(2026, 10, 6))
    assert result["metadata"]["budget_setup_available"] is False
    assert result["budget_available"] == {"MTD": False, "YTD": False}
    assert result["section_totals"]["ASSETS"]["MTD"]["current"] == Decimal("70")
    assert result["section_totals"]["ASSETS"]["MTD"]["budget"] is None
    with pytest.raises(HTTPException) as exc:
        FinancialReportingService(session).build(report_type="BALANCE_SHEET", as_of_date=date(2026, 10, 6), budget_id=1)
    assert exc.value.status_code == 503


@pytest.mark.parametrize("kind", ["BALANCE_SHEET", "PROFIT_LOSS"])
def test_missing_prior_year_is_unavailable_but_current_zero_is_real(db, kind):
    session, accounts = db
    activity(session, accounts, date(2026, 1, 10), "100", "30")
    result = FinancialReportingService(session).build(report_type=kind, as_of_date=date(2026, 10, 6))
    for period in ("MTD", "YTD"):
        assert result["comparison_available"][period]["last_year"] is False
        assert result["summary"][period]["last_year"] is None
        assert all(row["periods"][period]["last_year"] is None for row in result["rows"])
        assert all(section[period]["last_year"] is None for section in result["section_totals"].values())
    if kind == "PROFIT_LOSS":
        assert result["section_totals"]["REVENUE"]["MTD"]["current"] == Decimal("0")
    january = FinancialReportingService(session).build(report_type=kind, as_of_date=date(2026, 1, 31))
    assert january["comparison_available"]["MTD"]["last_month"] is False
    assert all(section["MTD"]["last_month"] is None for section in january["section_totals"].values())


def test_prior_year_evidence_is_scoped_and_can_have_genuine_zero(db):
    session, accounts = db
    activity(session, accounts, date(2026, 2, 1), "100", "30")
    activity(session, accounts, date(2025, 1, 1), "200", "60", location=2, unit=2)
    report = FinancialReportingService(session).build(report_type="PROFIT_LOSS", as_of_date=date(2026, 2, 28), location_id=1)
    assert report["section_totals"]["REVENUE"]["MTD"]["last_year"] is None
    report = FinancialReportingService(session).build(report_type="PROFIT_LOSS", as_of_date=date(2026, 2, 28), location_id=2)
    assert report["comparison_available"]["MTD"]["last_year"] is True
    assert report["section_totals"]["REVENUE"]["MTD"]["last_year"] == Decimal("0")
    assert report["section_totals"]["REVENUE"]["YTD"]["last_year"] == Decimal("200")


def test_header_dimensions_are_used_only_when_line_has_no_override(db):
    session, accounts = db
    header_only = activity(session, accounts, date(2026, 2, 1), "100", "30")
    for line in header_only.lines:
        line.location_id = line.business_unit_id = None
    overridden = activity(session, accounts, date(2026, 2, 2), "200", "60")
    for line in overridden.lines:
        line.location_id = line.business_unit_id = 2
    session.commit()
    for unit, expected in ((1, "100"), (2, "200")):
        report = FinancialReportingService(session).build(report_type="PROFIT_LOSS", as_of_date=date(2026, 2, 28), location_id=unit, business_unit_id=unit)
        assert report["section_totals"]["REVENUE"]["MTD"]["current"] == Decimal(expected)
    assert session.query(JournalEntry).count() == 2


def test_unit_opening_uses_only_explicitly_mapped_locations_and_named_exports(db):
    from app.models.location import Location
    from app.models.business_unit import BusinessUnit
    from app.models.department import Department
    from app.api.v1 import financial_reports
    from openpyxl import load_workbook
    from io import BytesIO
    session, accounts = db
    session.add_all([BusinessUnit(id=1, code="HOTEL", name="Hotel"), BusinessUnit(id=2, code="CATERING", name="Catering"),
                     Location(id=1, code="PKU", name="Pekanbaru", business_unit_id=1), Location(id=2, code="SITE", name="Site", business_unit_id=2),
                     Department(id=1, code="FIN", name="Finance", business_unit_id=1)])
    session.add_all([OpeningBalance(fiscal_year=2026, account_id=accounts[2].id, location_id=loc, amount=Decimal(amount), balance_type="DEBIT") for loc, amount in ((1, "100"), (2, "200"), (None, "300"))])
    session.commit()
    report = FinancialReportingService(session).build(report_type="BALANCE_SHEET", as_of_date=date(2026, 10, 6), business_unit_id=1, location_id=1)
    assert report["section_totals"]["ASSETS"]["YTD"]["current"] == Decimal("100")
    assert report["scope_labels"]["business_unit_id"] == "Hotel"
    assert report["scope_labels"]["location_id"] == "Pekanbaru"
    report = FinancialReportingService(session).build(report_type="BALANCE_SHEET", as_of_date=date(2026, 10, 6), business_unit_id=2)
    assert report["section_totals"]["ASSETS"]["YTD"]["current"] == Decimal("200")
    report = FinancialReportingService(session).build(report_type="BALANCE_SHEET", as_of_date=date(2026, 10, 6), business_unit_id=1, department_id=1)
    assert report["section_totals"]["ASSETS"]["YTD"]["current"] == Decimal("0")
    response = financial_reports.export_comparison(report_type="BALANCE_SHEET", as_of_date=date(2026, 10, 6),
        location_id=1, business_unit_id=1, department_id=None, budget_id=None, kap=False, include_zero=False, db=session, user=USER)
    async def body():
        return b"".join([chunk async for chunk in response.body_iterator])
    import asyncio
    workbook = load_workbook(BytesIO(asyncio.run(body())))
    rows = list(workbook["MTD"].values)
    assert rows[0][3] == "Pekanbaru"
    assert rows[0][5] == "Hotel"
    assets = next(row for row in rows if row[0] == "ASSETS")
    assert assets[5] is None  # unavailable prior-year data stays blank in Excel
    assert session.query(OpeningBalance).count() == 3


def test_year_close_uses_selected_year_not_execution_date(db):
    from calendar import monthrange
    from app.services.year_end_service import YearEndService
    session,a=db
    activity(session,a,date(2026,2,10),"100","30")
    for month in range(3,13):
        session.add(FiscalPeriod(fiscal_year=2026,fiscal_month=month,period_name=f"Month {month}",
            start_date=date(2026,month,1),end_date=date(2026,month,monthrange(2026,month)[1]),
            is_open=True,is_closed=False))
    session.commit()
    result=YearEndService.close_year(db=session,fiscal_year=2026,closing_date=date(2027,1,10),user=USER)
    assert result["closing_date"]==date(2026,12,31)
    assert session.query(FiscalPeriod).filter_by(is_closed=True).count()==12
    assert all(session.get(JournalEntry,jid).posting_date.year==2026 for jid in result["closing_journal_ids"])
    assert len(result["closing_journal_ids"])==1


def test_shared_budget_detail_exposes_only_assigned_location(db):
    from app.api.v1 import budgets
    session,a=db
    b=Budget(budget_number="SHARED",name="Plan",fiscal_year=2026,date_from=date(2026,1,1),date_to=date(2026,12,31),status="APPROVED")
    b.lines=[BudgetLine(account_id=a[0].id,business_unit_id=1,location_id=location,month=2,budget_amount=Decimal(amount)) for location,amount in ((1,"200"),(2,"900"))]
    session.add(b);session.commit()
    scoped=SimpleNamespace(id=1,roles=[],default_location_id=1,location_id=1)
    result=budgets.detail(budget_id=b.id,db=session,user=scoped)
    assert result["total_budget"]=="200.00"
    assert len(result["lines"])==1
    assert result["allowed_actions"]=={}
    with pytest.raises(HTTPException):
        budgets.get_budget(session,b.id,scoped)


def test_lightweight_excel_retains_numeric_values_and_treats_names_as_text():
    from app.services.xlsx_export import build_xlsx
    from openpyxl import load_workbook
    from io import BytesIO
    payload=build_xlsx([("MTD", [["Report"], ["Account", "Name", "Amount"], ["400", "=SUM(A1:A2)", Decimal("7316281.25")]])])
    sheet=load_workbook(BytesIO(payload))["MTD"]
    assert sheet["C3"].value==7316281.25
    assert sheet["C3"].data_type=="n"
    assert sheet["B3"].data_type=="s"
    assert sheet.freeze_panes=="C3"
    assert sheet.auto_filter.ref=="A2:C3"
