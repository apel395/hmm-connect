import copy
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from prepare_resiping_transactions import exact_date, extract_transactions, validate_database, write_xlsx, summary


class TransactionPreflightTests(unittest.TestCase):
    def workbook(self,path):
        w=Workbook();inv=w.active;inv.title='INVOICES';rows=w.create_sheet('ITEM_ENTRY')
        for s in (inv,rows):s.append(['Title']);s.append([]);s.append([])
        inv.append(['Invoice ID','Invoice Date','Supplier','Location','Invoice Number','Invoice Value','Tax'])
        rows.append(['Start Invoice ID','tanggal','Item','Unit','Qty','Unit Price','Diskon','Tax','Line Total'])
        inv.append(['INV-000001',date(2026,1,31),'Supplier','DURI','REAL-1',300,0])
        inv.append(['INV-000002','2-14 JAN 26','Supplier','DURI','REAL-2',100,None])
        inv.append(['INV-000003',date(2025,12,31),'Supplier','DURI','OLD',100,None])
        inv.append(['INV-000004',date(2026,1,31),'Supplier','DURI','MIXED',200,0])
        rows.append(['INV-000001',date(2026,1,2),'Ingredient','KG',1,100,None,None,'=E5*F5'])
        rows.append([None,None,'Ingredient','KG',2,100,None,None,'=E6*F6'])
        rows.append(['INV-000002','2-14 JAN 26','Ingredient','KG',1,100,None,None,'=E7*F7'])
        rows.append(['INV-000003',date(2025,12,31),'Old item','KG','invalid',100,None,None,None])
        rows.append(['INV-000004',date(2025,12,31),'Old mixed','KG',1,100,None,None,None])
        rows.append([None,date(2026,1,3),'Ingredient','KG',1,100,None,None,None])
        w.save(path)

    def report(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'source.xlsx';self.workbook(p);return extract_transactions(p)

    def test_exact_dates_do_not_guess_ranges_or_locale(self):
        for raw,expected in [('2-14 JAN 26',None),('14FEB&1MARET 26',None),('31MAR-14APRIL 26',None),
                             ('1 JAN 26',date(2026,1,1)),('03/01/2026',date(2026,1,3)),
                             (datetime(2026,3,1),date(2026,3,1)),('2026-01-09',date(2026,1,9))]:
            self.assertEqual(exact_date(raw),expected)

    def test_exclusion_arithmetic_date_carry_and_mixed_invoice(self):
        r=self.report();inv={x['key']:x for x in r['invoices']}
        self.assertEqual(inv['INV-000001']['status'],'SOURCE_VALIDATED')
        self.assertEqual(inv['INV-000001']['computed_gross'],Decimal(300))
        self.assertEqual(inv['INV-000001']['lines'][1]['received_date'],date(2026,1,2))
        self.assertEqual(inv['INV-000003']['status'],'EXCLUDED_PRE_2026')
        self.assertNotIn('INVALID_QUANTITY_PRICE_DISCOUNT_OR_TAX',inv['INV-000003']['excluded'][0]['errors'])
        self.assertIn('MIXED_2025_2026_INVOICE_REQUIRES_SPLIT',inv['INV-000004']['errors'])
        self.assertEqual(inv['INV-000004']['net_total'],100)
        self.assertIn('INVOICE_GROSS_TOTAL_MISMATCH',inv['INV-000004']['errors'])
        self.assertEqual(summary(r)['excluded_pre_2026_rows'],2)

    def test_blank_tax_is_unknown_not_non_taxable(self):
        x=self.report()['invoices'][1]
        self.assertIn('INVOICE_TAX_UNCONFIRMED',x['errors'])
        self.assertIn('TAX_EVIDENCE_MISSING',x['lines'][0]['errors'])
        self.assertIsNone(x['computed_gross'])
        self.assertIsNone(x['gross_difference'])
        self.assertNotIn('INVOICE_GROSS_TOTAL_MISMATCH',x['errors'])

    def test_discount_is_recomputed_without_executing_formula(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'source.xlsx';self.workbook(p)
            w=load_workbook(p);s=w['ITEM_ENTRY'];s['E5']=2;s['F5']=50;s['G5']=10;s['H5']=11;s['H6']=0;s['I5']='=1/0'
            w['INVOICES']['F5']=301;w['INVOICES']['G5']=11;w.save(p);w.close()
            x=extract_transactions(p)['invoices'][0]
        self.assertEqual(x['lines'][0]['net'],90)
        self.assertEqual(x['computed_gross'],301)
        self.assertIn('DISCOUNT_POSTING_REVIEW',x['lines'][0]['errors'])

    def catalog(self):
        return {'products':[{'id':1,'code':'P1','name':'Ingredient','item_type':'BUY','is_active':True,'uom':'KG','inventory_treatment':'STOCK','inventory_account_id':167,'tax_category':'NON_TAXABLE'}],
                'suppliers':[{'id':1,'name':'Supplier','code':'S1','is_active':True,'is_pkp':False}],
                'locations':[{'id':1,'name':'DURI','code':'DURI','is_active':True}],
                'departments':[{'id':1,'location_id':1,'name':'Kitchen','is_active':True}],
                'warehouses':[{'id':1,'location_id':1,'name':'Warehouse','is_active':True}],
                'accounts':[{'id':167,'is_active':True,'is_posting':True,'account_type':'ASSET','normal_balance':'DEBIT'}],
                'periods':[{'id':1,'start_date':date(2026,1,1),'end_date':date(2026,1,31),'is_open':True,'is_closed':False,'allow_backdated_entries':True}],
                'invoices':[],'receipts':[]}

    def test_db_mapping_respects_roles_units_periods_and_duplicates(self):
        report=self.report();catalog=self.catalog();original=copy.deepcopy(catalog)
        catalog['products'].append({**catalog['products'][0],'id':2,'item_type':'SELL'})
        catalog['invoices']=[{'id':4,'internal_invoice_number':'SI-1','supplier_invoice_number':'REAL-1','supplier_id':1,'invoice_date':date(2026,1,31),'total_amount':Decimal(300),'status':'POSTED'}]
        catalog['receipts']=[{'id':5,'receipt_number':'GR-1','supplier_id':1,'location_id':1,'received_at':datetime(2026,1,2),'status':'POSTED'}]
        validate_database(report,catalog,date(2026,10,7));x=report['invoices'][0]
        self.assertEqual(x['lines'][0]['product_id'],1)
        self.assertIn('EXISTING_INVOICE_REVIEW',x['errors'])
        self.assertIn('EXISTING_RECEIPT_REVIEW',x['lines'][0]['errors'])
        self.assertIn('WAREHOUSE_SELECTION_REQUIRED',x['lines'][0]['errors'])
        self.assertIn('DEPARTMENT_SELECTION_REQUIRED',x['errors'])
        self.assertFalse(report['seed_ready'])
        # The validation never changes catalog/master values.
        self.assertEqual(catalog['products'][0],original['products'][0])
        r=self.report();catalog=self.catalog();catalog['products'][0]['uom']='PACK'
        catalog['products'][0]['tax_category']='REVIEW_REQUIRED'
        catalog['periods'][0]['allow_backdated_entries']=False
        validate_database(r,catalog,date(2026,10,7));x=r['invoices'][0]
        self.assertIn('UNIT_OR_CONVERSION_REVIEW',x['lines'][0]['errors'])
        self.assertIn('ITEM_TAX_REVIEW_REQUIRED',x['lines'][0]['errors'])
        self.assertTrue(any(e.startswith('BACKDATE_NOT_ALLOWED') for e in x['errors']))

    def test_xlsx_source_text_is_not_executable_and_summary_is_compact(self):
        r=self.report();r['invoices'][0]['supplier']='=HYPERLINK("bad")'
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'review.xlsx';write_xlsx(r,p);w=load_workbook(p,data_only=False)
            self.assertEqual(w['Invoices']['C2'].data_type,'s')
            self.assertEqual(w['Invoices']['C2'].value,'=HYPERLINK("bad")')
            self.assertEqual(w.sheetnames,['Invoices','Receiving','Excluded 2025','Unlinked'])
            w.close()
        self.assertNotIn('invoices',summary(r))
        self.assertFalse(summary(r)['database_checked'])


def test_db_catalog_queries_use_actual_schema_without_writes():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.models import Base
    from prepare_resiping_transactions import db_catalog
    engine=create_engine('sqlite://');Base.metadata.create_all(engine)
    with Session(engine) as db:
        result=db_catalog(db)
        assert set(result)=={'products','suppliers','locations','departments','warehouses','accounts','periods','invoices','receipts'}
        assert all(rows==[] for rows in result.values())
