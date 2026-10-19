"""Editable Excel totals must recalculate from their visible inputs."""
from decimal import Decimal
from io import BytesIO
import re
from openpyxl import load_workbook
from app.services.profit_loss_template import detail_rows, summary_schedule
from app.services.xlsx_export import build_xlsx
from test_profit_loss_exact_template import hotel


def evaluate(sheet, coordinate, stack=()):
    assert coordinate not in stack, f'Circular formula: {stack} -> {coordinate}'
    value = sheet[coordinate].value
    if not isinstance(value,str) or not value.startswith('='):
        return Decimal(str(value or 0))
    expression = value[1:]
    assert re.fullmatch(r'[A-Z0-9+\-]+',expression), expression
    def number(match):
        return str(evaluate(sheet,match[0],(*stack,coordinate)))
    expression = re.sub(r'[A-Z]+[0-9]+',number,expression)
    return sum((Decimal(term) for term in re.findall(r'[+-]?\d+(?:\.\d+)?',expression)),Decimal(0))


def verify_formulas(sheets):
    exported = build_xlsx(sheets)
    formulas = load_workbook(BytesIO(exported))
    cached = load_workbook(BytesIO(exported),data_only=True)
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    actual = evaluate(sheet,cell.coordinate)
                    expected = Decimal(str(cached[sheet.title][cell.coordinate].value))
                    assert abs(actual-expected)<Decimal('0.00001'), (sheet.title,cell.coordinate,cell.value,actual,expected)
    assert formulas.calculation.calcMode == 'auto'
    return formulas,cached


def test_department_totals_follow_visible_inputs_and_preserve_blank_budget():
    source = hotel('Food and Beverage', [('Food Revenue','REVENUE',1000),('Service Charge Revenue','REVENUE',100),('Food Cost','COST_OF_REVENUE',200),('Salary','OPERATING_PAYROLL',80),('Service Charge Distribution','OPERATING_PAYROLL',60),('New expense','OPERATING_EXPENSE',30),('Bank income','NONOPERATING_INCOME',10),('BOD expense','UNDER_EXPENSE',20)])
    for a in source['related_accounts']:
        for period in a['periods'].values(): period['budget'] = None
    rows = detail_rows(source,'FB','2026-01-31')
    wb,cached = verify_formulas([('FB',rows)])
    sheet=wb.active
    net=next(c for row in sheet for c in row if c.column==1 and c.value=='NET RESULT').row
    cost=next(c for row in sheet for c in row if c.column==1 and isinstance(c.value,str) and c.value.strip()=='Food Cost').row
    assert cached.active[f'B{net}'].value==720
    before=evaluate(sheet,f'B{net}')
    sheet[f'B{cost}']=250
    assert evaluate(sheet,f'B{net}')==before-50
    budget=rows.headers.index('MTD Budget')+1
    assert sheet.cell(net,budget).value is None


def test_summary_gop_nop_formulas_recalculate_without_double_counting():
    room=hotel('Front Office',[('Room Sales','REVENUE',1000),('Salary','OPERATING_PAYROLL',80),('Expense','OPERATING_EXPENSE',20)])
    admin=hotel('Admin and General',[('Admin salary','OVERHEAD_PAYROLL',50),('Other deduction','UNDER_EXPENSE',10),('Interest','NONOPERATING_INCOME',5)])
    root=hotel('',[]);root['related_accounts']=room['related_accounts']+admin['related_accounts']
    rows=summary_schedule(root,'Hotel',[room,admin],'2026-01-31')
    wb,_=verify_formulas([('Summary',rows)])
    sheet=wb.active
    gop=next(c.row for row in sheet for c in row if c.column==1 and c.value=='GROSS OPERATING PROFIT')
    nop=next(c.row for row in sheet for c in row if c.column==1 and c.value=='NETT OPERATING PROFIT')
    assert evaluate(sheet,f'B{gop}')==850
    assert evaluate(sheet,f'B{nop}')==845
    revenue=next(c.row for row in sheet for c in row if c.column==1 and isinstance(c.value,str) and c.value.strip() in ('Room','Rooms') and sheet.cell(c.row,2).value==1000)
    sheet[f'B{revenue}']=1200
    assert evaluate(sheet,f'B{gop}')==1050
    assert evaluate(sheet,f'B{nop}')==1045


def test_utilities_total_uses_supplied_revenue_base():
    from test_hotel_pl_template_departments import amounts
    source=hotel('Utilities',[('Electricity','OPERATING_EXPENSE',100)])
    source['template_revenue']=amounts(1000)
    verify_formulas([('Utilities',detail_rows(source,'Utilities','2026-01-31'))])


def test_summary_without_department_breakdown_keeps_all_posted_costs():
    source=hotel('',[('Room Sales','REVENUE',1000),('Salary','OPERATING_PAYROLL',80),('Expense','OVERHEAD_EXPENSE',20)])
    rows=summary_schedule(source,'Hotel',[],'2026-01-31')
    verify_formulas([('Summary',rows)])
    assert next(r[1] for r in rows if r and r[0]=='GROSS OPERATING PROFIT')==900
