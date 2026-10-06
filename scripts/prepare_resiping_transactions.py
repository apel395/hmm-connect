#!/usr/bin/env python3
"""Read-only receiving/invoice preflight from RESIPING. Never seeds or posts.

Recompute literal arithmetic; do not execute macros or cached Excel formulas.
Exclude pre-2026 receiving, retain unresolved evidence in JSON/XLSX reports.
"""
import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
import json
from pathlib import Path
import re
import sys
import tempfile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from prepare_resiping_item_seed import download_public_drive, invoice_id, normalized, source_year, MAX_SOURCE_BYTES
CENT=Decimal('0.01')
# User-confirmed operational location aliases for this historical import.
# Keep the source label on every invoice; resolve to an existing canonical site.
LOCATION_ALIASES={
    'RUMBAI':'RUMIN','MINAS':'RUMIN','HULU CAFE':'RUMIN','HULU CAFÉ':'RUMIN',
    'LIBO':'LIPET','PETAPAHAN':'LIPET',
    'ZAMRUD':'ZAMPED','PEDADA':'ZAMPED',
}
AGRO_BOGA_SOURCE_NAMES={
    'ARGOBOGAUTAMA', 'AGROBOGAUTAMA',
    'PTAGROBOGAUTAMA', 'PTARGOBOGAUTAMA',
}


def number(value, blank=None):
    if value is None or value=='':return blank
    if isinstance(value,bool) or str(value).startswith('='):raise ValueError('Expected literal number')
    try:v=Decimal(str(value))
    except InvalidOperation as exc:raise ValueError('Invalid number') from exc
    if not v.is_finite():raise ValueError('Non-finite number')
    return v


def money(value):return value.quantize(CENT,rounding=ROUND_HALF_UP)


def exact_date(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    text=str(value or '').strip()
    formats=('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y','%d/%m/%y','%d-%m-%y')
    for fmt in formats:
        try:return datetime.strptime(text,fmt).date()
        except ValueError:pass
    # An explicit year is required. For ranges/lists use their latest date,
    # as authorized for this historical import; never change Excel date values.
    months={'JAN':1,'JANUARI':1,'JANUARY':1,'FEB':2,'FEBRUARI':2,'FEBRUARY':2,
            'MAR':3,'MARET':3,'MARCH':3,'APR':4,'APRI':4,'APRIL':4,'MEI':5,'MAY':5,
            'JUN':6,'JUNI':6,'JUNE':6,'JUL':7,'JULI':7,'JULY':7,
            'AGU':8,'AGUSTUS':8,'AUG':8,'AUGUST':8,'SEP':9,'SEPTEMBER':9,
            'OKT':10,'OKTOBER':10,'OCT':10,'OCTOBER':10,'NOV':11,'NOVEMBER':11,
            'DES':12,'DESEMBER':12,'DEC':12,'DECEMBER':12}
    named=re.fullmatch(r'(.+?)\s+(20\d{2}|25|26)',text.upper())
    if named:
        year=int(named[2]);year=year+2000 if year<100 else year
        parts=re.split(r'\s*(?:[-–—,&]|\bDAN\b|\bTO\b)\s*',named[1])
        parsed=[];month=None
        # Shared month/year apply to earlier days: "2,4 APRIL 26".
        for part in reversed(parts):
            match=re.fullmatch(r'(\d{1,2})\s*([A-Z]+)?',part.strip())
            if not match:return None
            if match[2]:
                month=months.get(match[2])
                if month is None:return None
            if month is None:return None
            try:parsed.append(date(year,month,int(match[1])))
            except ValueError:return None
        if parsed:return max(parsed)
    return None


def text_number(value):
    if isinstance(value,(int,float,Decimal)) and not isinstance(value,bool):
        d=number(value)
        return str(int(d)) if d==d.to_integral_value() else str(d)
    return str(value or '').strip()


def extract_transactions(path):
    from openpyxl import load_workbook
    if path.stat().st_size>MAX_SOURCE_BYTES:raise ValueError('Workbook exceeds 30 MiB')
    w=load_workbook(path,read_only=True,data_only=False)
    try:
        if not {'INVOICES','ITEM_ENTRY'}.issubset(w.sheetnames):raise ValueError('Expected INVOICES and ITEM_ENTRY')
        ih=list(next(w['INVOICES'].iter_rows(min_row=4,max_row=4,max_col=12,values_only=True)))
        rh=list(next(w['ITEM_ENTRY'].iter_rows(min_row=4,max_row=4,max_col=14,values_only=True)))
        if ih[:7]!=['Invoice ID','Invoice Date','Supplier','Location','Invoice Number','Invoice Value','Tax']:
            raise ValueError('Invoice columns changed')
        if rh[:9]!=['Start Invoice ID','tanggal','Item','Unit','Qty','Unit Price','Diskon','Tax','Line Total']:
            raise ValueError('Receiving columns changed; review discount/tax layout')
        invoices={}
        for row,c in enumerate(w['INVOICES'].iter_rows(min_row=5,max_col=12,values_only=True),5):
            if not any(c[1:7]):continue
            key=invoice_id(c[0],row)
            if not key or key in invoices:raise ValueError(f'Invoice row {row}: missing/duplicate/unsupported ID')
            errors=[]
            if not c[2] or str(c[2]).startswith('='):errors.append('INVALID_SUPPLIER_LITERAL')
            if not c[3] or str(c[3]).startswith('='):errors.append('INVALID_LOCATION_LITERAL')
            if c[4] is not None and str(c[4]).startswith('='):errors.append('INVALID_INVOICE_NUMBER_LITERAL')
            try:value=number(c[5]);tax=number(c[6])
            except ValueError:value=tax=None;errors.append('INVALID_INVOICE_AMOUNT_OR_TAX')
            invoices[key]={'key':key,'row':row,'supplier':str(c[2] or '').strip(),'location':str(c[3] or '').strip(),
                'invoice_number':text_number(c[4]),'raw_invoice_date':c[1],'invoice_date':exact_date(c[1]),
                'date_policy':'LATEST_DATE_IN_RANGE_OR_LIST',
                'declared_total':value,'declared_tax':tax,'errors':errors,'lines':[], 'excluded':[]}
        sequence=bool(invoices) and all(x['invoice_number']==str(x['row']-4) for x in invoices.values())
        orphan=[];current=None;carried=None
        for row,c in enumerate(w['ITEM_ENTRY'].iter_rows(min_row=5,max_col=14,values_only=True),5):
            if not any(c[:8]):current=None;carried=None;continue
            if c[0] is not None:current=invoice_id(c[0],row,first_item=True);carried=None
            if current not in invoices:
                orphan.append({'row':row,'item':str(c[2] or ''),'errors':['MISSING_OR_UNSUPPORTED_INVOICE_LINK']});continue
            header=invoices[current]
            if c[1] is not None:carried=c[1]
            raw=carried if carried is not None else header['raw_invoice_date']
            year=source_year(raw)
            line={'row':row,'invoice_key':current,'item':normalized(c[2]),'unit':normalized(c[3]),
                  'raw_received_date':raw,'received_date':exact_date(raw),
                  'date_policy':'LATEST_DATE_IN_RANGE_OR_LIST',
                  'date_source':'item/carried' if carried is not None else 'invoice header',
                  'raw_quantity':c[4],'raw_unit_price':c[5],'raw_discount':c[6],'raw_tax':c[7],'errors':[]}
            if year is not None and year<2026:
                line['exclusion']='PRE_2026_RECEIVING';header['excluded'].append(line);continue
            if year!=2026:line['errors'].append('RECEIVING_YEAR_UNRESOLVED_OR_OUTSIDE_2026')
            if line['received_date'] is None:line['errors'].append('RECEIVING_DATE_NOT_SINGLE_DATE')
            if not line['item'] or line['item'].startswith('='):line['errors'].append('INVALID_ITEM_NAME')
            if not line['unit'] or line['unit'].startswith('='):line['errors'].append('INVALID_UNIT')
            try:
                qty=number(c[4]);price=number(c[5]);discount=number(c[6],Decimal(0));tax=number(c[7])
                if qty is None or price is None or qty<=0 or price<0 or discount<0 or (tax is not None and tax<0):raise ValueError('Invalid amount')
                base=money(qty*price)
                if discount>base:raise ValueError('Discount exceeds subtotal')
                line.update(quantity=qty,unit_price=price,discount=discount,net=money(base-discount),tax=tax)
                if qty!=money(qty) or price!=money(price):line['errors'].append('RECEIVING_MODEL_PRECISION_REVIEW')
                if discount>0:line['errors'].append('DISCOUNT_POSTING_REVIEW')
                literal_total=number(c[8]) if c[8] is not None and not str(c[8]).startswith('=') else None
                line['literal_line_total']=literal_total
                if tax is None and header['declared_tax']==0:line['tax']=Decimal(0);line['tax_evidence']='Explicit zero invoice-header tax'
                elif tax is not None:line['tax_evidence']='Explicit ITEM_ENTRY tax amount'
                else:line['errors'].append('TAX_EVIDENCE_MISSING')
                line['gross']=money(line['net']+line['tax']) if line['tax'] is not None else None
                if literal_total is not None and line['gross'] is not None and literal_total!=line['gross']:
                    line['errors'].append('LITERAL_LINE_TOTAL_MISMATCH')
            except ValueError:line['errors'].append('INVALID_QUANTITY_PRICE_DISCOUNT_OR_TAX')
            header['lines'].append(line)
        for header in invoices.values():
            if not header['lines']:
                header['status']='EXCLUDED_PRE_2026' if header['excluded'] else 'EMPTY_OR_UNLINKED'
                continue
            if header['excluded']:header['errors'].append('MIXED_2025_2026_INVOICE_REQUIRES_SPLIT')
            if not header['invoice_date']:header['errors'].append('INVOICE_DATE_NOT_SINGLE_DATE')
            elif header['invoice_date'].year!=2026:header['errors'].append('INVOICE_DATE_OUTSIDE_2026')
            if not header['invoice_number']:header['errors'].append('MISSING_SUPPLIER_INVOICE_NUMBER')
            elif sequence:header['errors'].append('INVOICE_NUMBER_IS_TEMPLATE_SEQUENCE_REVIEW')
            if header['declared_total'] is None or header['declared_total']<=0:header['errors'].append('INVALID_INVOICE_TOTAL')
            if header['declared_tax'] is None:header['errors'].append('INVOICE_TAX_UNCONFIRMED')
            header['net_total']=sum((x.get('net',Decimal(0)) for x in header['lines']),Decimal(0))
            all_tax=all(x.get('gross') is not None for x in header['lines'])
            header['computed_gross']=sum((x['gross'] for x in header['lines']),Decimal(0)) if all_tax else None
            header['gap_against_net']=header['declared_total']-header['net_total'] if header['declared_total'] is not None else None
            header['gross_difference']=header['computed_gross']-header['declared_total'] if all_tax and header['declared_total'] is not None else None
            if header['gross_difference'] is not None and abs(header['gross_difference'])>=CENT:header['errors'].append('INVOICE_GROSS_TOTAL_MISMATCH')
            if all_tax and header['declared_tax'] is not None and sum((x['tax'] for x in header['lines']),Decimal(0))!=header['declared_tax']:
                header['errors'].append('INVOICE_TAX_TOTAL_MISMATCH')
            header['status']='REVIEW' if header['errors'] or any(x['errors'] for x in header['lines']) else 'SOURCE_VALIDATED'
        return {'mode':'READ_ONLY_PREFLIGHT','source_sha256':sha256(path.read_bytes()).hexdigest(),'source_file':path.name,
                'cutoff':'2026-01-01','database_checked':False,'seed_ready':False,
                'invoices':list(invoices.values()),'unlinked_rows':orphan}
    finally:w.close()


def exact_match(rows,label):
    candidates=[x for x in rows if normalized(x.get('name'))==normalized(label) or (x.get('code') and normalized(x['code'])==normalized(label))]
    return candidates[0] if len(candidates)==1 and candidates[0].get('is_active',True) else None


def validate_database(report,catalog,today=None):
    today=today or date.today()
    for inv in report['invoices']:
        if inv['status']=='EXCLUDED_PRE_2026':continue
        supplier_key=re.sub(r'[^A-Z0-9]','',normalized(inv['supplier']))
        if supplier_key in AGRO_BOGA_SOURCE_NAMES:
            # User selected DB supplier 225. Do not fall back to duplicate 242.
            matches=[s for s in catalog['suppliers'] if s['id']==225 and s.get('is_active',True)
                     and re.sub(r'[^A-Z0-9]','',normalized(s['name']))=='PTAGROBOGAUTAMA']
            supplier=matches[0] if len(matches)==1 else None
            inv['supplier_mapping_rule']='USER_CONFIRMED_SUPPLIER_225'
        else:
            supplier=exact_match(catalog['suppliers'],inv['supplier'])
            inv['supplier_mapping_rule']='EXACT_SOURCE_LABEL'
        inv['mapped_supplier_name']=supplier['name'] if supplier else None
        location_label=LOCATION_ALIASES.get(normalized(inv['location']),inv['location'])
        location=exact_match(catalog['locations'],location_label)
        inv['location_mapping_label']=location_label
        inv['location_mapping_rule']='USER_CONFIRMED_ALIAS' if normalized(inv['location']) in LOCATION_ALIASES else 'EXACT_SOURCE_LABEL'
        inv['mapped_location_name']=location['name'] if location else None
        inv['supplier_id']=supplier['id'] if supplier else None;inv['location_id']=location['id'] if location else None
        if not supplier:inv['errors'].append('SUPPLIER_NOT_UNIQUELY_MATCHED')
        if not location:inv['errors'].append('LOCATION_NOT_UNIQUELY_MATCHED')
        if supplier:
            if not supplier.get('is_pkp') and (inv.get('declared_tax') or Decimal(0))>0:
                inv['errors'].append('PPN_REQUIRES_PKP_SUPPLIER_REVIEW')
            possible=[x for x in catalog['invoices'] if x['supplier_id']==supplier['id'] and (str(x.get('supplier_invoice_number') or '')==inv['invoice_number'] or (x['invoice_date']==inv['invoice_date'] and x['total_amount']==inv['declared_total']))]
            inv['existing_invoice_candidates']=[{'id':x['id'],'number':x['internal_invoice_number'],'status':x['status']} for x in possible]
            if possible:inv['errors'].append('EXISTING_INVOICE_REVIEW')
        if location:
            inv['department_candidates']=[x for x in catalog['departments'] if x['location_id']==location['id'] and x['is_active']]
            inv['warehouse_candidates']=[x for x in catalog['warehouses'] if x['location_id']==location['id'] and x['is_active']]
            operational=[x for x in inv['department_candidates'] if normalized(x['name'])=='OPERASIONAL']
            is_head_office=normalized(location['name']) in ('HEAD OFFICE','HEAD OFFICE LOCATION','KANTOR PUSAT')
            if len(operational)==1 and not is_head_office:
                inv['department_id']=operational[0]['id']
                inv['department_mapping_rule']='USER_CONFIRMED_OPERATIONAL'
            else:inv['errors'].append('DEPARTMENT_SELECTION_REQUIRED')
        dates=[inv['invoice_date']]
        for line in inv['lines']:
            products=[p for p in catalog['products'] if p['item_type']=='BUY' and normalized(p['name'])==line['item']]
            p=products[0] if len(products)==1 and products[0]['is_active'] else None
            if not p:line['errors'].append('BUY_ITEM_NOT_UNIQUELY_MATCHED');continue
            line['product_id']=p['id'];line['product_code']=p['code'];line['treatment']=p['inventory_treatment']
            if normalized(p['uom'])!=line['unit']:line['errors'].append('UNIT_OR_CONVERSION_REVIEW')
            if p['inventory_treatment'] not in ('STOCK','DIRECT_USE','FIXED_ASSET'):line['errors'].append('ITEM_TREATMENT_UNRESOLVED')
            field={'STOCK':'inventory_account_id','DIRECT_USE':'purchase_account_id','FIXED_ASSET':'asset_account_id'}.get(p['inventory_treatment'])
            account=next((a for a in catalog['accounts'] if a['id']==p.get(field)),None) if field else None
            kind='EXPENSE' if p['inventory_treatment']=='DIRECT_USE' else 'ASSET'
            if not account or not account['is_active'] or not account['is_posting'] or account['account_type']!=kind or account['normal_balance']!='DEBIT':line['errors'].append('ITEM_POSTING_ACCOUNT_INVALID')
            if p.get('tax_category')=='REVIEW_REQUIRED':line['errors'].append('ITEM_TAX_REVIEW_REQUIRED')
            line['department_id']=inv.get('department_id')
            line['warehouse_id']=None
            if p['inventory_treatment']=='STOCK':
                warehouses=inv.get('warehouse_candidates',[])
                if len(warehouses)==1:
                    line['warehouse_id']=warehouses[0]['id']
                    line['warehouse_mapping_rule']='UNIQUE_ACTIVE_LOCATION_WAREHOUSE'
                else:line['errors'].append('WAREHOUSE_SELECTION_REQUIRED')
            dates.append(line['received_date'])
            if supplier and line['received_date']:
                possible=[x for x in catalog['receipts'] if x['supplier_id']==supplier['id'] and x['received_at'].date()==line['received_date'] and (not location or x['location_id']==location['id'])]
                line['existing_receipt_candidates']=[{'id':x['id'],'number':x['receipt_number'],'status':x['status']} for x in possible]
                if possible:line['errors'].append('EXISTING_RECEIPT_REVIEW')
        for day in {d for d in dates if d is not None and d.year==2026}:
            periods=[p for p in catalog['periods'] if p['start_date']<=day<=p['end_date']]
            if len(periods)!=1 or not periods[0]['is_open'] or periods[0]['is_closed']:
                inv['errors'].append(f'ACCOUNTING_PERIOD_NOT_OPEN:{day}')
            elif day<today and not periods[0]['allow_backdated_entries']:
                inv['errors'].append(f'BACKDATE_NOT_ALLOWED:{day}')
        inv['status']='REVIEW' if inv['errors'] or any(x['errors'] for x in inv['lines']) else 'DB_VALIDATED'
    report['database_checked']=True
    report['seed_ready']=False  # Mapping choices require explicit source/DB verification; this tool cannot seed.
    return report


def db_catalog(db):
    from sqlalchemy import text
    queries={'products':'SELECT id,code,name,item_type,is_active,uom,usage_uom,conversion,inventory_treatment,tax_category,inventory_account_id,purchase_account_id,asset_account_id FROM products',
      'suppliers':'SELECT id,code,name,is_active,is_pkp FROM suppliers',
      'locations':'SELECT id,code,name,is_active FROM locations',
      'departments':'SELECT d.id,d.code,d.name,d.is_active,ld.location_id FROM departments d JOIN location_departments ld ON ld.department_id=d.id',
      'warehouses':'SELECT id,code,name,location_id,is_active FROM warehouses',
      'accounts':'SELECT id,account_number,account_type,normal_balance,is_active,is_posting FROM chart_of_accounts',
      'periods':'SELECT id,start_date,end_date,is_open,is_closed,allow_backdated_entries FROM fiscal_periods WHERE fiscal_year=2026',
      'invoices':'SELECT id,internal_invoice_number,supplier_invoice_number,supplier_id,invoice_date,total_amount,status FROM supplier_invoices',
      'receipts':'SELECT id,receipt_number,supplier_id,location_id,received_at,status FROM goods_receipts'}
    return {k:[dict(r) for r in db.execute(text(q)).mappings()] for k,q in queries.items()}


def summary(report):
    active=[x for x in report['invoices'] if x['status']!='EXCLUDED_PRE_2026']
    errors=Counter(e for x in active for e in x['errors'])
    errors.update(e for x in active for line in x['lines'] for e in line['errors'])
    return {k:report[k] for k in ('mode','source_sha256','database_checked','seed_ready')} | {
        'invoice_headers':len(report['invoices']),'invoice_statuses':dict(Counter(x['status'] for x in report['invoices'])),
        'receiving_rows':sum(len(x['lines']) for x in report['invoices']),
        'excluded_pre_2026_rows':sum(len(x['excluded']) for x in report['invoices']),
        'unlinked_rows':len(report['unlinked_rows']),'review_counts':dict(errors)}


def write_xlsx(report,path):
    from openpyxl import Workbook
    w=Workbook(write_only=True)
    def sheet(name,columns,rows):
        ws=w.create_sheet(name);ws.append(columns)
        for record in rows:
            values=[]
            for key in columns:
                value=record.get(key)
                if isinstance(value,(list,dict)):value=json.dumps(value,default=str,ensure_ascii=False)
                if isinstance(value,Decimal):value=str(value)
                if isinstance(value,datetime):value=value.isoformat()
                # Force text for source/formula-like strings; no executable XLSX formula cells.
                from openpyxl.cell import WriteOnlyCell
                cell=WriteOnlyCell(ws,value=value)
                if isinstance(value,str):cell.data_type='s'
                values.append(cell)
            ws.append(values)
    sheet('Invoices',['key','row','supplier','location','invoice_number','raw_invoice_date','invoice_date','declared_total','declared_tax','net_total','computed_gross','gap_against_net','gross_difference','status','errors','supplier_id','mapped_supplier_name','supplier_mapping_rule','location_id','location_mapping_label','mapped_location_name','location_mapping_rule','department_id','department_mapping_rule','existing_invoice_candidates','department_candidates','warehouse_candidates'],report['invoices'])
    sheet('Receiving',['invoice_key','row','item','unit','raw_received_date','received_date','quantity','unit_price','discount','net','tax','gross','product_id','product_code','treatment','department_id','warehouse_id','warehouse_mapping_rule','errors','existing_receipt_candidates'],[r for x in report['invoices'] for r in x['lines']])
    sheet('Excluded 2025',['invoice_key','row','item','unit','raw_received_date','exclusion'],[r for x in report['invoices'] for r in x['excluded']])
    sheet('Unlinked',['row','item','errors'],report['unlinked_rows'])
    w.save(path)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    src=p.add_mutually_exclusive_group(required=True);src.add_argument('--file',type=Path);src.add_argument('--drive-id')
    p.add_argument('--output',type=Path,required=True);p.add_argument('--xlsx',type=Path);p.add_argument('--db-check',action='store_true')
    args=p.parse_args()
    if args.drive_id:
        with tempfile.TemporaryDirectory(prefix='hmm-transaction-preflight-') as tmp:
            path=Path(tmp)/'source.xlsm';path.write_bytes(download_public_drive(args.drive_id));report=extract_transactions(path)
        report['source_drive_id']=args.drive_id
    else:report=extract_transactions(args.file)
    if args.db_check:
        from sqlalchemy import text
        from app.core.database import SessionLocal
        with SessionLocal() as db:
            db.execute(text('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY'))
            db.execute(text("SET LOCAL statement_timeout='60s'"))
            try:validate_database(report,db_catalog(db))
            finally:db.rollback()
    args.output.write_text(json.dumps(report,indent=2,default=str,ensure_ascii=False)+'\n')
    if args.xlsx:write_xlsx(report,args.xlsx)
    print(json.dumps(summary(report),indent=2))
    print(f'Full JSON: {args.output}')
    if args.xlsx:print(f'Review workbook: {args.xlsx}')
    print('READ ONLY: no master, receiving, invoice, stock, journal, tax or opening-balance writes.')

if __name__=='__main__':main()
