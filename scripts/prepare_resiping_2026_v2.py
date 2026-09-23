#!/usr/bin/env python3
"""Read-only preparation report for RESIPING2026(2).xlsx historical seeding.

Business rules encoded from source review:
- TOTAL INVOICE (column M) starts one supplier invoice.
- Later supplier/date blocks with no TOTAL INVOICE continue the current invoice
  when supplier + mapped location are the same.
- Rows without a positive quantity AND positive unit price are not received.
- RUMBAI / MINAS / RUMBAI MINAS / RUMBAI MINUS -> Rumin (005).
- SANGASANGA / NKL -> Kalimantan (007).
- OPERATIONAL invoices are skipped for now.
- No product treatment is guessed. Existing REVIEW_REQUIRED items and new items
  are reported for Item Master review before transactional posting.

This script writes nothing to the database.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from sqlalchemy import func

from app.core.database import SessionLocal
from app.models.location import Location
from app.models.product import Product
from app.models.supplier import Supplier


ZERO = Decimal("0")

LOCATION_ALIASES = {
    "MINAS": ("005", "Rumin"),
    "RUMBAI": ("005", "Rumin"),
    "RUMBAI MINAS": ("005", "Rumin"),
    "RUMBAI MINUS": ("005", "Rumin"),
    "HEAD OFFICE": ("002", "Pekanbaru"),
    "BSP PEDADA": ("003", "Zamped"),
    "PEDADA": ("003", "Zamped"),
    "DURI": ("006", "Duri"),
    "NKL": ("007", "Kalimantan"),
    "SANGASANGA": ("007", "Kalimantan"),
}

SUPPLIER_ALIASES = {
    "IBU ERMY": "IBU ERNY",
    "PEMPEKL WENNY": "PEMPEK WENNY",
}

PRODUCT_ALIASES = {
    "BAKSO IKAN 1 KG": "FISH BALL",
    "BEEF SAUSAGE 1KG": "BEEF SAUSAGE",
    "CHEESE MOZARELLA {250GR]": "MOZZARELLA CHEESE",
    "CHICKEN SAUSAGE 1 KG": "CHICKEN SAUSAGE",
    "CORN KERNEL 1 KG": "CORN KERNEL",
    "FRENCH FRIES CRINKLE CUT 1 KG": "FRENCH FRIES CRINKLE CUT",
    "FRENCH FRIES SHOESTRING 1 GR": "FRENCH FRIES SHOESTRING",
    "ICE CREAM VANILLA 8 L": "VANILLA ICE CREAM",
    "ICE CREAM REG. CUP VANILLA 75 ML @24": "VANILLA ICE CREAM CUP",
    "ICE CREAM REG. CUP VANILLA 75 ML @25": "VANILLA ICE CREAM CUP",
    "ICE CREAM REG. CUP VANILLA 75 ML @26": "VANILLA ICE CREAM CUP",
    "MIX VEGETABLE FROZEN 1 KG": "MIXED VEGETABLES FROZEN",
    "PUFF PASTRY SHEET 750 GR": "PUFF PASTRY SHEET",
    "BAKSO SAPI HALUS": "BEEF BALL",
    "BUTTER UNSALTED 200 GR": "UNSALTED BUTTER",
    "GREEN PEAS 1 KG": "GREEN PEAS",
    "MAYONNAISE 4L": "MAYONNAISE",
    "SUSU FULLCREAM 1 LITER @12": "FULL CREAM MILK",
    "BEEF MINCED": "MINCED BEEF",
    "ICE CREAM 3 RASA 8 L": "ICE CREAM 3 FLAVORS",
    "BEEF BALL 1 KG": "BEEF BALL",
}


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def upper(value) -> str:
    return clean(value).upper()


def supplier_key(value) -> str:
    name = upper(value)
    name = SUPPLIER_ALIASES.get(name, name)
    name = re.sub(r"^(PT\.?|CV\.?|UD\.?|TOKO)\s*", "", name)
    return re.sub(r"[^A-Z0-9]+", "", name)


def canonical_supplier_name(value) -> str:
    name = upper(value)
    return SUPPLIER_ALIASES.get(name, name)


def source_date(*values) -> date | None:
    for value in values:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
    return None


def period_start(value) -> date | None:
    d = source_date(value)
    if d:
        return d
    raw = upper(value)
    if not raw:
        return None
    months = {
        "JAN": 1, "JANUARY": 1, "JANUARI": 1,
        "FEB": 2, "FEBRUARY": 2, "FEBRUARI": 2,
        "MAR": 3, "MARCH": 3, "MARET": 3,
        "APR": 4, "APRIL": 4,
        "MAY": 5, "MEI": 5,
        "JUN": 6, "JUNE": 6, "JUNI": 6,
        "JUL": 7, "JULY": 7, "JULI": 7,
        "AUG": 8, "AUGUST": 8, "AGUSTUS": 8,
        "SEP": 9, "SEPT": 9, "SEPTEMBER": 9,
        "OCT": 10, "OCTOBER": 10, "OKTOBER": 10,
        "NOV": 11, "NOVEMBER": 11,
        "DEC": 12, "DESEMBER": 12, "DECEMBER": 12, "DES": 12,
    }
    m = re.search(r"(\d{1,2})(?:\s*-\s*\d{1,2})?\s*[- ]?([A-Z]+)\s*(20)?(\d{2})", raw)
    if not m:
        return None
    month = months.get(m.group(2))
    if not month:
        return None
    year = int(m.group(4)) + 2000
    try:
        return date(year, month, int(m.group(1)))
    except ValueError:
        return None


def decimal_number(value) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    raw = clean(value)
    if not raw:
        return None
    # Quantity source can use Indonesian decimal comma, e.g. 14,36.
    if raw.count(",") == 1 and "." not in raw:
        left, right = raw.split(",")
        if len(right) <= 4 and len(left) <= 4:
            raw = left + "." + right
        else:
            raw = raw.replace(",", "")
    else:
        raw = raw.replace(",", "")
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def rupiah(value) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, Decimal)):
        return Decimal(value)
    if isinstance(value, float):
        # Tax cells in this workbook use 333.896 to mean 333,896.
        if 0 < abs(value) < 1000:
            return Decimal(str(value)) * 1000
        return Decimal(str(value))
    raw = clean(value)
    if not raw:
        return None
    # Remove a trailing decimal ",00" / ".00" after a thousands-formatted value.
    raw = re.sub(r"([,.])00$", "", raw)
    digits = re.sub(r"[^0-9-]", "", raw)
    if not digits or digits == "-":
        return None
    try:
        return Decimal(digits)
    except InvalidOperation:
        return None


def tax_rupiah(value) -> Decimal:
    parsed = rupiah(value)
    if parsed is None:
        return ZERO
    # Integer 262 in this source means 262,000; 2500 is really 2,500.
    if isinstance(value, int) and 0 < abs(value) < 1000:
        return parsed * 1000
    return parsed


def location_key(raw: str) -> str:
    return upper(raw)


def mapped_location(raw: str) -> tuple[str | None, str | None, str]:
    key = location_key(raw)
    if key == "OPERATIONAL":
        return None, None, "SKIP_OPERATIONAL"
    if key in LOCATION_ALIASES:
        code, name = LOCATION_ALIASES[key]
        return code, name, "MAPPED"
    if not key:
        return None, None, "MISSING"
    return None, key, "LOOKUP"


@dataclass
class Item:
    row: int
    received_date: date | None
    name: str
    raw_name: str
    brand_or_group: str
    uom: str
    qty: Decimal | None
    price: Decimal | None
    tax: Decimal = ZERO

    @property
    def received(self) -> bool:
        return bool(self.qty is not None and self.qty > ZERO and self.price is not None and self.price > ZERO)

    @property
    def amount(self) -> Decimal:
        return (self.qty or ZERO) * (self.price or ZERO)


@dataclass
class Block:
    row: int
    supplier_raw: str
    supplier_name: str
    supplier_key: str
    location_raw: str
    location_code: str | None
    location_name: str | None
    location_state: str
    header_date: date | None
    invoice_date: date | None
    total_invoice: Decimal | None
    items: list[Item] = field(default_factory=list)


@dataclass
class InvoiceGroup:
    start_row: int
    supplier_name: str
    supplier_key: str
    location_code: str | None
    location_name: str | None
    location_state: str
    total_invoice: Decimal | None
    blocks: list[Block] = field(default_factory=list)

    @property
    def received_items(self) -> list[Item]:
        return [item for block in self.blocks for item in block.items if item.received]

    @property
    def not_received_items(self) -> list[Item]:
        return [item for block in self.blocks for item in block.items if not item.received]

    @property
    def receiving_subtotal(self) -> Decimal:
        return sum((item.amount for item in self.received_items), ZERO)

    @property
    def source_tax(self) -> Decimal:
        return sum((item.tax for block in self.blocks for item in block.items), ZERO)


def parse_blocks(ws) -> list[Block]:
    blocks: list[Block] = []
    current: Block | None = None
    carry_date: date | None = None

    for row_no in range(2, ws.max_row + 1):
        values = [ws.cell(row_no, col).value for col in range(1, 14)]
        a, b, c, d, e, f, g, h, i, _j, tax, _total, total_invoice = values

        supplier_header = isinstance(b, str) and clean(b) != ""
        if supplier_header:
            code, loc_name, loc_state = mapped_location(clean(c))
            hdr_date = period_start(a)
            current = Block(
                row=row_no,
                supplier_raw=clean(b),
                supplier_name=canonical_supplier_name(b),
                supplier_key=supplier_key(b),
                location_raw=clean(c),
                location_code=code,
                location_name=loc_name,
                location_state=loc_state,
                header_date=hdr_date,
                invoice_date=source_date(i),
                total_invoice=rupiah(total_invoice),
            )
            blocks.append(current)
            carry_date = hdr_date

        if current is None:
            continue

        row_date = source_date(a, b) or period_start(a)
        if row_date:
            carry_date = row_date

        raw_name = clean(e) or clean(d)
        brand_or_group = clean(d)
        if not raw_name:
            continue

        qty = decimal_number(g)
        price = decimal_number(h)
        item_name = PRODUCT_ALIASES.get(upper(raw_name), upper(raw_name))
        current.items.append(Item(
            row=row_no,
            received_date=carry_date,
            name=item_name,
            raw_name=upper(raw_name),
            brand_or_group=brand_or_group,
            uom=upper(f),
            qty=qty,
            price=price,
            tax=tax_rupiah(tax),
        ))

    return blocks


def group_invoices(blocks: list[Block]) -> list[InvoiceGroup]:
    groups: list[InvoiceGroup] = []
    active: dict[tuple[str, str | None], InvoiceGroup] = {}

    for block in blocks:
        key = (block.supplier_key, block.location_code or upper(block.location_raw))
        if block.total_invoice is not None:
            group = InvoiceGroup(
                start_row=block.row,
                supplier_name=block.supplier_name,
                supplier_key=block.supplier_key,
                location_code=block.location_code,
                location_name=block.location_name,
                location_state=block.location_state,
                total_invoice=block.total_invoice,
                blocks=[block],
            )
            groups.append(group)
            active[key] = group
            continue

        group = active.get(key)
        if group is None:
            group = InvoiceGroup(
                start_row=block.row,
                supplier_name=block.supplier_name,
                supplier_key=block.supplier_key,
                location_code=block.location_code,
                location_name=block.location_name,
                location_state=block.location_state,
                total_invoice=None,
                blocks=[block],
            )
            groups.append(group)
            active[key] = group
        else:
            group.blocks.append(block)

    return groups


def treatment_ready(product: Product) -> tuple[bool, str]:
    treatment = str(getattr(product, "inventory_treatment", "REVIEW_REQUIRED") or "REVIEW_REQUIRED").upper()
    if treatment == "REVIEW_REQUIRED":
        return False, "TREATMENT_REVIEW"
    if treatment == "STOCK" and not product.inventory_account_id:
        return False, "MISSING_INVENTORY_COA"
    if treatment in {"DIRECT_USE", "NON_STOCK"} and not product.purchase_account_id:
        return False, "MISSING_EXPENSE_COA"
    if treatment == "FIXED_ASSET" and not getattr(product, "asset_account_id", None):
        return False, "MISSING_ASSET_COA"
    return True, treatment


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--from-row", type=int, default=2)
    parser.add_argument("--to-row", type=int)
    parser.add_argument("--supplier")
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=False, data_only=True)
    ws = workbook[workbook.sheetnames[0]]
    blocks = parse_blocks(ws)
    groups = group_invoices(blocks)

    supplier_filter = supplier_key(args.supplier) if args.supplier else None
    groups = [
        group for group in groups
        if group.start_row >= args.from_row
        and (args.to_row is None or group.start_row <= args.to_row)
        and (supplier_filter is None or group.supplier_key == supplier_filter)
    ]

    db = SessionLocal()
    try:
        products = {upper(row.name): row for row in db.query(Product).all()}
        suppliers = {supplier_key(row.name): row for row in db.query(Supplier).all()}
        locations_by_code = {upper(row.code): row for row in db.query(Location).all() if row.code}

        summary = defaultdict(int)
        print("=== RESIPING 2026 V2 PREPARE ===")
        print(f"Workbook: {args.workbook}")
        print(f"Invoice/receipt groups in scope: {len(groups)}")
        print()

        for group in groups:
            summary["groups"] += 1
            rows = ",".join(str(block.row) for block in group.blocks)
            location = locations_by_code.get(upper(group.location_code)) if group.location_code else None

            if group.location_state == "SKIP_OPERATIONAL":
                summary["skip_operational"] += 1
                print(f"[SKIP OPERATIONAL] row {group.start_row} {group.supplier_name} blocks={rows}")
                continue

            location_ok = bool(location) if group.location_code else False
            supplier_exists = group.supplier_key in suppliers
            received = group.received_items
            not_received = group.not_received_items
            missing_products: list[str] = []
            treatment_issues: list[str] = []

            for item in received:
                product = products.get(item.name)
                if not product:
                    if item.name not in missing_products:
                        missing_products.append(item.name)
                    continue
                ready, state = treatment_ready(product)
                if not ready:
                    label = f"{product.code} {product.name}: {state}"
                    if label not in treatment_issues:
                        treatment_issues.append(label)

            tax = group.source_tax
            subtotal = group.receiving_subtotal
            invoice_total = group.total_invoice
            variance = invoice_total - subtotal - tax if invoice_total is not None else None

            status = "READY"
            reasons: list[str] = []
            if not location_ok:
                status = "BLOCK"
                reasons.append(f"location={group.location_raw if group.blocks else ''}")
            if missing_products:
                status = "BLOCK"
                reasons.append(f"new_products={len(missing_products)}")
            if treatment_issues:
                status = "BLOCK"
                reasons.append(f"item_master_review={len(treatment_issues)}")
            if invoice_total is None:
                reasons.append("receipt_only_no_TOTAL_INVOICE")

            if status == "READY":
                summary["ready"] += 1
            else:
                summary["blocked"] += 1
            summary["received_lines"] += len(received)
            summary["not_received_lines"] += len(not_received)
            summary["missing_products"] += len(missing_products)
            summary["treatment_issues"] += len(treatment_issues)

            print(
                f"[{status}] row {group.start_row} {group.supplier_name} "
                f"loc={group.location_name or group.location_raw or 'MISSING'} blocks={rows} "
                f"received={len(received)} not_received={len(not_received)} "
                f"TOTAL_INVOICE={invoice_total if invoice_total is not None else 'NONE'} "
                f"recv_subtotal={subtotal} tax={tax} variance={variance if variance is not None else 'N/A'}"
            )
            if reasons:
                print("  reasons:", "; ".join(reasons))
            if not supplier_exists:
                print("  supplier: NEW MASTER")
            if missing_products:
                print("  products needing master/review:")
                for name in missing_products[:25]:
                    print("   -", name)
                if len(missing_products) > 25:
                    print(f"   ... {len(missing_products)-25} more")
            if treatment_issues:
                print("  item master blockers:")
                for issue in treatment_issues[:25]:
                    print("   -", issue)
                if len(treatment_issues) > 25:
                    print(f"   ... {len(treatment_issues)-25} more")
            if not_received:
                print("  not received (no positive qty + price):")
                for item in not_received[:20]:
                    print(f"   - row {item.row}: {item.raw_name} qty={item.qty} price={item.price}")
                if len(not_received) > 20:
                    print(f"   ... {len(not_received)-20} more")
            print()

        print("=== SUMMARY ===")
        for key in (
            "groups", "ready", "blocked", "skip_operational", "received_lines",
            "not_received_lines", "missing_products", "treatment_issues",
        ):
            print(f"{key}: {summary[key]}")
        print("No database changes were made.")
    finally:
        db.close()
        workbook.close()


if __name__ == "__main__":
    main()
