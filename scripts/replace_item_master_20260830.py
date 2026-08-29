#!/usr/bin/env python3
"""Replace the active Item Master from HMM-Master-Products-Cleaned-20260830.xlsx.

Safe replacement strategy:
- Workbook Products sheet is the canonical active master.
- Existing canonical codes are updated in-place so their database IDs and
  accounting/supplier links survive.
- Merge Map is used to repoint every FK that references products.id from a
  duplicate/source product to its canonical/master product.
- Inventory collisions are consolidated per warehouse (quantity summed,
  average cost recalculated).
- Non-canonical products are deleted only after all product FKs are clear.
- Existing preferred supplier / direct COA mappings on retained canonical
  products are preserved because the workbook intentionally contains no COA IDs.
- One database transaction is used. Dry-run is default; --apply commits.

This script is intended to run on the EC2 worker/admin environment where
openpyxl is available. It is not packaged into Lambda.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import inspect, text
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.inventory import Inventory
from app.models.item_group import ItemGroup
from app.models.item_subgroup import ItemSubgroup
from app.models.product import Product

EXPECTED_CANONICAL_PRODUCTS = 1148
EXPECTED_SOURCE_ROWS = 1341
PRODUCTS_SHEET = "Products"
MERGE_SHEET = "Merge Map"

PRODUCT_HEADERS = [
    "Code", "Name", "Group Code", "Group Name", "Subgroup Code", "Subgroup Name",
    "Purchase UOM", "Usage UOM", "Conversion", "Cost Price", "Selling Price",
    "Minimum Stock", "Reorder Level", "Maximum Stock", "Lead Time Days",
    "Tax Category", "Taxable", "Tax Rate", "Active",
]
MERGE_HEADERS = [
    "Source Code", "Source Name", "Master Code", "Canonical Name", "Action",
]

ZERO = Decimal("0")


def norm_text(value) -> str:
    return " ".join(str(value or "").strip().split())


def upper(value) -> str:
    return norm_text(value).upper()


def dec(value, field: str) -> Decimal:
    if value in (None, ""):
        return ZERO
    try:
        return Decimal(str(value))
    except InvalidOperation as exc:
        raise RuntimeError(f"Invalid decimal in {field}: {value!r}") from exc


def integer(value, field: str) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"Invalid integer in {field}: {value!r}") from exc


def boolean(value) -> bool:
    if isinstance(value, bool):
        return value
    token = upper(value)
    if token in {"TRUE", "1", "YES", "Y"}:
        return True
    if token in {"FALSE", "0", "NO", "N", ""}:
        return False
    raise RuntimeError(f"Invalid boolean value: {value!r}")


def sheet_rows(ws, required_headers: list[str]) -> list[dict]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    header_map = {
        norm_text(name): idx
        for idx, name in enumerate(header_row)
        if norm_text(name)
    }
    missing = [name for name in required_headers if name not in header_map]
    if missing:
        raise RuntimeError(f"{ws.title}: missing required headers: {missing}")

    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        row = {name: raw[header_map[name]] for name in required_headers}
        if not any(value not in (None, "") for value in row.values()):
            continue
        rows.append(row)
    return rows


def load_workbook_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if PRODUCTS_SHEET not in wb.sheetnames or MERGE_SHEET not in wb.sheetnames:
        raise RuntimeError(
            f"Workbook must contain sheets {PRODUCTS_SHEET!r} and {MERGE_SHEET!r}; "
            f"found {wb.sheetnames}"
        )

    products = sheet_rows(wb[PRODUCTS_SHEET], PRODUCT_HEADERS)
    merge_rows = sheet_rows(wb[MERGE_SHEET], MERGE_HEADERS)

    if len(products) != EXPECTED_CANONICAL_PRODUCTS:
        raise RuntimeError(
            f"Canonical product count changed: expected {EXPECTED_CANONICAL_PRODUCTS}, "
            f"got {len(products)}"
        )

    codes = [upper(row["Code"]) for row in products]
    names = [upper(row["Name"]) for row in products]
    if "" in codes or "" in names:
        raise RuntimeError("Products sheet contains blank Code or Name")
    if len(codes) != len(set(codes)):
        raise RuntimeError("Products sheet contains duplicate canonical Code values")
    if len(names) != len(set(names)):
        duplicates = sorted({name for name in names if names.count(name) > 1})
        raise RuntimeError(f"Products sheet contains duplicate canonical Name values: {duplicates[:10]}")

    source_codes = {upper(row["Source Code"]) for row in merge_rows if upper(row["Source Code"])}
    if len(source_codes) != EXPECTED_SOURCE_ROWS:
        raise RuntimeError(
            f"Merge Map source count changed: expected {EXPECTED_SOURCE_ROWS}, got {len(source_codes)}"
        )

    master_codes = set(codes)
    bad_targets = sorted({
        upper(row["Master Code"])
        for row in merge_rows
        if upper(row["Master Code"]) and upper(row["Master Code"]) not in master_codes
    })
    if bad_targets:
        raise RuntimeError(f"Merge Map references unknown master codes: {bad_targets[:10]}")

    return products, merge_rows


def ensure_classification(db: Session, products: list[dict]) -> None:
    groups: dict[str, tuple[str, set[tuple[str, str]]]] = {}
    for row in products:
        group_code = upper(row["Group Code"])
        group_name = upper(row["Group Name"])
        subgroup_code = upper(row["Subgroup Code"])
        subgroup_name = upper(row["Subgroup Name"])
        if not all((group_code, group_name, subgroup_code, subgroup_name)):
            raise RuntimeError(f"Incomplete classification for product {row['Code']}")
        groups.setdefault(group_code, (group_name, set()))[1].add((subgroup_code, subgroup_name))

    for group_code, (group_name, subgroups) in groups.items():
        group = db.query(ItemGroup).filter(ItemGroup.code == group_code).one_or_none()
        if group is None:
            group = ItemGroup(code=group_code, name=group_name, is_active=True)
            db.add(group)
            db.flush()
        else:
            group.name = group_name
            group.is_active = True

        for subgroup_code, subgroup_name in subgroups:
            subgroup = db.query(ItemSubgroup).filter(ItemSubgroup.code == subgroup_code).one_or_none()
            if subgroup is None:
                db.add(ItemSubgroup(
                    code=subgroup_code,
                    name=subgroup_name,
                    group_id=group.id,
                    is_active=True,
                ))
            else:
                subgroup.name = subgroup_name
                subgroup.group_id = group.id
                subgroup.is_active = True
    db.flush()


def tax_reference(category: str) -> tuple[str | None, str | None]:
    if category == "EXEMPT":
        return "Cleaned Item Master 20260830", "PP 49 Tahun 2022"
    if category == "STANDARD":
        return "Cleaned Item Master 20260830", "PMK 131 Tahun 2024"
    if category == "REVIEW_REQUIRED":
        return "Cleaned Item Master 20260830", "Finance review required"
    return "Cleaned Item Master 20260830", None


def canonical_values(row: dict) -> dict:
    category = upper(row["Tax Category"]) or "REVIEW_REQUIRED"
    taxable = boolean(row["Taxable"])
    rate = dec(row["Tax Rate"], "Tax Rate") if taxable else ZERO
    if category in {"EXEMPT", "NON_VAT", "REVIEW_REQUIRED"}:
        taxable = False
        rate = ZERO
    reason, reference = tax_reference(category)
    subgroup_name = upper(row["Subgroup Name"])
    return {
        "code": upper(row["Code"]),
        "name": upper(row["Name"]),
        "group_code": upper(row["Group Code"]),
        "group_name": upper(row["Group Name"]),
        "subgroup_code": upper(row["Subgroup Code"]),
        "subgroup_name": subgroup_name,
        "uom": norm_text(row["Purchase UOM"]) or "Pcs",
        "usage_uom": norm_text(row["Usage UOM"]) or norm_text(row["Purchase UOM"]) or "Pcs",
        "conversion": dec(row["Conversion"], "Conversion") or Decimal("1"),
        "cost_price": dec(row["Cost Price"], "Cost Price"),
        "selling_price": dec(row["Selling Price"], "Selling Price"),
        "minimum_stock": dec(row["Minimum Stock"], "Minimum Stock"),
        "reorder_level": dec(row["Reorder Level"], "Reorder Level"),
        "maximum_stock": dec(row["Maximum Stock"], "Maximum Stock"),
        "lead_time_days": integer(row["Lead Time Days"], "Lead Time Days"),
        "tax_category": category,
        "is_taxable": taxable,
        "tax_rate": rate,
        "tax_reason": reason,
        "tax_reference": reference,
        "is_market_item": subgroup_name in {"VEGETABLES", "FRUIT"},
        "is_active": boolean(row["Active"]),
    }


def upsert_canonical_products(db: Session, rows: list[dict], stats: Counter) -> dict[str, Product]:
    existing = {
        upper(product.code): product
        for product in db.query(Product).order_by(Product.id).all()
    }
    canonical: dict[str, Product] = {}

    for row in rows:
        values = canonical_values(row)
        code = values["code"]
        product = existing.get(code)
        if product is None:
            product = Product(**values)
            db.add(product)
            db.flush()
            stats["created_canonical"] += 1
        else:
            for field, value in values.items():
                setattr(product, field, value)
            # Intentionally preserve these current mappings because the cleaned
            # workbook has no supplier/COA IDs:
            # preferred_supplier_id, inventory_account_id,
            # purchase_account_id, cogs_account_id.
            stats["updated_canonical"] += 1
        canonical[code] = product

    db.flush()
    return canonical


def consolidate_inventory(db: Session, source_id: int, target_id: int, stats: Counter) -> None:
    rows = db.query(Inventory).filter(Inventory.product_id == source_id).all()
    for source in rows:
        target = (
            db.query(Inventory)
            .filter(
                Inventory.product_id == target_id,
                Inventory.warehouse_id == source.warehouse_id,
            )
            .one_or_none()
        )
        if target is None:
            source.product_id = target_id
            stats["inventory_repointed"] += 1
            continue

        source_qty = Decimal(source.quantity or 0)
        target_qty = Decimal(target.quantity or 0)
        total_qty = source_qty + target_qty
        if total_qty:
            total_value = (
                source_qty * Decimal(source.average_cost or 0)
                + target_qty * Decimal(target.average_cost or 0)
            )
            target.average_cost = total_value / total_qty
        target.quantity = total_qty
        db.delete(source)
        stats["inventory_consolidated"] += 1
    db.flush()


def product_fk_columns(db: Session) -> list[tuple[str, str]]:
    inspector = inspect(db.get_bind())
    refs = []
    for table_name in inspector.get_table_names():
        for fk in inspector.get_foreign_keys(table_name):
            if fk.get("referred_table") != "products":
                continue
            constrained = fk.get("constrained_columns") or []
            referred = fk.get("referred_columns") or []
            for column, remote in zip(constrained, referred):
                if remote == "id":
                    refs.append((table_name, column))
    return sorted(set(refs))


def count_fk_refs(db: Session, product_id: int, refs: list[tuple[str, str]]) -> dict[str, int]:
    counts = {}
    for table_name, column_name in refs:
        if table_name == "inventory":
            continue
        count = db.execute(
            text(f'SELECT count(*) FROM "{table_name}" WHERE "{column_name}" = :pid'),
            {"pid": product_id},
        ).scalar_one()
        if count:
            counts[f"{table_name}.{column_name}"] = int(count)
    return counts


def repoint_generic_fks(
    db: Session,
    source_id: int,
    target_id: int,
    refs: list[tuple[str, str]],
    stats: Counter,
) -> None:
    for table_name, column_name in refs:
        if table_name == "inventory":
            continue
        result = db.execute(
            text(
                f'UPDATE "{table_name}" '
                f'SET "{column_name}" = :target '
                f'WHERE "{column_name}" = :source'
            ),
            {"target": target_id, "source": source_id},
        )
        if result.rowcount:
            stats["fk_rows_repointed"] += int(result.rowcount)
    db.flush()


def inherit_accounting_links(source: Product, target: Product) -> None:
    for field in (
        "preferred_supplier_id",
        "inventory_account_id",
        "purchase_account_id",
        "cogs_account_id",
    ):
        if getattr(target, field) is None and getattr(source, field) is not None:
            setattr(target, field, getattr(source, field))


def apply_merge_map(
    db: Session,
    merge_rows: list[dict],
    canonical: dict[str, Product],
    refs: list[tuple[str, str]],
    stats: Counter,
) -> set[int]:
    existing = {
        upper(product.code): product
        for product in db.query(Product).order_by(Product.id).all()
    }
    merged_source_ids: set[int] = set()
    seen_sources: set[str] = set()

    for row in merge_rows:
        source_code = upper(row["Source Code"])
        master_code = upper(row["Master Code"])
        if not source_code or not master_code or source_code in seen_sources:
            continue
        seen_sources.add(source_code)

        source = existing.get(source_code)
        target = canonical[master_code]
        if source is None or source.id == target.id:
            continue

        inherit_accounting_links(source, target)
        consolidate_inventory(db, source.id, target.id, stats)
        repoint_generic_fks(db, source.id, target.id, refs, stats)
        merged_source_ids.add(source.id)
        stats["duplicate_products_merged"] += 1

    db.flush()
    return merged_source_ids


def remove_noncanonical(
    db: Session,
    canonical_codes: set[str],
    refs: list[tuple[str, str]],
    stats: Counter,
) -> None:
    noncanonical = [
        product
        for product in db.query(Product).order_by(Product.id).all()
        if upper(product.code) not in canonical_codes
    ]

    blockers = []
    for product in noncanonical:
        inventory_count = db.query(Inventory).filter(Inventory.product_id == product.id).count()
        ref_counts = count_fk_refs(db, product.id, refs)
        if inventory_count:
            ref_counts["inventory.product_id"] = inventory_count
        if ref_counts:
            blockers.append({
                "id": product.id,
                "code": product.code,
                "name": product.name,
                "references": ref_counts,
            })

    if blockers:
        sample = json.dumps(blockers[:20], indent=2, default=str)
        raise RuntimeError(
            "Cannot fully replace Item Master because non-canonical products still "
            "have references not resolved by Merge Map. No changes were committed. "
            f"First blockers:\n{sample}"
        )

    for product in noncanonical:
        db.delete(product)
        stats["deleted_noncanonical"] += 1
    db.flush()


def verify_final(db: Session, canonical_codes: set[str]) -> None:
    active_codes = {
        upper(code)
        for (code,) in db.query(Product.code).filter(Product.is_active.is_(True)).all()
    }
    if active_codes != canonical_codes:
        missing = sorted(canonical_codes - active_codes)
        extra = sorted(active_codes - canonical_codes)
        raise RuntimeError(
            f"Final active Item Master mismatch. missing={missing[:10]} extra={extra[:10]}"
        )

    duplicates = db.execute(text(
        """
        SELECT lower(trim(name)) AS normalized_name, count(*)
        FROM products
        WHERE is_active = true
        GROUP BY lower(trim(name))
        HAVING count(*) > 1
        LIMIT 10
        """
    )).all()
    if duplicates:
        raise RuntimeError(f"Final active Item Master still contains duplicate names: {duplicates}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default="/opt/hmm-connect/imports/HMM-Master-Products-Cleaned-20260830.xlsx",
        help="Path to cleaned workbook on EC2",
    )
    parser.add_argument("--apply", action="store_true", help="Commit replacement")
    args = parser.parse_args()

    path = Path(args.input)
    if not path.is_file():
        raise RuntimeError(f"Workbook not found: {path}")

    products, merge_rows = load_workbook_rows(path)
    canonical_codes = {upper(row["Code"]) for row in products}
    stats = Counter()

    db = SessionLocal()
    try:
        print(f"Input workbook: {path}")
        print(f"Canonical products: {len(products)}")
        print(f"Merge Map rows: {len(merge_rows)}")
        print(f"Existing products before seed: {db.query(Product).count()}")

        ensure_classification(db, products)
        canonical = upsert_canonical_products(db, products, stats)
        refs = product_fk_columns(db)
        print(f"Detected product FK columns: {len(refs)}")
        apply_merge_map(db, merge_rows, canonical, refs, stats)
        remove_noncanonical(db, canonical_codes, refs, stats)
        verify_final(db, canonical_codes)

        final_total = db.query(Product).count()
        final_active = db.query(Product).filter(Product.is_active.is_(True)).count()
        print("Item Master replacement summary:")
        for key in sorted(stats):
            print(f"  {key}={stats[key]}")
        print(f"  final_total={final_total}")
        print(f"  final_active={final_active}")

        if args.apply:
            db.commit()
            print("Item Master replacement committed successfully.")
        else:
            db.rollback()
            print("Dry-run successful; transaction rolled back. Re-run with --apply to commit.")
        return 0
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
