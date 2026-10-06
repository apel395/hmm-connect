#!/usr/bin/env python3
"""Prepare BUY item identities from 2026 receiving rows in RESIPING XLSM/XLSX.

Read-only. Formula/macro execution, item creation and transaction seeding are
not performed. --db-check matches existing active/inactive BUY items exactly.
"""
import argparse
from collections import Counter
from datetime import date, datetime
from difflib import get_close_matches
from hashlib import sha256
from html.parser import HTMLParser
from http.cookiejar import CookieJar
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import urlencode, urlparse
from urllib.request import HTTPCookieProcessor, Request, build_opener
from zipfile import is_zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


MAX_SOURCE_BYTES = 30 * 1024 * 1024


class DownloadConfirmation(HTMLParser):
    """Read Google's download form without executing HTML or JavaScript."""
    def __init__(self):
        super().__init__()
        self.action = None
        self.fields = {}
        self.in_form = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'form':
            target = urlparse(attrs.get('action',''))
            self.in_form = (target.scheme == 'https' and target.netloc == 'drive.usercontent.google.com'
                            and target.path == '/download' and attrs.get('method','get').lower() == 'get')
            if self.in_form:
                self.action = attrs['action']
        elif tag == 'input' and self.in_form and attrs.get('type') == 'hidden':
            if attrs.get('name') in ('id','export','confirm','uuid'):
                self.fields[attrs['name']] = attrs.get('value','')

    def handle_endtag(self, tag):
        if tag == 'form':
            self.in_form = False


def download_public_drive(drive_id, opener=None):
    if not re.fullmatch(r'[A-Za-z0-9_-]{10,200}',drive_id):
        raise ValueError('Use the exact Drive file ID, not a URL')
    opener = opener or build_opener(HTTPCookieProcessor(CookieJar()))
    url = 'https://drive.google.com/uc?' + urlencode({'export':'download','id':drive_id})
    for attempt in range(2):
        try:
            with opener.open(Request(url,headers={'User-Agent':'HMM-Connect-item-source/1.0'}),timeout=60) as response:
                content = response.read(MAX_SOURCE_BYTES + 1)
        except OSError as exc:
            raise ValueError('Public Drive download failed. Check sharing and EC2 network access.') from exc
        if len(content) > MAX_SOURCE_BYTES:
            raise ValueError('Drive workbook exceeds 30 MiB limit')
        if is_zipfile(BytesIO(content)):
            return content
        confirmation = DownloadConfirmation()
        confirmation.feed(content.decode('utf-8',errors='replace'))
        if attempt == 0 and confirmation.action and confirmation.fields.get('id') == drive_id:
            url = confirmation.action + '?' + urlencode(confirmation.fields)
        else:
            raise ValueError('Drive returned a sign-in/error page instead of an Excel workbook. Check anyone-with-link access and download permission.')
    raise ValueError('Drive workbook download did not complete')


def load_source(file=None, drive_id=None):
    if drive_id:
        content = download_public_drive(drive_id)
        with tempfile.TemporaryDirectory(prefix='hmm-item-source-') as directory:
            path = Path(directory) / 'drive-source.xlsm'
            path.write_bytes(content)
            report = extract(path)
        report['source_file'] = 'Google Drive Excel file'
        report['source_drive_id'] = drive_id
        return report
    if file is None:
        raise ValueError('Provide --file or --drive-id')
    if file.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError('Workbook exceeds 30 MiB limit')
    return extract(file)


def normalized(value):
    return ' '.join(str(value or '').strip().upper().split())


def source_year(value):
    if isinstance(value, (date, datetime)):
        return value.year
    text = normalized(value)
    # Only explicit years qualify. Unyearred ranges cannot change scope.
    four = set(int(year) for year in re.findall(r'(?<!\d)(20\d{2})(?!\d)', text))
    if len(four) == 1:
        return next(iter(four))
    if four:
        return None
    if not re.search(r'[A-Z]',text) and not re.fullmatch(r'\d{1,2}[-/]\d{1,2}[-/]\d{2}',text):
        return None
    match = re.search(r'(?:[-/\s])(25|26)\s*$', text)
    return 2000 + int(match.group(1)) if match else None


def invoice_id(value, row, first_item=False):
    text = str(value or '').strip()
    if re.fullmatch(r'INV-\d{6}', text):
        return text
    # Decode only this template's row-number ID; never evaluate arbitrary Excel.
    expected = [f'=IF(COUNTA({left}{row}:{right}{row})=0,"","INV-"&TEXT(ROW()-4,"000000"))' for left,right in [('B','G'),('C','I')]]
    if text in expected:
        if not first_item or row == 5:
            return f'INV-{row-4:06d}'
    return None


def extract(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if not {'INVOICES','ITEM_ENTRY'}.issubset(workbook.sheetnames):
            raise ValueError('Expected INVOICES and ITEM_ENTRY tabs')
        invoice_headers = list(next(workbook['INVOICES'].iter_rows(min_row=4,max_row=4,max_col=12,values_only=True)))
        item_headers = list(next(workbook['ITEM_ENTRY'].iter_rows(min_row=4,max_row=4,max_col=14,values_only=True)))
        if invoice_headers[:4] != ['Invoice ID','Invoice Date','Supplier','Location'] or item_headers[:6] != ['Start Invoice ID','tanggal','Item','Unit','Qty','Unit Price']:
            raise ValueError('Workbook columns changed; update parser before using source')
        invoices = {}
        for row, cells in enumerate(workbook['INVOICES'].iter_rows(min_row=5,max_col=12,values_only=True),5):
            if not any(cells[1:7]):
                continue
            key = invoice_id(cells[0],row)
            if not key or key in invoices:
                raise ValueError(f'Invoice row {row}: missing, unsupported or duplicate ID')
            invoices[key] = {'row':row,'raw_date':cells[1], 'supplier':cells[2], 'location':cells[3]}
        items, exclusions, review = {}, [], []
        current_id = None
        current_date = None
        counts = Counter()
        for row, cells in enumerate(workbook['ITEM_ENTRY'].iter_rows(min_row=5,max_col=14,values_only=True),5):
            if not any(cells[:8]):
                current_id = None; current_date = None
                continue
            if not cells[2]:
                review.append({'row':row,'reason':'Missing item name'}); continue
            if cells[0] is not None:
                current_id = invoice_id(cells[0],row,first_item=True)
                current_date = None
            if not current_id or current_id not in invoices:
                review.append({'row':row,'item':str(cells[2]),'reason':'Missing or unsupported invoice link'}); continue
            header = invoices[current_id]
            if cells[1] is not None:
                current_date = cells[1]
            raw_date = current_date if current_date is not None else header['raw_date']
            year = source_year(raw_date)
            evidence = {'row':row,'invoice_id':current_id,'raw_receiving_date':raw_date,'date_source':'item/carried' if current_date is not None else 'invoice period', 'supplier':header['supplier'],'location':header['location']}
            if year is not None and year < 2026:
                exclusions.append({**evidence,'item':str(cells[2]),'reason':'Pre-2026 receiving'});counts['excluded_pre_2026_rows'] += 1;continue
            if year != 2026:
                review.append({**evidence,'item':str(cells[2]),'reason':'Receiving year missing/ambiguous or outside 2026'});continue
            if not cells[3] or str(cells[2]).startswith('=') or str(cells[3]).startswith('='):
                review.append({**evidence,'item':str(cells[2]),'reason':'Item/unit must be literal and populated'});continue
            name, unit = normalized(cells[2]), normalized(cells[3])
            item = items.setdefault(name, {'name':name,'source_names':set(),'units':set(),'source_rows':[],'item_type':'BUY'})
            item['source_names'].add(str(cells[2]).strip());item['units'].add(unit)
            if not isinstance(raw_date, (datetime,date)):
                evidence['receiving_date_requires_review'] = True
            if source_year(header['raw_date']) not in (None,year):
                evidence['invoice_receiving_year_mismatch'] = True
            item['source_rows'].append(evidence);counts['included_2026_rows'] += 1
        for item in items.values():
            item['source_names'] = sorted(item['source_names']);item['units'] = sorted(item['units'])
            item['unit_review_required'] = len(item['units']) != 1
        return {'source_sha256':sha256(path.read_bytes()).hexdigest(),'source_file':path.name,
                'mode':'PREPARATION_ONLY','cutoff':'2026-01-01','invoice_headers':len(invoices),
                'counts':{**counts,'unique_2026_items':len(items),'date_or_link_review_rows':len(review),
                          'multiple_unit_items':sum(item['unit_review_required'] for item in items.values())},
                'items':list(items.values()),'excluded_rows':exclusions,'review_rows':review,
                'database_checked':False,'seed_ready':False}
    finally:
        workbook.close()


def match_existing(items, products):
    by_name = {}
    for product in products:
        if product['item_type'] == 'BUY':
            by_name.setdefault(normalized(product['name']),[]).append(product)
    counts = Counter()
    for item in items:
        matches = by_name.get(item['name'],[])
        if len(matches) > 1:
            item['match_status'] = 'AMBIGUOUS_BUY_NAME'
            item['existing_ids'] = [p['id'] for p in matches]
        elif not matches:
            similar = get_close_matches(item['name'], list(by_name), n=3, cutoff=0.86)
            slash_aliases = [part.strip() for part in item['name'].split('/') if part.strip() in by_name]
            similar = sorted(set(similar + slash_aliases))
            if similar:
                item['match_status'] = 'POSSIBLE_BUY_ALIAS_REVIEW'
                item['possible_existing_items'] = [{'id':p['id'],'code':p['code'],'name':p['name']} for name in similar for p in by_name[name]]
            elif item['unit_review_required']:
                item['match_status'] = 'UNIT_REVIEW'
            else:
                item['match_status'] = 'NEW_BUY_ITEM_REQUIRES_CLASSIFICATION'
        else:
            product = matches[0]
            item['existing_product'] = dict(product)
            if not product['is_active']:
                item['match_status'] = 'INACTIVE_BUY_ITEM_REVIEW'
            elif item['unit_review_required'] or normalized(product['uom']) not in item['units']:
                item['match_status'] = 'UNIT_REVIEW'
            else:
                item['match_status'] = 'EXISTING_BUY_ITEM'
        counts[item['match_status']] += 1
    return dict(counts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument('--file',type=Path)
    source.add_argument('--drive-id',help='Public Google Drive Excel file ID')
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--db-check',action='store_true')
    args = parser.parse_args()
    report = load_source(args.file,args.drive_id)
    if args.db_check:
        from sqlalchemy import text
        from app.core.database import engine
        with engine.connect() as conn:
            with conn.begin():
                conn.execute(text('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY'))
                conn.execute(text("SET LOCAL statement_timeout='30s'"))
                products = [dict(row) for row in conn.execute(text('''
                    SELECT id,code,name,item_type,is_active,uom,usage_uom,conversion,
                           group_code,group_name,subgroup_code,subgroup_name,inventory_treatment,
                           inventory_account_id,purchase_account_id,cogs_account_id,asset_account_id,
                           is_taxable,tax_rate,tax_category
                    FROM products ORDER BY id
                ''')).mappings()]
                report['db_matches'] = match_existing(report['items'],products)
                report['database_checked'] = True
    args.output.write_text(json.dumps(report,indent=2,default=str,ensure_ascii=False)+'\n')
    print(json.dumps({key:value for key,value in report.items() if key not in ('items','excluded_rows','review_rows')},indent=2,default=str,ensure_ascii=False))
    print(f'Full item preparation: {args.output}')
    print('READ ONLY: no master, stock, receiving, invoice, tax, journal or opening balance writes.')


if __name__ == '__main__':
    main()
