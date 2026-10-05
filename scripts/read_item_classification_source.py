#!/usr/bin/env python3
"""Read the authoritative local classification workbook. No database writes.

Default source: item-classification-review-237.xlsx beside this script.
Local override: --file scripts/item-classification-review-237.xlsx
Public native Sheet: --drive-id FILE_ID
Private Drive file/Sheet: --drive-id FILE_ID --credentials /secure/service-account.json
Private access requires google-auth and sharing the file with the service account.
"""
import argparse
from collections import Counter
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from zipfile import is_zipfile

MAX_BYTES = 20 * 1024 * 1024
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def fetch(url, token=None):
    headers = {"Authorization": "Bearer " + token} if token else {}
    try:
        with urlopen(Request(url, headers=headers), timeout=60) as response:
            content = response.read(MAX_BYTES + 1)
    except HTTPError as exc:
        raise ValueError(f"Google returned HTTP {exc.code}. Check sharing and credentials.") from None
    except URLError:
        raise ValueError("Could not reach Google. Check EC2 network access.") from None
    if len(content) > MAX_BYTES:
        raise ValueError("Source exceeds the 20 MiB reader limit.")
    return content


def download(drive_id, credentials_path):
    if not re.fullmatch(r"[A-Za-z0-9_-]+", drive_id):
        raise ValueError("Use the Drive file ID, not a full URL.")
    if not credentials_path:
        return fetch(f"https://docs.google.com/spreadsheets/d/{drive_id}/export?format=xlsx")
    try:
        from google.oauth2 import service_account
        from google.auth.transport.requests import Request as AuthRequest
    except ImportError:
        raise ValueError("Private Drive access requires google-auth and requests in .venv.") from None
    try:
        credentials = service_account.Credentials.from_service_account_file(
            credentials_path,
            scopes=["https://www.googleapis.com/auth/drive.readonly"],
        )
        credentials.refresh(AuthRequest())
    except Exception:
        raise ValueError("Service-account authentication failed. Check the credential file.") from None
    base = f"https://www.googleapis.com/drive/v3/files/{drive_id}"
    metadata = json.loads(fetch(base + "?fields=id,mimeType&supportsAllDrives=true", credentials.token))
    mime = metadata.get("mimeType")
    if mime == "application/vnd.google-apps.spreadsheet":
        url = base + "/export?" + urlencode({"mimeType": XLSX_MIME})
    elif mime == XLSX_MIME:
        url = base + "?alt=media&supportsAllDrives=true"
    else:
        raise ValueError("The Drive file must be a native Google Sheet or XLSX workbook.")
    return fetch(url, credentials.token)


def inspect(content, sheet_name, expected_count):
    if not is_zipfile(BytesIO(content)):
        raise ValueError("Source is not XLSX. A Google sign-in page is not a spreadsheet.")
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("The reader requires openpyxl in .venv.") from None
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception:
        raise ValueError("Source could not be opened as an XLSX workbook.") from None
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Missing review tab {sheet_name!r}. Available tabs: {wb.sheetnames}")
        ws = wb[sheet_name]
        required = ["ID", "Code", "Item", "Treatment recommendation",
                    "Primary COA recommendation", "Tax decision", "Finding", "Stock qty"]
        header = None
        for number, cells in enumerate(ws.iter_rows(max_row=40, values_only=True), 1):
            names = [str(v).strip() if v is not None else "" for v in cells]
            if set(required).issubset(names):
                header = (number, names)
                break
        if header is None:
            raise ValueError("Review headers were not found within the first 40 rows.")
        number, names = header
        populated_headers = [n for n in names if n]
        if len(populated_headers) != len(set(populated_headers)):
            raise ValueError("Duplicate column headers in the review tab.")
        records, ids, codes = [], set(), set()
        for row_number, cells in enumerate(ws.iter_rows(min_row=number + 1, values_only=True), number + 1):
            if not any(v is not None for v in cells):
                continue
            row = {name: cells[i] for i, name in enumerate(names) if name}
            if any(isinstance(v, str) and v.startswith("=") for v in row.values()):
                raise ValueError(f"Row {row_number}: review data contains a formula; literal decisions required.")
            raw_id = row.get("ID")
            if isinstance(raw_id, bool) or not isinstance(raw_id, (int, float)) or raw_id <= 0 or int(raw_id) != raw_id:
                raise ValueError(f"Row {row_number}: invalid product ID.")
            row["ID"] = int(raw_id)
            row["Code"] = str(row.get("Code") or "").strip()
            if not row["Code"] or not str(row.get("Item") or "").strip():
                raise ValueError(f"Row {row_number}: missing code or item name.")
            if row["ID"] in ids or row["Code"] in codes:
                raise ValueError(f"Row {row_number}: duplicate product ID or code.")
            ids.add(row["ID"])
            codes.add(row["Code"])
            records.append(row)
        if len(records) != expected_count:
            raise ValueError(f"Expected {expected_count} items; found {len(records)}. Review scope before seeding.")
        return {
            "sha256": sha256(content).hexdigest(),
            "sheet": sheet_name,
            "item_count": len(records),
            "source_ready": True,
            "source_role": "Authoritative item classification source",
            "findings": dict(Counter(str(r["Finding"]) for r in records)),
            "treatment_recommendations": dict(Counter(str(r["Treatment recommendation"]) for r in records)),
            "tax_decisions": dict(Counter(str(r["Tax decision"]) for r in records)),
            "seed_ready": False,
            "reason": "The XLSX source is valid and accepted. Classification seeding remains pending final treatment, exact COA mappings, tax evidence and a correction plan for affected posted stock. No alternate source spreadsheet is required.",
            "records": records,
        }
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group()
    source.add_argument("--file", type=Path)
    source.add_argument("--drive-id")
    parser.add_argument("--credentials", help="Service-account JSON path; never paste its contents into chat")
    parser.add_argument("--sheet", default="Classification review")
    parser.add_argument("--expected-count", type=int, default=237)
    parser.add_argument("--show-items", action="store_true")
    args = parser.parse_args()
    try:
        if not args.drive_id:
            path = args.file or Path(__file__).resolve().with_name("item-classification-review-237.xlsx")
            if path.stat().st_size > MAX_BYTES:
                raise ValueError("Source exceeds 20 MiB.")
            content = path.read_bytes()
        else:
            content = download(args.drive_id, args.credentials)
        report = inspect(content, args.sheet, args.expected_count)
        report["source_file"] = path.name if not args.drive_id else "Google Drive file " + args.drive_id
        if not args.show_items:
            report.pop("records")
        print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
        print("READ ONLY: no database access, no seeding, no source changes.")
    except (ValueError, OSError) as exc:
        # Do not display filesystem/credential error details that may contain secrets.
        message = str(exc) if isinstance(exc, ValueError) else "Source file could not be read; check path and permissions."
        parser.exit(1, "ERROR: " + message + "\n")


if __name__ == "__main__":
    main()
